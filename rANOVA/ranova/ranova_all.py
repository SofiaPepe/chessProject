import os
import pandas as pd
import pingouin as pg
import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.formatting.rule import CellIsRule

# --- CONFIGURAZIONE ---
dir_base = os.path.dirname(os.path.abspath(__file__))
db_path = os.path.join(dir_base, '../descrittive/database_compositi.xlsx')
out_path = os.path.join(dir_base, 'ranova_e_posthoc.xlsx')

# 1. CARICAMENTO DATI
df = pd.read_excel(db_path)
id_col = 'ID'
cols = list(df.columns)

# 2. IDENTIFICAZIONE VARIABILI
analisi_pairs = []
for c in cols:
    c_upper = c.upper()
    base = None
    suffix_out = ""
    
    if c_upper.endswith('_PRE_LN'):
        base = c[:-7]
        suffix_out = "_ln"
        col_post = next((x for x in cols if x.upper() == (base + "_POST_LN").upper()), None)
    elif c_upper.endswith('_PRE'):
        base = c[:-4]
        suffix_out = ""
        col_post = next((x for x in cols if x.upper() == (base + "_POST").upper()), None)
        
    if base and col_post:
        analisi_pairs.append((base + suffix_out, c, col_post))

analisi_pairs = sorted(list(set(analisi_pairs)))
anova_results = []
posthoc_results = []

# 3. ANALISI
for base_name, col_pre, col_post in analisi_pairs:
    current_cols = [id_col, 'gruppo', col_pre, col_post]
    temp_df = df[current_cols].dropna()
    if temp_df.empty or temp_df['gruppo'].nunique() < 2:
        continue

    long_df = pd.melt(temp_df, id_vars=[id_col, 'gruppo'], value_vars=[col_pre, col_post], 
                      var_name='tempo', value_name='valore')
    long_df['tempo'] = long_df['tempo'].apply(lambda x: 'PRE' if 'PRE' in x.upper() else 'POST')

    try:
        aov = pg.mixed_anova(dv='valore', within='tempo', between='gruppo', subject=id_col, data=long_df)
        aov.insert(0, 'Variabile', base_name)
        anova_results.append(aov)

        p_col = [c for c in aov.columns if 'p-' in c or 'p_unc' in c][0]
        if (aov[p_col] < 0.05).any():
            ph = pg.pairwise_tests(dv='valore', within='tempo', between='gruppo', 
                                    subject=id_col, data=long_df, padjust='bonf')
            ph.insert(0, 'Variabile', base_name)
            posthoc_results.append(ph)
    except Exception as e:
        print(f"Errore su {base_name}: {e}")

# 4. SALVATAGGIO
with pd.ExcelWriter(out_path, engine='openpyxl') as writer:
    pd.concat(anova_results, ignore_index=True).to_excel(writer, sheet_name='ANOVA', index=False)
    if posthoc_results:
        pd.concat(posthoc_results, ignore_index=True).to_excel(writer, sheet_name='POST-HOC', index=False)

# 5. RIPRISTINO FORMATTAZIONE VERDE (CORRETTO)
wb = openpyxl.load_workbook(out_path)
green_fill = PatternFill(start_color='90EE90', end_color='90EE90', fill_type='solid')

# Lista completa di possibili nomi per le colonne p-value
p_targets = ['p-unc', 'p-GG', 'p-corr', 'p-adj', 'p-GG-corr', 'p_unc']

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    for col in ws.iter_cols(1, ws.max_column):
        header = str(col[0].value).strip() # .strip() rimuove eventuali spazi bianchi
        
        # Se il nome della colonna contiene uno dei p_targets o inizia con "p-"
        if any(target in header for target in p_targets) or header.startswith('p-'):
            col_letter = col[0].column_letter
            # Applica la regola: se il valore è < 0.05, colora di verde
            ws.conditional_formatting.add(
                f'{col_letter}2:{col_letter}{ws.max_row}',
                CellIsRule(operator='lessThan', formula=['0.05'], fill=green_fill)
            )
            # Formattazione decimale per leggibilità
            for cell in col[1:]:
                cell.number_format = '0.00000'

wb.save(out_path)
print(f"Fatto! Analisi completata. Il file '{out_path}' ha di nuovo le evidenziazioni verdi.")
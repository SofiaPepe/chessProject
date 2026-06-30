from __future__ import annotations

import hashlib
import math
import re
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from scipy import stats
from statsmodels.stats.multitest import multipletests

from config import (
    GROUP_COLUMN,
    GROUP_LABELS,
    ID_COLUMN,
    INPUT_DATABASE,
    INPUT_SHEET,
    LOG_BASES,
    MINEFIELD_COMPOSITES,
)


def to_numeric(series: pd.Series) -> pd.Series:
    if pd.api.types.is_numeric_dtype(series):
        return pd.to_numeric(series, errors="coerce")
    cleaned = series.astype(str).str.replace(",", ".", regex=False)
    cleaned = cleaned.replace({"nan": np.nan, "None": np.nan, "": np.nan})
    return pd.to_numeric(cleaned, errors="coerce")


def safe_filename(value: str, max_length: int = 120) -> str:
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", str(value)).strip("_")
    return (safe or "item")[:max_length]


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_source_database() -> pd.DataFrame:
    if not INPUT_DATABASE.exists():
        raise FileNotFoundError(f"Input database not found: {INPUT_DATABASE}")
    df = pd.read_excel(INPUT_DATABASE, sheet_name=INPUT_SHEET)
    df.columns = df.columns.astype(str).str.strip()
    df = df.rename(
        columns={
            "PWM_ MATT_MIN_PRE": "PWM_MATT_MIN_PRE",
            "PLANNING_span_POST": "PLANNING_SPAN_POST",
            "PLANNING_trial_POST": "PLANNING_TRIAL_POST",
            "PLANNING_TOT_POST": "PLANNING_TR_POST",
            "PWM_T_POST": "PWM_TR_POST",
        }
    )
    if ID_COLUMN not in df or GROUP_COLUMN not in df:
        raise KeyError(f"Required columns not found: {ID_COLUMN}, {GROUP_COLUMN}")
    if df[ID_COLUMN].duplicated().any():
        duplicates = df.loc[df[ID_COLUMN].duplicated(), ID_COLUMN].tolist()
        raise ValueError(f"Duplicate participant IDs: {duplicates}")
    df[GROUP_COLUMN] = to_numeric(df[GROUP_COLUMN]).astype("Int64")
    df["group_label"] = df[GROUP_COLUMN].map(GROUP_LABELS)
    df["group_experimental"] = (df[GROUP_COLUMN] == 1).astype("Int64")
    return df


def add_questionnaire_totals(df: pd.DataFrame) -> pd.DataFrame:
    result = df.copy()
    abas_standard = [
        "ABAS_comm_use",
        "ABAS_fun_acc",
        "ABAS_school_liv",
        "ABAS_Health_saf_sic",
        "ABAS_Leisure",
        "ABAS_Selfcare",
        "ABAS_Selfdirection",
        "ABAS_soc",
    ]
    abas_supplement = [
        "ABAS_comm_use_supp",
        "ABAS_fun_acc_suppongo",
        "ABAS_school_liv_suppongo",
        "ABAS_Health_saf_suppongo",
        "ABAS_Leisure_suppongo",
        "ABAS_Selfcare_suppongo",
        "ABAS_Selfdirection_suppongo",
        "ABAS_soc_suppongo",
    ]
    brief = [
        "BRIEF_Inhibit",
        "BRIEF_SelfMonitor",
        "BRIEF_shift",
        "BRIEF_Em_Con",
        "BRIEF_Initiate",
        "BRIEF_wm",
        "BRIEF_plann",
        "BRIEF_TaskMonitor",
        "BRIEF_OrganizMaterial",
    ]
    for columns, output in [
        (abas_standard, "ABAS_TOT"),
        ([*abas_standard, *abas_supplement], "ABAS_ASSUMED_TOTAL"),
        (brief, "BRIEF_TOT"),
    ]:
        existing = [column for column in columns if column in result]
        if existing:
            numeric = result[existing].apply(to_numeric)
            result[output] = numeric.sum(axis=1, min_count=1)
    return result


def add_minefield_composites(df: pd.DataFrame) -> pd.DataFrame:
    result = df.copy()
    for composite, components in MINEFIELD_COMPOSITES.items():
        for occasion in ("PRE", "POST"):
            accuracy_col = f"{components['accuracy']}_{occasion}"
            efficiency_col = f"{components['efficiency']}_{occasion}"
            if accuracy_col not in result or efficiency_col not in result:
                continue
            accuracy = to_numeric(result[accuracy_col]).where(lambda x: x.between(0, 16))
            efficiency = to_numeric(result[efficiency_col]).where(lambda x: x.between(0, 100))
            result[f"{composite}_{occasion}"] = pd.concat(
                [(accuracy / 16) * 100, efficiency], axis=1
            ).mean(axis=1, skipna=False)
    return result


def add_log_transforms(df: pd.DataFrame) -> pd.DataFrame:
    result = df.copy()
    for base in LOG_BASES:
        for occasion in ("PRE", "POST"):
            source = f"{base}_{occasion}"
            if source not in result:
                continue
            values = to_numeric(result[source])
            result[f"{base}_LN_{occasion}"] = np.log(values.where(values > 0))
    return result


def build_analysis_dataset() -> pd.DataFrame:
    df = load_source_database()
    df = add_questionnaire_totals(df)
    df = add_minefield_composites(df)
    df = add_log_transforms(df)
    return df


def normalize_column(value: str) -> str:
    return re.sub(r"\s+", "", str(value).upper())


def find_prepost_pairs(df: pd.DataFrame) -> list[dict]:
    lookup = {normalize_column(column): column for column in df.columns}
    pairs = []
    seen = set()
    for pre_col in df.columns:
        normalized = normalize_column(pre_col)
        if not normalized.endswith("_PRE"):
            continue
        base = normalized[:-4]
        post_col = lookup.get(f"{base}_POST")
        if post_col is None or base in seen:
            continue
        seen.add(base)
        pairs.append(
            {
                "variable": str(pre_col)[:-4],
                "pre_col": pre_col,
                "post_col": post_col,
                "domain": (
                    "minefield"
                    if str(pre_col).upper().startswith(("WM_", "PLANNING_", "PWM_"))
                    else "neuropsychological"
                ),
                "transformation": "log" if "_LN_" in str(pre_col) else "raw",
            }
        )
    return sorted(pairs, key=lambda row: row["variable"].upper())


def fdr_bh(frame: pd.DataFrame, p_column: str, q_column: str = "q_fdr_bh") -> pd.DataFrame:
    result = frame.copy()
    result[q_column] = np.nan
    if p_column not in result.columns:
        result[p_column] = np.nan
    p_values = pd.to_numeric(result[p_column], errors="coerce")
    valid = p_values.notna()
    if valid.any():
        result.loc[valid, q_column] = multipletests(
            p_values.loc[valid], method="fdr_bh"
        )[1]
    return result


def hedges_g(first: pd.Series, second: pd.Series) -> float:
    first = to_numeric(first).dropna()
    second = to_numeric(second).dropna()
    n1, n2 = len(first), len(second)
    if n1 < 2 or n2 < 2:
        return np.nan
    pooled_var = ((n1 - 1) * first.var(ddof=1) + (n2 - 1) * second.var(ddof=1)) / (n1 + n2 - 2)
    if pooled_var <= 0 or pd.isna(pooled_var):
        return np.nan
    d = (first.mean() - second.mean()) / math.sqrt(pooled_var)
    correction = 1 - 3 / (4 * (n1 + n2) - 9)
    return d * correction


def cohen_dz(change: pd.Series) -> float:
    values = to_numeric(change).dropna()
    if len(values) < 2 or values.std(ddof=1) == 0:
        return np.nan
    return values.mean() / values.std(ddof=1)


def numeric_summary(values: pd.Series, prefix: str = "") -> dict:
    x = to_numeric(values).dropna()
    key = f"{prefix}_" if prefix else ""
    return {
        f"{key}n": int(x.count()),
        f"{key}missing": int(len(values) - x.count()),
        f"{key}mean": x.mean(),
        f"{key}sd": x.std(ddof=1),
        f"{key}median": x.median(),
        f"{key}q1": x.quantile(0.25),
        f"{key}q3": x.quantile(0.75),
        f"{key}min": x.min(),
        f"{key}max": x.max(),
        f"{key}skew": x.skew(),
        f"{key}kurtosis": x.kurtosis(),
    }


def write_workbook(path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        wrote = False
        for sheet_name, frame in sheets.items():
            if frame is None:
                continue
            pd.DataFrame(frame).to_excel(writer, sheet_name=sheet_name[:31], index=False)
            wrote = True
        if not wrote:
            pd.DataFrame({"status": ["no results"]}).to_excel(writer, sheet_name="results", index=False)
    style_workbook(path)


def style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    significant_fill = PatternFill("solid", fgColor="C6EFCE")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        headers = [cell.value for cell in sheet[1]]
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        for index, header in enumerate(headers, 1):
            letter = sheet.cell(1, index).column_letter
            lengths = [len(str(header or ""))]
            lengths.extend(
                len(str(sheet.cell(row, index).value))
                for row in range(2, min(sheet.max_row, 300) + 1)
                if sheet.cell(row, index).value is not None
            )
            sheet.column_dimensions[letter].width = min(max(lengths, default=10) + 2, 45)
            header_lower = str(header or "").lower()
            if header_lower in {"p", "p_value", "q_fdr_bh", "welch_p", "mann_whitney_p", "interaction_p"} or header_lower.endswith("_p"):
                if sheet.max_row >= 2:
                    sheet.conditional_formatting.add(
                        f"{letter}2:{letter}{sheet.max_row}",
                        CellIsRule(operator="lessThan", formula=["0.05"], fill=significant_fill),
                    )
                for cell in sheet[letter][1:]:
                    cell.number_format = "0.00000"
    workbook.save(path)


def plot_prepost(df: pd.DataFrame, pair: dict, path: Path, interaction_p=np.nan) -> bool:
    columns = [ID_COLUMN, GROUP_COLUMN, pair["pre_col"], pair["post_col"]]
    data = df[columns].copy()
    data[pair["pre_col"]] = to_numeric(data[pair["pre_col"]])
    data[pair["post_col"]] = to_numeric(data[pair["post_col"]])
    data = data.dropna(subset=[pair["pre_col"], pair["post_col"], GROUP_COLUMN])
    if data.empty:
        return False
    rows = []
    for group_code, group_label in GROUP_LABELS.items():
        subset = data[data[GROUP_COLUMN] == group_code]
        for occasion, column in [("PRE", pair["pre_col"]), ("POST", pair["post_col"])]:
            values = subset[column].dropna()
            if values.empty:
                continue
            se = values.std(ddof=1) / math.sqrt(len(values)) if len(values) > 1 else np.nan
            rows.append({"group": group_label, "time": occasion, "mean": values.mean(), "ci": 1.96 * se})
    summary = pd.DataFrame(rows)
    if summary.empty:
        return False
    path.parent.mkdir(parents=True, exist_ok=True)
    colors = {"experimental": "#3566A5", "control": "#C45A11"}
    fig, ax = plt.subplots(figsize=(7.2, 4.8))
    for label in ["experimental", "control"]:
        group = summary[summary["group"] == label].copy()
        if group.empty:
            continue
        group["x"] = group["time"].map({"PRE": 0, "POST": 1})
        group = group.sort_values("x")
        ax.errorbar(group["x"], group["mean"], yerr=group["ci"], marker="o", linewidth=2, capsize=4, label=label, color=colors[label])
    ax.set_xticks([0, 1], ["PRE", "POST"])
    ax.set_title(pair["variable"])
    ax.set_ylabel("Score")
    ax.grid(axis="y", alpha=0.25)
    ax.legend(frameon=False)
    if pd.notna(interaction_p):
        label = "p < .001" if interaction_p < 0.001 else f"p = {interaction_p:.3f}"
        ax.text(0.98, 0.98, f"group x time {label}", transform=ax.transAxes, ha="right", va="top", fontsize=9)
    fig.tight_layout()
    fig.savefig(path, dpi=220)
    plt.close(fig)
    return True


def normality_safe_wilcoxon(pre: pd.Series, post: pd.Series) -> float:
    try:
        if len(pre) < 2 or np.allclose(pre, post):
            return np.nan
        return stats.wilcoxon(pre, post).pvalue
    except ValueError:
        return np.nan

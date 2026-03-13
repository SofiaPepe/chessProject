from scipy.stats import ttest_ind, f_oneway

import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.formatting.rule import CellIsRule
import os

# Percorso base della cartella descrittive
base_dir = os.path.dirname(os.path.abspath(__file__))


# Leggi il file Excel dalla cartella descrittive
db_path = os.path.join(base_dir, 'database_compositi.xlsx')
df = pd.read_excel(db_path)

# Raggruppa le classi: 1, 1A, 1B -> '1'; 2, 2A, 2B -> '2'; INF rimane 'INF'
if 'classe' in df.columns:
	df['classe'] = df['classe'].astype(str).str.strip().str.upper()
	df['classe'] = df['classe'].replace({
		'1A': '1', '1B': '1', '1': '1',
		'2A': '2', '2B': '2', '2': '2',
		'INF': 'INF', 'INFANZIA': 'INF'
	})

# Seleziona solo le colonne numeriche
df_num = df.select_dtypes(include='number')

# Stampa elenco variabili numeriche
print("Elenco variabili numeriche:")
for col in df_num.columns:
	print(col)
print("\nTotale variabili numeriche:", len(df_num.columns))

# Calcolo statistiche descrittive di base
statistiche = df_num.describe().T

# Calcolo curtosi e asimmetria
statistiche['skewness'] = df_num.skew()
statistiche['kurtosis'] = df_num.kurtosis()

# Stampa risultati
print("\nStatistiche descrittive con asimmetria e curtosi (solo colonne numeriche):\n")
print(statistiche)

# Statistiche per gruppo 1
statistiche_g1 = df[df['gruppo'] == 1][df_num.columns].describe().T
statistiche_g1['skewness'] = df[df['gruppo'] == 1][df_num.columns].skew()
statistiche_g1['kurtosis'] = df[df['gruppo'] == 1][df_num.columns].kurtosis()

# Statistiche per gruppo 2
statistiche_g2 = df[df['gruppo'] == 2][df_num.columns].describe().T
statistiche_g2['skewness'] = df[df['gruppo'] == 2][df_num.columns].skew()
statistiche_g2['kurtosis'] = df[df['gruppo'] == 2][df_num.columns].kurtosis()

# Tabella di contingenza gruppo x età
contingenza = pd.crosstab(df['gruppo'], df['età'])
print("\nTabella di contingenza (gruppo x età):")
print(contingenza)

# Tabella di contingenza classe x età
if 'classe' in df.columns:
	contingenza_classe_eta = pd.crosstab(df['classe'], df['età'])
	print("\nTabella di contingenza (classe x età):")
	print(contingenza_classe_eta)
else:
	contingenza_classe_eta = None
	print("\nColonna 'classe' non trovata per la tabella di contingenza classe x età.")

# Tabella di contingenza gruppo x età x sesso
if 'sesso' in df.columns:
	contingenza_3d = pd.crosstab([df['gruppo'], df['età']], df['sesso'])
	print("\nTabella di contingenza (gruppo x età x sesso):")
	print(contingenza_3d)
else:
	contingenza_3d = None
	print("\nColonna 'sesso' non trovata per la tabella di contingenza 3D.")

# Analisi differenze per sesso e per età
if 'sesso' in df.columns:
	risultati_sesso = []
	for col in df_num.columns:
		gruppi = df[['sesso', col]].dropna()
		if gruppi['sesso'].nunique() == 2:
			g1, g2 = gruppi['sesso'].unique()
			vals1 = gruppi[gruppi['sesso'] == g1][col]
			vals2 = gruppi[gruppi['sesso'] == g2][col]
			stat, p = ttest_ind(vals1, vals2, equal_var=False)
			risultati_sesso.append({'variabile': col, f'media_{g1}': vals1.mean(), f'media_{g2}': vals2.mean(), 'p_value': p})
	df_ris_sesso = pd.DataFrame(risultati_sesso)
	print("\nRisultati confronto per sesso (t-test):")
	print(df_ris_sesso)
else:
	df_ris_sesso = None
	print("\nColonna 'sesso' non trovata nel database.")

# Analisi per età (ANOVA)
risultati_eta = []
for col in df_num.columns:
	gruppi = [g[col].dropna() for _, g in df.groupby('età') if len(g[col].dropna()) > 0]
	if len(gruppi) > 1:
		stat, p = f_oneway(*gruppi)
		medie = df.groupby('età')[col].mean().to_dict()
		risultato = {'variabile': col, 'p_value': p}
		for eta, media in medie.items():
			risultato[f'media_eta_{eta}'] = media
		risultati_eta.append(risultato)
df_ris_eta = pd.DataFrame(risultati_eta)
print("\nRisultati confronto per età (ANOVA):")
print(df_ris_eta)

# Analisi per classe (ANOVA o t-test)
if 'classe' in df.columns:
	risultati_classe = []
	n_classi = df['classe'].nunique()
	for col in df_num.columns:
		gruppi = [g[col].dropna() for _, g in df.groupby('classe') if len(g[col].dropna()) > 0]
		if n_classi == 2 and len(gruppi) == 2:
			# t-test se solo 2 classi
			stat, p = ttest_ind(gruppi[0], gruppi[1], equal_var=False)
			medie = df.groupby('classe')[col].mean().to_dict()
			risultato = {'variabile': col, 'p_value': p}
			for classe, media in medie.items():
				risultato[f'media_classe_{classe}'] = media
			risultati_classe.append(risultato)
		elif len(gruppi) > 1:
			# ANOVA se più di 2 classi
			stat, p = f_oneway(*gruppi)
			medie = df.groupby('classe')[col].mean().to_dict()
			risultato = {'variabile': col, 'p_value': p}
			for classe, media in medie.items():
				risultato[f'media_classe_{classe}'] = media
			risultati_classe.append(risultato)
	df_ris_classe = pd.DataFrame(risultati_classe)
	print("\nRisultati confronto per classe:")
	print(df_ris_classe)
else:
	df_ris_classe = None
	print("\nColonna 'classe' non trovata per il confronto tra classi.")

# Salva tutto in un unico file Excel con più fogli nella cartella descrittive
excel_path = os.path.join(base_dir, 'statistiche_descrittive.xlsx')
with pd.ExcelWriter(excel_path) as writer:
	statistiche.to_excel(writer, sheet_name='Tutti')
	statistiche_g1.to_excel(writer, sheet_name='Gruppo 1')
	statistiche_g2.to_excel(writer, sheet_name='Gruppo 2')
	contingenza.to_excel(writer, sheet_name='Contingenza gruppo-età')
	if contingenza_3d is not None:
		contingenza_3d.to_excel(writer, sheet_name='Contingenza gruppo-eta-sesso')
	if contingenza_classe_eta is not None:
		contingenza_classe_eta.to_excel(writer, sheet_name='Contingenza classe-eta')
	if df_ris_sesso is not None:
		df_ris_sesso.to_excel(writer, sheet_name='Differenze Sesso', index=False)
	df_ris_eta.to_excel(writer, sheet_name='Differenze Età', index=False)
	if df_ris_classe is not None:
		df_ris_classe.to_excel(writer, sheet_name='Differenze Classe', index=False)

# Evidenzia i p-value significativi (<0.05) in verde nei fogli Differenze Sesso e Differenze Età
wb = openpyxl.load_workbook(excel_path)
fill = PatternFill(start_color='90EE90', end_color='90EE90', fill_type='solid')  # verde chiaro
for sheet_name in ['Differenze Sesso', 'Differenze Età']:
	if sheet_name in wb.sheetnames:
		ws = wb[sheet_name]
		# Trova la colonna del p-value
		for col in ws.iter_cols(1, ws.max_column):
			if col[0].value == 'p_value':
				col_letter = col[0].column_letter
				# Applica la regola di formattazione condizionale
				ws.conditional_formatting.add(f'{col_letter}2:{col_letter}{ws.max_row}',
					CellIsRule(operator='lessThan', formula=['0.05'], fill=fill))
				break

# Evidenzia i p-value significativi (<0.05) anche nel foglio Differenze Classe
for sheet_name in ['Differenze Sesso', 'Differenze Età', 'Differenze Classe']:
	if sheet_name in wb.sheetnames:
		ws = wb[sheet_name]
		# Trova la colonna del p-value
		for col in ws.iter_cols(1, ws.max_column):
			if col[0].value == 'p_value':
				col_letter = col[0].column_letter
				# Applica la regola di formattazione condizionale
				ws.conditional_formatting.add(f'{col_letter}2:{col_letter}{ws.max_row}',
					CellIsRule(operator='lessThan', formula=['0.05'], fill=fill))
				break
wb.save(excel_path)

print(f"\nTutte le statistiche (tutti, gruppo 1, gruppo 2), le tabelle di contingenza e i confronti per sesso/età sono stati salvati in '{excel_path}'.\nI p-value significativi sono evidenziati in verde.")

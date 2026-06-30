from pathlib import Path
import sys
import warnings

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from scipy import stats
import statsmodels.formula.api as smf


ROOT = Path(__file__).resolve().parents[2]
sys.path.append(str(ROOT / "analysis" / "descriptives"))

from add_composites import build_composite_dataframe, to_numeric  # noqa: E402


METADATA_PATH = ROOT / "analysis" / "shared" / "variable_metadata.csv"
COMPOSITE_METADATA_PATH = ROOT / "analysis" / "shared" / "composite_metadata.csv"
OUT_DIR = ROOT / "output" / "training_effect" / "treatment_effectiveness"
OUT_PATH = OUT_DIR / "treatment_effectiveness_results.xlsx"
ALPHA = 0.05

GROUP_LABELS = {
    1: "experimental",
    2: "control",
}


def bh_fdr(p_values: pd.Series) -> pd.Series:
    p = pd.to_numeric(p_values, errors="coerce")
    q = pd.Series(np.nan, index=p.index, dtype=float)
    valid = p.dropna().sort_values()
    m = len(valid)
    if m == 0:
        return q
    adjusted = valid * m / np.arange(1, m + 1)
    adjusted = adjusted.iloc[::-1].cummin().iloc[::-1].clip(upper=1)
    q.loc[adjusted.index] = adjusted
    return q


def hedges_g(group_a: pd.Series, group_b: pd.Series) -> float:
    group_a = group_a.dropna()
    group_b = group_b.dropna()
    n_a = group_a.count()
    n_b = group_b.count()
    if n_a < 2 or n_b < 2:
        return np.nan

    pooled_den = n_a + n_b - 2
    pooled_sd = np.sqrt(((n_a - 1) * group_a.var(ddof=1) + (n_b - 1) * group_b.var(ddof=1)) / pooled_den)
    if pooled_sd == 0 or pd.isna(pooled_sd):
        return np.nan

    cohen_d = (group_a.mean() - group_b.mean()) / pooled_sd
    correction = 1 - (3 / (4 * (n_a + n_b) - 9))
    return cohen_d * correction


def load_metadata(df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    metadata = pd.read_csv(METADATA_PATH)
    metadata = metadata[metadata["include_effectiveness"].astype(str).str.lower().eq("yes")].copy()
    metadata["min_possible"] = pd.to_numeric(metadata["min_possible"], errors="coerce")
    metadata["max_possible"] = pd.to_numeric(metadata["max_possible"], errors="coerce")
    metadata["valid_min"] = pd.to_numeric(metadata["valid_min"], errors="coerce")
    metadata["valid_max"] = pd.to_numeric(metadata["valid_max"], errors="coerce")

    validation_rows = []
    valid_rows = []
    for _, row in metadata.iterrows():
        missing = [col for col in [row["pre_col"], row["post_col"]] if col not in df.columns]
        has_required_bound = (
            pd.notna(row["max_possible"])
            if row["improvement_direction"] == "higher_better"
            else pd.notna(row["min_possible"])
        )
        status = "ok" if not missing and has_required_bound else "skipped"
        reason = ""
        if missing:
            reason = f"missing columns: {', '.join(missing)}"
        elif not has_required_bound:
            reason = "missing required theoretical bound"
        else:
            valid_rows.append(row)

        validation_rows.append(
            {
                "variable": row["variable"],
                "pre_col": row["pre_col"],
                "post_col": row["post_col"],
                "improvement_direction": row["improvement_direction"],
                "status": status,
                "reason": reason,
            }
        )

    return pd.DataFrame(valid_rows), pd.DataFrame(validation_rows)


def calculate_effectiveness(pre: pd.Series, post: pd.Series, metadata_row: pd.Series) -> tuple[pd.Series, pd.Series]:
    pre = to_numeric(pre)
    post = to_numeric(post)
    direction = metadata_row["improvement_direction"]
    min_possible = metadata_row["min_possible"]
    max_possible = metadata_row["max_possible"]
    valid_min = metadata_row.get("valid_min", np.nan)
    valid_max = metadata_row.get("valid_max", np.nan)

    status = pd.Series("ok", index=pre.index, dtype=object)
    effectiveness = pd.Series(np.nan, index=pre.index, dtype=float)
    complete = pre.notna() & post.notna()
    status.loc[~complete] = "missing_pre_or_post"

    valid_range = complete.copy()
    if pd.notna(valid_min):
        valid_range &= pre >= valid_min
        valid_range &= post >= valid_min
    if pd.notna(valid_max):
        valid_range &= pre <= valid_max
        valid_range &= post <= valid_max
    status.loc[complete & ~valid_range] = "invalid_pre_or_post_outside_valid_range"

    if direction == "higher_better":
        room = max_possible - pre
        no_room = valid_range & (room <= 0)
        valid = valid_range & (room > 0)
        effectiveness.loc[valid] = ((post.loc[valid] - pre.loc[valid]) / room.loc[valid]) * 100
        status.loc[no_room] = "ceiling_no_room_to_improve"
    elif direction == "lower_better":
        room = pre - min_possible
        no_room = valid_range & (room <= 0)
        valid = valid_range & (room > 0)
        effectiveness.loc[valid] = ((pre.loc[valid] - post.loc[valid]) / room.loc[valid]) * 100
        status.loc[no_room] = "floor_no_room_to_improve"
    else:
        status.loc[complete] = "invalid_direction"

    return effectiveness, status


def build_individual_effectiveness(df: pd.DataFrame, metadata: pd.DataFrame) -> pd.DataFrame:
    base_cols = ["ID", "group", "sex", "age", "class"]
    rows = []
    for _, meta in metadata.iterrows():
        effectiveness, status = calculate_effectiveness(df[meta["pre_col"]], df[meta["post_col"]], meta)
        temp = df[[col for col in base_cols if col in df.columns]].copy()
        temp["group_label"] = temp["group"].map(GROUP_LABELS)
        temp["variable"] = meta["variable"]
        temp["pre_col"] = meta["pre_col"]
        temp["post_col"] = meta["post_col"]
        temp["improvement_direction"] = meta["improvement_direction"]
        temp["min_possible"] = meta["min_possible"]
        temp["max_possible"] = meta["max_possible"]
        temp["valid_min"] = meta["valid_min"]
        temp["valid_max"] = meta["valid_max"]
        temp["pre_value"] = to_numeric(df[meta["pre_col"]])
        temp["post_value"] = to_numeric(df[meta["post_col"]])
        temp["raw_change"] = temp["post_value"] - temp["pre_value"]
        temp["effectiveness_pct"] = effectiveness
        temp["effectiveness_status"] = status
        rows.append(temp)
    return pd.concat(rows, ignore_index=True) if rows else pd.DataFrame()


def summarize_by_group(individual: pd.DataFrame) -> pd.DataFrame:
    rows = []
    ok = individual[individual["effectiveness_status"].eq("ok")].copy()
    for (variable, group_label), subset in ok.groupby(["variable", "group_label"]):
        values = subset["effectiveness_pct"].dropna()
        rows.append(
            {
                "variable": variable,
                "group": group_label,
                "n": int(values.count()),
                "mean": values.mean(),
                "sd": values.std(ddof=1),
                "median": values.median(),
                "q1": values.quantile(0.25),
                "q3": values.quantile(0.75),
                "min": values.min(),
                "max": values.max(),
            }
        )
    return pd.DataFrame(rows)


def compare_groups(individual: pd.DataFrame) -> pd.DataFrame:
    rows = []
    ok = individual[individual["effectiveness_status"].eq("ok")].copy()
    for variable, subset in ok.groupby("variable"):
        exp = subset.loc[subset["group"] == 1, "effectiveness_pct"].dropna()
        ctrl = subset.loc[subset["group"] == 2, "effectiveness_pct"].dropna()
        if exp.empty or ctrl.empty:
            rows.append({"variable": variable, "status": "skipped_missing_group"})
            continue

        t_p = np.nan
        t_stat = np.nan
        if exp.count() >= 2 and ctrl.count() >= 2:
            with warnings.catch_warnings():
                warnings.filterwarnings("ignore", category=RuntimeWarning)
                t_stat, t_p = stats.ttest_ind(exp, ctrl, equal_var=False, nan_policy="omit")

        try:
            mw_p = stats.mannwhitneyu(exp, ctrl, alternative="two-sided").pvalue
        except ValueError:
            mw_p = np.nan

        rows.append(
            {
                "variable": variable,
                "status": "ok",
                "experimental_n": int(exp.count()),
                "control_n": int(ctrl.count()),
                "experimental_mean": exp.mean(),
                "control_mean": ctrl.mean(),
                "mean_difference_exp_minus_control": exp.mean() - ctrl.mean(),
                "hedges_g": hedges_g(exp, ctrl),
                "welch_t": t_stat,
                "welch_p": t_p,
                "mann_whitney_p": mw_p,
            }
        )

    results = pd.DataFrame(rows)
    if not results.empty and "welch_p" in results.columns:
        results["welch_q_fdr_bh"] = bh_fdr(results["welch_p"])
        results["mann_whitney_q_fdr_bh"] = bh_fdr(results["mann_whitney_p"])
        results["significant_p05"] = (results["welch_p"] < ALPHA) | (results["mann_whitney_p"] < ALPHA)
        results["significant_fdr05"] = (results["welch_q_fdr_bh"] < ALPHA) | (
            results["mann_whitney_q_fdr_bh"] < ALPHA
        )
    return results


def age_adjusted_group_comparison(individual: pd.DataFrame) -> pd.DataFrame:
    rows = []
    ok = individual[individual["effectiveness_status"].eq("ok")].copy()
    if ok.empty:
        return pd.DataFrame()

    for variable, subset in ok.groupby("variable"):
        model_df = subset[["effectiveness_pct", "group", "age"]].copy()
        model_df["age"] = to_numeric(model_df["age"])
        model_df["group_experimental"] = (model_df["group"] == 1).astype(int)
        model_df = model_df.dropna()
        if len(model_df) < 20 or model_df["group_experimental"].nunique() < 2:
            rows.append({"variable": variable, "status": "skipped_insufficient_data", "n": int(len(model_df))})
            continue

        try:
            fit = smf.ols("effectiveness_pct ~ group_experimental + age", data=model_df).fit()
            conf = fit.conf_int()
            rows.append(
                {
                    "variable": variable,
                    "status": "ok",
                    "model": "effectiveness_pct ~ group_experimental + age",
                    "n": int(len(model_df)),
                    "experimental_n": int(model_df["group_experimental"].sum()),
                    "control_n": int((model_df["group_experimental"] == 0).sum()),
                    "group_effect_exp_minus_control_age_adjusted": fit.params.get("group_experimental", np.nan),
                    "group_effect_se": fit.bse.get("group_experimental", np.nan),
                    "group_effect_t": fit.tvalues.get("group_experimental", np.nan),
                    "group_effect_p": fit.pvalues.get("group_experimental", np.nan),
                    "group_effect_ci_low": conf.loc["group_experimental", 0]
                    if "group_experimental" in conf.index
                    else np.nan,
                    "group_effect_ci_high": conf.loc["group_experimental", 1]
                    if "group_experimental" in conf.index
                    else np.nan,
                    "age_effect": fit.params.get("age", np.nan),
                    "age_effect_se": fit.bse.get("age", np.nan),
                    "age_effect_t": fit.tvalues.get("age", np.nan),
                    "age_effect_p": fit.pvalues.get("age", np.nan),
                    "r_squared": fit.rsquared,
                }
            )
        except Exception as exc:
            rows.append({"variable": variable, "status": f"error: {type(exc).__name__}", "n": int(len(model_df))})

    results = pd.DataFrame(rows)
    if not results.empty and "group_effect_p" in results.columns:
        results["group_effect_q_fdr_bh"] = bh_fdr(results["group_effect_p"])
        results["significant_age_adjusted_p05"] = results["group_effect_p"] < ALPHA
        results["significant_age_adjusted_fdr05"] = results["group_effect_q_fdr_bh"] < ALPHA
    return results


def covariate_models(individual: pd.DataFrame) -> pd.DataFrame:
    rows = []
    ok = individual[individual["effectiveness_status"].eq("ok")].copy()
    if ok.empty:
        return pd.DataFrame()

    for variable, subset in ok.groupby("variable"):
        model_df = subset[["effectiveness_pct", "group", "age", "sex", "class"]].copy()
        model_df["age"] = to_numeric(model_df["age"])
        model_df["group_experimental"] = (model_df["group"] == 1).astype(int)
        model_df["sex"] = model_df["sex"].astype(str)
        model_df["class_cov"] = model_df["class"].astype(str)
        model_df = model_df.dropna()
        if len(model_df) < 20 or model_df["group"].nunique() < 2:
            rows.append({"variable": variable, "status": "skipped_insufficient_data", "n": int(len(model_df))})
            continue

        try:
            fit = smf.ols(
                "effectiveness_pct ~ group_experimental + age + C(sex) + C(class_cov)",
                data=model_df,
            ).fit()
            conf = fit.conf_int()
            for term in fit.params.index:
                rows.append(
                    {
                        "variable": variable,
                        "status": "ok",
                        "term": term,
                        "estimate": fit.params[term],
                        "se": fit.bse[term],
                        "t": fit.tvalues[term],
                        "p_value": fit.pvalues[term],
                        "ci_low": conf.loc[term, 0],
                        "ci_high": conf.loc[term, 1],
                        "n": int(len(model_df)),
                        "r_squared": fit.rsquared,
                    }
                )
        except Exception as exc:
            rows.append({"variable": variable, "status": f"error: {type(exc).__name__}", "n": int(len(model_df))})

    results = pd.DataFrame(rows)
    if not results.empty and "p_value" in results.columns:
        group_mask = results["term"].eq("group_experimental") & results["p_value"].notna()
        results["group_q_fdr_bh"] = np.nan
        if group_mask.any():
            results.loc[group_mask, "group_q_fdr_bh"] = bh_fdr(results.loc[group_mask, "p_value"])
    return results


def write_workbook(
    metadata: pd.DataFrame,
    validation: pd.DataFrame,
    individual: pd.DataFrame,
    summary: pd.DataFrame,
    comparisons: pd.DataFrame,
    age_adjusted: pd.DataFrame,
    covariates: pd.DataFrame,
) -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    readme = pd.DataFrame(
        [
            {"item": "input_database", "value": "data/FINAL_DATABASE.xlsx"},
            {"item": "metadata", "value": str(METADATA_PATH.relative_to(ROOT))},
            {"item": "composite_metadata", "value": str(COMPOSITE_METADATA_PATH.relative_to(ROOT))},
            {"item": "formula_higher_better", "value": "((post - pre) / (max_possible - pre)) * 100"},
            {"item": "formula_lower_better", "value": "((pre - post) / (pre - min_possible)) * 100"},
            {"item": "validity_rule", "value": "Rows outside metadata valid_min/valid_max are excluded before effectiveness"},
            {"item": "age_adjusted_model", "value": "effectiveness_pct ~ group_experimental + age"},
            {"item": "full_covariate_model", "value": "effectiveness_pct ~ group_experimental + age + C(sex) + C(class_cov)"},
            {"item": "scope", "value": "Only user-confirmed variables; response times and unconfirmed variables excluded"},
        ]
    )
    with pd.ExcelWriter(OUT_PATH, engine="openpyxl") as writer:
        readme.to_excel(writer, sheet_name="readme", index=False)
        metadata.to_excel(writer, sheet_name="variable_metadata", index=False)
        if COMPOSITE_METADATA_PATH.exists():
            pd.read_csv(COMPOSITE_METADATA_PATH).to_excel(writer, sheet_name="composite_metadata", index=False)
        validation.to_excel(writer, sheet_name="metadata_validation", index=False)
        individual.to_excel(writer, sheet_name="individual_effectiveness", index=False)
        summary.to_excel(writer, sheet_name="group_summary", index=False)
        comparisons.to_excel(writer, sheet_name="group_comparison", index=False)
        age_adjusted.to_excel(writer, sheet_name="age_adjusted_group", index=False)
        covariates.to_excel(writer, sheet_name="covariate_models", index=False)

    style_workbook(OUT_PATH)


def style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    green_fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        headers = [cell.value for cell in ws[1]]
        for col_idx, header in enumerate(headers, start=1):
            col_letter = ws.cell(row=1, column=col_idx).column_letter
            max_len = len(str(header or ""))
            for cell in ws[col_letter][1:]:
                if cell.value is not None:
                    max_len = max(max_len, min(len(str(cell.value)), 45))
            ws.column_dimensions[col_letter].width = min(max_len + 2, 48)
            header_text = str(header or "").lower()
            if header_text in {"welch_p", "mann_whitney_p", "p_value"} or header_text.endswith("_p"):
                ws.conditional_formatting.add(
                    f"{col_letter}2:{col_letter}{ws.max_row}",
                    CellIsRule(operator="lessThan", formula=["0.05"], fill=green_fill),
                )
            if header_text in {"significant_p05", "significant_fdr05"}:
                ws.conditional_formatting.add(
                    f"{col_letter}2:{col_letter}{ws.max_row}",
                    CellIsRule(operator="equal", formula=["TRUE"], fill=yellow_fill),
                )
    wb.save(path)


def main() -> None:
    df = build_composite_dataframe()
    df["group"] = pd.to_numeric(df["group"], errors="coerce")
    metadata, validation = load_metadata(df)
    individual = build_individual_effectiveness(df, metadata)
    summary = summarize_by_group(individual)
    comparisons = compare_groups(individual)
    age_adjusted = age_adjusted_group_comparison(individual)
    covariates = covariate_models(individual)
    write_workbook(metadata, validation, individual, summary, comparisons, age_adjusted, covariates)

    tested = int(metadata.shape[0])
    comparison_sig = int(comparisons["significant_p05"].sum()) if "significant_p05" in comparisons.columns else 0
    comparison_fdr = int(comparisons["significant_fdr05"].sum()) if "significant_fdr05" in comparisons.columns else 0
    age_sig = (
        int(age_adjusted["significant_age_adjusted_p05"].sum())
        if "significant_age_adjusted_p05" in age_adjusted.columns
        else 0
    )
    age_fdr = (
        int(age_adjusted["significant_age_adjusted_fdr05"].sum())
        if "significant_age_adjusted_fdr05" in age_adjusted.columns
        else 0
    )
    print("Treatment effectiveness analysis complete")
    print(f"Variables included: {tested}")
    print(f"Group comparisons p < .05: {comparison_sig}; FDR < .05: {comparison_fdr}")
    print(f"Age-adjusted group effects p < .05: {age_sig}; FDR < .05: {age_fdr}")
    print(f"Results workbook: {OUT_PATH}")


if __name__ == "__main__":
    main()

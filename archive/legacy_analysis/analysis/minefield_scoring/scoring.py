from __future__ import annotations

import csv
import hashlib
import re
import statistics
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.append(str(Path(__file__).resolve().parents[1] / "common"))

from shortest_path import TrialConfiguration, build_trial_configuration  # noqa: E402


ROOT = Path(__file__).resolve().parents[2]
RAW_DIR = ROOT / "data" / "Minefield_raw"
CONFIGURATION_FILE = ROOT / "data" / "configuration_trial.xlsx"
OFFICIAL_DATABASE = ROOT / "data" / "FINAL_DATABASE.xlsx"
COMPLETE_DATABASE = ROOT / "data" / "FINAL_DATABASE_RECALCULATED_MINEFIELD.xlsx"
FULL_BEST_AVAILABLE_DATABASE = ROOT / "data" / "FINAL_DATABASE_MINEFIELD_FULL_BEST_AVAILABLE.xlsx"
OUTPUT_DIR = ROOT / "output" / "minefield_scoring"
DEDICATED_OLDER_FILE = "_i due partecipanti anziani ORN E ANTO.xlsx"

EXCLUDED_PARTICIPANTS = {"VERDI", "PROVA BAMBINI", "LINVER"}
OLDER_PARTICIPANT_PREFIXES = ("ORNFLO", "ANTONO")
ANTO_CANONICAL_ID = "ANTONO17081939"
PARTICIPANT_ID_ALIASES = {
    "SOFIA": "SOFPEP23071998",
}
MAX_LEVEL = 9

DESCRIPTION_RE = re.compile(r"^(PLWM|PLN|WM)(\d+)([A-Z]*)$")
DESCRIPTION_TASK = {
    "WM": "wm",
    "PLN": "planning",
    "PLWM": "wmplanning",
}

SCORE_FIELDS = [
    "participant_id",
    "span_wm",
    "accuracy_wm",
    "span_planning",
    "accuracy_planning",
    "span_wmplanning",
    "accuracy_wmplanning",
    "route_efficiency_pct_planning",
    "route_efficiency_pct_wmplanning",
    "wm_mean_total_time",
    "wm_mean_planning_time",
    "wm_mean_execution_time",
    "planning_mean_total_time",
    "planning_mean_planning_time",
    "planning_mean_execution_time",
    "wmplanning_mean_total_time",
    "wmplanning_mean_planning_time",
    "wmplanning_mean_execution_time",
]

LONG_SCORE_FIELDS = ["participant_id", "occasion", *SCORE_FIELDS[1:]]

AUDIT_FIELDS = [
    "participant_id",
    "task",
    "test_type_raw",
    "id_test",
    "test_soggetto.time_stamp",
    "descrizione",
    "level",
    "attempt",
    "source_file",
    "source_row",
    "start",
    "stop",
    "bombs",
    "percorso_soggetto",
    "n_percorso_soggetto",
    "historical_percorso_minimo_raw",
    "historical_n_percorso_minimo",
    "bfs_minimum_path",
    "bfs_minimum_moves",
    "bfs_minimum_intermediate_cells",
    "efficency",
    "route_efficiency_pct",
    "total_time",
    "planning_time",
    "execution_time",
    "level_outcome",
    "accuracy_points",
    "included_in_span",
    "included_in_efficency",
    "included_in_timing",
    "timing_exclusion_reason",
    "exclusion_reason",
]

PROJECT_AUDIT_FIELDS = [AUDIT_FIELDS[0], "occasion", *AUDIT_FIELDS[1:]]

@dataclass
class RawTrial:
    source_file: str
    source_row: int
    participant_id: str
    surname_raw: str
    name_raw: str
    birth_date_raw: str
    task: str
    test_type_raw: str
    id_test: str
    test_timestamp_raw: str
    test_timestamp: datetime | None
    row_timestamp_raw: str
    row_timestamp: datetime | None
    description: str
    level: int
    attempt: str
    participant_path: str
    participant_cells: int | None
    historical_minimum_path: str
    historical_minimum_cells: int | None
    total_time: float | None
    planning_time: float | None
    execution_time: float | None
    configuration: TrialConfiguration


def clean(value) -> str:
    return str(value or "").strip()


def normalize_id(value) -> str:
    return re.sub(r"[^A-Z0-9]", "", clean(value).upper())


def to_float(value) -> float | None:
    text = clean(value).replace(",", ".")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def rel(path: Path) -> str:
    try:
        return str(path.relative_to(ROOT))
    except ValueError:
        return str(path)


def parse_timestamp(value) -> datetime | None:
    text = clean(value)
    if not text:
        return None
    if isinstance(value, datetime):
        return value
    for fmt in (
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y %H:%M",
        "%Y-%m-%d %H:%M:%S",
    ):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def id_test_sort_value(value: str) -> int:
    try:
        return int(clean(value))
    except ValueError:
        return -1


def count_path_cells(path_text) -> int | None:
    text = clean(path_text)
    if not text or text == "-":
        return None
    tokens = [token.strip() for token in text.split("-") if token.strip()]
    return len(tokens) if tokens else None


def parse_description(value) -> tuple[str, int, str] | None:
    text = clean(value).upper()
    match = DESCRIPTION_RE.fullmatch(text)
    if not match:
        return None
    prefix, level, attempt = match.groups()
    return DESCRIPTION_TASK[prefix], int(level), attempt


def read_xlsx_rows(path: Path) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    values = list(sheet.iter_rows(values_only=True))
    workbook.close()
    if not values:
        return []
    headers = [clean(value) for value in values[0]]
    rows = []
    for raw_row in values[1:]:
        row = {
            header: value
            for header, value in zip(headers, raw_row)
            if header
        }
        if any(value not in {None, ""} for value in row.values()):
            rows.append(row)
    return rows


def load_configurations() -> dict[str, TrialConfiguration]:
    configurations: dict[str, TrialConfiguration] = {}
    for row in read_xlsx_rows(CONFIGURATION_FILE):
        configuration = build_trial_configuration(
            row.get("descrizione"),
            row.get("tipo_test"),
            row.get("coordinate"),
            row.get("percorso_impostato_start"),
            row.get("percorso_impostato_stop"),
        )
        if not configuration.description:
            continue
        if configuration.description in configurations:
            raise ValueError(f"Duplicate configuration for {configuration.description}")
        configurations[configuration.description] = configuration
    return configurations


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def is_older_participant(participant_id: str) -> bool:
    normalized = normalize_id(participant_id)
    return normalized.startswith(OLDER_PARTICIPANT_PREFIXES)


def canonical_participant_id(row: dict, dedicated_file: bool) -> str:
    surname_code = normalize_id(row.get("cognome"))
    name_code = normalize_id(row.get("nome"))
    raw_participant_id = max(
        (surname_code, name_code),
        key=lambda value: (len(value), value == surname_code),
    )
    if dedicated_file and normalize_id(raw_participant_id).startswith("ANTONO"):
        return ANTO_CANONICAL_ID
    return PARTICIPANT_ID_ALIASES.get(raw_participant_id, raw_participant_id)


def source_record_key(row: dict, participant_id: str) -> tuple[str, ...]:
    return (
        normalize_id(participant_id),
        clean(row.get("id_perc")),
        clean(row.get("id_test")),
        clean(row.get("id_soggetto")),
        clean(row.get("descrizione")).upper(),
        clean(row.get("percorso_soggetto.time_stamp")),
        clean(row.get("test_soggetto.time_stamp")),
    )


def iter_raw_sources(counters: Counter):
    seen_hashes: dict[str, str] = {}
    for path in sorted(RAW_DIR.glob("*.csv")):
        digest = file_sha256(path)
        if digest in seen_hashes:
            counters["duplicate_files_skipped"] += 1
            counters[f"duplicate_file:{path.name}"] += 1
            continue
        seen_hashes[digest] = path.name
        counters["csv_files_read"] += 1
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            for row_number, row in enumerate(csv.DictReader(handle, delimiter=";"), start=2):
                yield path, row_number, row, False

    dedicated_path = RAW_DIR / DEDICATED_OLDER_FILE
    if dedicated_path.exists():
        counters["dedicated_xlsx_files_read"] += 1
        for row_number, row in enumerate(read_xlsx_rows(dedicated_path), start=2):
            yield dedicated_path, row_number, row, True


def load_raw_trials(
    configurations: dict[str, TrialConfiguration],
) -> tuple[list[RawTrial], Counter]:
    trials: list[RawTrial] = []
    counters: Counter = Counter()
    seen_records: set[tuple[str, ...]] = set()

    for path, row_number, row, dedicated_file in iter_raw_sources(counters):
        counters["raw_rows_seen"] += 1
        participant_id = canonical_participant_id(row, dedicated_file)

        if not dedicated_file and is_older_participant(participant_id):
            counters["global_rows_excluded_orn_anto"] += 1
            continue
        if participant_id in EXCLUDED_PARTICIPANTS:
            counters["rows_excluded_demo_or_nonstandard_id"] += 1
            continue

        record_key = source_record_key(row, participant_id)
        if record_key in seen_records:
            counters["duplicate_records_skipped"] += 1
            continue
        seen_records.add(record_key)

        description = clean(row.get("descrizione")).upper()
        parsed = parse_description(description)
        if parsed is None:
            counters["rows_excluded_unparseable_description"] += 1
            continue
        task, level, attempt = parsed

        configuration = configurations.get(description)
        if configuration is None:
            counters["rows_excluded_missing_configuration"] += 1
            continue

        participant_path = clean(row.get("percorso_soggetto"))
        historical_path = clean(row.get("percorso_minimo"))
        total_time = to_float(row.get("tempo_tot"))
        planning_time = to_float(row.get("tempo_pianifica"))
        execution_time = (
            total_time - planning_time
            if total_time is not None and planning_time is not None
            else None
        )
        trials.append(
            RawTrial(
                source_file=path.name,
                source_row=row_number,
                participant_id=participant_id,
                surname_raw=clean(row.get("cognome")),
                name_raw=clean(row.get("nome")),
                birth_date_raw=clean(row.get("data_nascita")),
                task=task,
                test_type_raw=clean(row.get("test_soggetto.tipo_test")),
                id_test=clean(row.get("id_test")),
                test_timestamp_raw=clean(row.get("test_soggetto.time_stamp")),
                test_timestamp=parse_timestamp(row.get("test_soggetto.time_stamp")),
                row_timestamp_raw=clean(row.get("percorso_soggetto.time_stamp")),
                row_timestamp=parse_timestamp(row.get("percorso_soggetto.time_stamp")),
                description=description,
                level=level,
                attempt=attempt,
                participant_path=participant_path,
                participant_cells=count_path_cells(participant_path),
                historical_minimum_path=historical_path,
                historical_minimum_cells=count_path_cells(historical_path),
                total_time=total_time,
                planning_time=planning_time,
                execution_time=execution_time,
                configuration=configuration,
            )
        )
        counters["candidate_task_rows"] += 1

    counters["unique_source_records"] = len(seen_records)
    return trials, counters


def session_sort_key(rows: list[RawTrial]) -> tuple[datetime, int]:
    timestamps = [row.test_timestamp for row in rows if row.test_timestamp is not None]
    return (
        max(timestamps) if timestamps else datetime.min,
        id_test_sort_value(rows[0].id_test),
    )


def select_latest_sessions(trials: list[RawTrial], counters: Counter) -> list[RawTrial]:
    sessions: dict[tuple[str, str, str], list[RawTrial]] = defaultdict(list)
    for trial in trials:
        sessions[(trial.participant_id, trial.task, trial.id_test)].append(trial)

    participant_tasks: dict[tuple[str, str], list[list[RawTrial]]] = defaultdict(list)
    for (participant_id, task, _), rows in sessions.items():
        participant_tasks[(participant_id, task)].append(rows)

    selected = []
    for session_groups in participant_tasks.values():
        if len(session_groups) > 1:
            counters["participant_task_multiple_sessions_resolved"] += 1
        chosen = max(session_groups, key=session_sort_key)
        selected.extend(chosen)
        counters["selected_sessions"] += 1
    counters["nonselected_session_rows"] = len(trials) - len(selected)
    return selected


def row_recency_key(trial: RawTrial) -> tuple[datetime, int]:
    return (
        trial.row_timestamp or trial.test_timestamp or datetime.min,
        trial.source_row,
    )


def blank_audit_row(trial: RawTrial) -> dict:
    configuration = trial.configuration
    minimum_path = "-".join(configuration.minimum_path)
    if minimum_path:
        minimum_path += "-"
    efficiency = ""
    route_efficiency_pct = ""
    if (
        trial.participant_cells is not None
        and configuration.minimum_intermediate_cells is not None
    ):
        efficiency = trial.participant_cells - configuration.minimum_intermediate_cells
        if efficiency >= 0 and trial.participant_cells > 0:
            route_efficiency_pct = round(
                100
                * configuration.minimum_intermediate_cells
                / trial.participant_cells,
                6,
            )
        elif efficiency >= 0 and configuration.minimum_intermediate_cells == 0:
            route_efficiency_pct = 100.0
    return {
        "participant_id": trial.participant_id,
        "task": trial.task,
        "test_type_raw": trial.test_type_raw,
        "id_test": trial.id_test,
        "test_soggetto.time_stamp": trial.test_timestamp_raw,
        "descrizione": trial.description,
        "level": trial.level,
        "attempt": trial.attempt or "base",
        "source_file": trial.source_file,
        "source_row": trial.source_row,
        "start": configuration.start_raw,
        "stop": configuration.stop_raw,
        "bombs": configuration.bombs_raw,
        "percorso_soggetto": trial.participant_path,
        "n_percorso_soggetto": trial.participant_cells if trial.participant_cells is not None else "",
        "historical_percorso_minimo_raw": trial.historical_minimum_path,
        "historical_n_percorso_minimo": (
            trial.historical_minimum_cells
            if trial.historical_minimum_cells is not None
            else ""
        ),
        "bfs_minimum_path": minimum_path,
        "bfs_minimum_moves": (
            configuration.minimum_moves
            if configuration.minimum_moves is not None
            else ""
        ),
        "bfs_minimum_intermediate_cells": (
            configuration.minimum_intermediate_cells
            if configuration.minimum_intermediate_cells is not None
            else ""
        ),
        "efficency": efficiency,
        "route_efficiency_pct": route_efficiency_pct,
        "total_time": trial.total_time if trial.total_time is not None else "",
        "planning_time": (
            trial.planning_time if trial.planning_time is not None else ""
        ),
        "execution_time": (
            trial.execution_time if trial.execution_time is not None else ""
        ),
        "level_outcome": "",
        "accuracy_points": "",
        "included_in_span": "no",
        "included_in_efficency": "no",
        "included_in_timing": "no",
        "timing_exclusion_reason": "",
        "exclusion_reason": "",
    }


def apply_efficiency_status(row: dict) -> None:
    if row["task"] == "wm":
        row["exclusion_reason"] = "task_has_no_spatial_efficiency"
        return
    if row["n_percorso_soggetto"] == "":
        row["exclusion_reason"] = "missing_participant_path"
        return
    if row["bfs_minimum_intermediate_cells"] == "":
        row["exclusion_reason"] = "missing_bfs_path"
        return
    if row["efficency"] != "" and float(row["efficency"]) < 0:
        row["exclusion_reason"] = "negative_efficency"
        return
    row["included_in_efficency"] = "yes"
    row["exclusion_reason"] = ""


def apply_timing_status(row: dict) -> None:
    if row["total_time"] == "" or row["planning_time"] == "":
        row["timing_exclusion_reason"] = "missing_time"
        return
    if row["execution_time"] == "":
        row["timing_exclusion_reason"] = "missing_execution_time"
        return
    if float(row["execution_time"]) < 0:
        row["timing_exclusion_reason"] = "negative_execution_time"
        return
    row["included_in_timing"] = "yes"


def score_selected_trials(
    selected_trials: list[RawTrial],
    counters: Counter,
) -> list[dict]:
    sessions: dict[tuple[str, str, str], list[RawTrial]] = defaultdict(list)
    for trial in selected_trials:
        sessions[(trial.participant_id, trial.task, trial.id_test)].append(trial)

    audit_rows = []
    for session_rows in sessions.values():
        by_description: dict[str, list[RawTrial]] = defaultdict(list)
        for trial in session_rows:
            by_description[trial.description].append(trial)

        chosen_by_description: dict[str, RawTrial] = {}
        for description, rows in by_description.items():
            chosen_by_description[description] = max(rows, key=row_recency_key)

        by_level: dict[int, dict[str, RawTrial]] = defaultdict(dict)
        for trial in chosen_by_description.values():
            attempt_key = "retry" if trial.attempt else "base"
            current = by_level[trial.level].get(attempt_key)
            if current is None or row_recency_key(trial) > row_recency_key(current):
                by_level[trial.level][attempt_key] = trial

        level_results: dict[int, tuple[str, int | str, RawTrial | None]] = {}
        levels = set(by_level)
        for level in sorted(levels):
            attempts = by_level[level]
            has_next_level = (level + 1) in levels
            base = attempts.get("base")
            retry = attempts.get("retry")
            if retry is not None:
                if has_next_level:
                    level_results[level] = ("retry_success", 1, retry)
                else:
                    level_results[level] = ("failed_second_attempt", 0, None)
            elif base is not None and (has_next_level or level == MAX_LEVEL):
                level_results[level] = ("first_attempt_success", 2, base)
            else:
                level_results[level] = ("terminal_base_unconfirmed", "", None)

        for description, rows in by_description.items():
            chosen = chosen_by_description[description]
            for trial in rows:
                output = blank_audit_row(trial)
                if trial is not chosen:
                    output["level_outcome"] = "duplicate_attempt_superseded"
                    output["exclusion_reason"] = "duplicate_attempt_superseded"
                    counters["trial_rows_duplicate_attempt_superseded"] += 1
                    audit_rows.append(output)
                    continue

                outcome, points, successful_trial = level_results[trial.level]
                output["level_outcome"] = outcome
                if successful_trial is trial:
                    output["accuracy_points"] = points
                    output["included_in_span"] = "yes"
                    apply_efficiency_status(output)
                    apply_timing_status(output)
                elif outcome == "retry_success":
                    output["exclusion_reason"] = "superseded_by_successful_retry"
                    output["timing_exclusion_reason"] = "superseded_by_successful_retry"
                elif outcome == "failed_second_attempt":
                    if trial.attempt:
                        output["accuracy_points"] = 0
                    output["exclusion_reason"] = "failed_second_attempt"
                    output["timing_exclusion_reason"] = "failed_second_attempt"
                elif outcome == "terminal_base_unconfirmed":
                    output["exclusion_reason"] = "terminal_base_unconfirmed"
                    output["timing_exclusion_reason"] = "terminal_base_unconfirmed"
                elif successful_trial is not None:
                    output["exclusion_reason"] = "non_scoring_attempt"
                    output["timing_exclusion_reason"] = "non_scoring_attempt"
                audit_rows.append(output)

    audit_rows.sort(
        key=lambda row: (
            row["participant_id"],
            row["task"],
            id_test_sort_value(row["id_test"]),
            int(row["level"]),
            row["attempt"],
            int(row["source_row"]),
        )
    )
    for row in audit_rows:
        counters[f"level_outcome:{row['level_outcome']}"] += 1
        if row["included_in_efficency"] == "yes":
            counters["trial_rows_included_in_efficency"] += 1
        if row["included_in_timing"] == "yes":
            counters["trial_rows_included_in_timing"] += 1
        if row["timing_exclusion_reason"] == "negative_execution_time":
            counters["trial_rows_negative_execution_time"] += 1
    counters["audit_rows_written"] = len(audit_rows)
    return audit_rows


def mean_or_blank(values: list[float]) -> float | str:
    return round(statistics.mean(values), 6) if values else ""


def build_participant_outputs(
    audit_rows: list[dict],
) -> list[dict]:
    participants = sorted({row["participant_id"] for row in audit_rows})
    grouped: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for row in audit_rows:
        grouped[(row["participant_id"], row["task"])].append(row)

    csv_rows = []
    for participant_id in participants:
        csv_row = {"participant_id": participant_id}
        for task in ("wm", "planning", "wmplanning"):
            rows = grouped.get((participant_id, task), [])
            if rows:
                successful = [row for row in rows if row["included_in_span"] == "yes"]
                span = max((int(row["level"]) for row in successful), default=0)
                points_by_level: dict[int, int] = {}
                for row in rows:
                    if row["accuracy_points"] == "":
                        continue
                    level = int(row["level"])
                    points_by_level[level] = max(
                        points_by_level.get(level, 0),
                        int(row["accuracy_points"]),
                    )
                accuracy = sum(points_by_level.values())
                csv_row[f"span_{task}"] = span
                csv_row[f"accuracy_{task}"] = accuracy
            else:
                csv_row[f"span_{task}"] = ""
                csv_row[f"accuracy_{task}"] = ""

        for task in ("planning", "wmplanning"):
            rows = grouped.get((participant_id, task), [])
            included_values = [
                float(row["route_efficiency_pct"])
                for row in rows
                if row["included_in_efficency"] == "yes"
                and row["route_efficiency_pct"] != ""
            ]
            route_efficiency_pct = mean_or_blank(included_values)
            csv_row[f"route_efficiency_pct_{task}"] = route_efficiency_pct

        for task in ("wm", "planning", "wmplanning"):
            rows = grouped.get((participant_id, task), [])
            timing_rows = [
                row for row in rows if row["included_in_timing"] == "yes"
            ]
            csv_row[f"{task}_mean_total_time"] = mean_or_blank(
                [float(row["total_time"]) for row in timing_rows]
            )
            csv_row[f"{task}_mean_planning_time"] = mean_or_blank(
                [float(row["planning_time"]) for row in timing_rows]
            )
            csv_row[f"{task}_mean_execution_time"] = mean_or_blank(
                [float(row["execution_time"]) for row in timing_rows]
            )
        csv_rows.append(csv_row)
    return csv_rows


def write_csv(path: Path, rows: list[dict], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_table_sheet(
    workbook: Workbook,
    title: str,
    rows: list[dict],
    fieldnames: list[str],
    highlight_negative: bool = False,
) -> None:
    if workbook.active.title == "Sheet" and workbook.active.max_row == 1:
        sheet = workbook.active
        sheet.title = title
    else:
        sheet = workbook.create_sheet(title)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    negative_fill = PatternFill("solid", fgColor="FFC7CE")
    negative_font = Font(color="9C0006")

    for column, field in enumerate(fieldnames, start=1):
        cell = sheet.cell(row=1, column=column, value=field)
        cell.fill = header_fill
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for row_number, row in enumerate(rows, start=2):
        negative = (
            highlight_negative
            and (
                (
                    row.get("efficency") not in {"", None}
                    and float(row["efficency"]) < 0
                )
                or (
                    row.get("execution_time") not in {"", None}
                    and float(row["execution_time"]) < 0
                )
            )
        )
        for column, field in enumerate(fieldnames, start=1):
            cell = sheet.cell(row=row_number, column=column, value=row.get(field, ""))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if negative:
                cell.fill = negative_fill
                cell.font = negative_font

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, field in enumerate(fieldnames, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = min(
            max(len(field) + 4, 14),
            48,
        )


def configuration_rows(
    configurations: dict[str, TrialConfiguration],
) -> list[dict]:
    rows = []
    for configuration in configurations.values():
        minimum_path = "-".join(configuration.minimum_path)
        if minimum_path:
            minimum_path += "-"
        rows.append(
            {
                "description": configuration.description,
                "test_type": configuration.test_type,
                "bombs": configuration.bombs_raw,
                "start": configuration.start_raw,
                "stop": configuration.stop_raw,
                "minimum_moves": (
                    configuration.minimum_moves
                    if configuration.minimum_moves is not None
                    else ""
                ),
                "minimum_intermediate_cells": (
                    configuration.minimum_intermediate_cells
                    if configuration.minimum_intermediate_cells is not None
                    else ""
                ),
                "minimum_path": minimum_path,
            }
        )
    return sorted(rows, key=lambda row: row["description"])


def source_file_rows() -> list[dict]:
    paths = [CONFIGURATION_FILE, *sorted(RAW_DIR.glob("*"))]
    return [
        {
            "source_file": rel(path),
            "sha256": file_sha256(path),
            "size_bytes": path.stat().st_size,
        }
        for path in paths
        if path.is_file()
    ]


def write_scoring_workbook(
    path: Path,
    participant_rows: list[dict],
    audit_rows: list[dict],
    configurations: dict[str, TrialConfiguration],
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    write_table_sheet(
        workbook,
        "participant_scores",
        participant_rows,
        SCORE_FIELDS,
    )
    write_table_sheet(
        workbook,
        "trial_audit",
        audit_rows,
        AUDIT_FIELDS,
        highlight_negative=True,
    )
    config_rows = configuration_rows(configurations)
    write_table_sheet(
        workbook,
        "configurations",
        config_rows,
        [
            "description",
            "test_type",
            "bombs",
            "start",
            "stop",
            "minimum_moves",
            "minimum_intermediate_cells",
            "minimum_path",
        ],
    )
    exclusion_rows = [
        row
        for row in audit_rows
        if row["included_in_span"] != "yes"
        or row["included_in_timing"] != "yes"
        or (
            row["task"] != "wm"
            and row["included_in_efficency"] != "yes"
        )
    ]
    write_table_sheet(
        workbook,
        "exclusions",
        exclusion_rows,
        AUDIT_FIELDS,
        highlight_negative=True,
    )
    sources = source_file_rows()
    write_table_sheet(
        workbook,
        "source_files",
        sources,
        ["source_file", "sha256", "size_bytes"],
    )
    workbook.save(path)


def write_summary(
    counters: Counter,
    configurations: dict[str, TrialConfiguration],
    audit_rows: list[dict],
    participant_rows: list[dict],
) -> None:
    outcomes = Counter(row["level_outcome"] for row in audit_rows)
    exclusions = Counter(
        row["exclusion_reason"]
        for row in audit_rows
        if row["exclusion_reason"]
    )
    older_sources = Counter(
        row["source_file"]
        for row in audit_rows
        if is_older_participant(row["participant_id"])
    )
    lines = [
        "# Minefield raw score recalculation",
        "",
        "## Method",
        "",
        "- Minimum routes are recalculated from `configuration_trial.xlsx` with breadth-first search on an 8x8 grid.",
        "- Movement is orthogonal only; bomb coordinates are blocked.",
        "- Start and stop are excluded from the minimum intermediate-cell count.",
        "- Historical raw `percorso_minimo` values are retained only in the audit output and are not used for scoring.",
        "- Audit `efficency = participant intermediate cells - BFS minimum intermediate cells`.",
        "- Participant route efficiency is the mean of `(BFS minimum intermediate cells / participant intermediate cells) x 100` across eligible successful trials.",
        "- Timing means use only successful trials that contribute to span.",
        "- `execution_time = tempo_tot - tempo_pianifica`; negative execution times are excluded.",
        "- ORN and ANTO are read only from the dedicated older-participant workbook.",
        "",
        "## Counts",
        "",
        f"- Configurations read: {len(configurations)}",
        f"- CSV files read after file-level deduplication: {counters['csv_files_read']}",
        f"- Duplicate files skipped: {counters['duplicate_files_skipped']}",
        f"- Raw rows seen: {counters['raw_rows_seen']}",
        f"- Duplicate records skipped: {counters['duplicate_records_skipped']}",
        f"- ORN/ANTO rows excluded from global CSV files: {counters['global_rows_excluded_orn_anto']}",
        f"- Selected sessions: {counters['selected_sessions']}",
        f"- Audit rows written: {len(audit_rows)}",
        f"- Participant rows written: {len(participant_rows)}",
        f"- Successful trial rows included in timing means: {counters['trial_rows_included_in_timing']}",
        f"- Negative execution-time rows excluded: {counters['trial_rows_negative_execution_time']}",
        "",
        "## ORN/ANTO source verification",
        "",
    ]
    for source, count in sorted(older_sources.items()):
        lines.append(f"- `{source}`: {count} selected trial rows")
    lines.extend(["", "## Level outcomes", ""])
    for outcome, count in sorted(outcomes.items()):
        lines.append(f"- `{outcome}`: {count}")
    lines.extend(["", "## Exclusion reasons", ""])
    for reason, count in sorted(exclusions.items()):
        lines.append(f"- `{reason}`: {count}")
    lines.extend(
        [
            "",
            "## Output files",
            "",
            "- `minefield_scores.csv`",
            "- `minefield_scoring.xlsx`",
        ]
    )
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    (OUTPUT_DIR / "summary.md").write_text(
        "\n".join(lines) + "\n",
        encoding="utf-8",
    )


def validate_outputs(
    configurations: dict[str, TrialConfiguration],
    audit_rows: list[dict],
    participant_rows: list[dict],
) -> None:
    route_configs = [
        configuration
        for configuration in configurations.values()
        if configuration.description.startswith(("PLN", "PLWM"))
    ]
    unreachable = [
        configuration.description
        for configuration in route_configs
        if configuration.minimum_intermediate_cells is None
    ]
    if unreachable:
        raise ValueError(f"Unreachable configured trials: {', '.join(unreachable)}")

    missing_config_rows = [
        row["descrizione"]
        for row in audit_rows
        if row["task"] in {"planning", "wmplanning"}
        and row["bfs_minimum_intermediate_cells"] == ""
    ]
    if missing_config_rows:
        raise ValueError(f"Rows without BFS minimum: {missing_config_rows[:5]}")

    wrong_older_sources = [
        row
        for row in audit_rows
        if is_older_participant(row["participant_id"])
        and row["source_file"] != DEDICATED_OLDER_FILE
    ]
    if wrong_older_sources:
        raise ValueError("ORN/ANTO selected rows were found outside the dedicated workbook")

    for row in participant_rows:
        if set(row) != set(SCORE_FIELDS):
            raise ValueError(
                f"Unexpected participant score columns for {row['participant_id']}"
            )
        for task in ("wm", "planning", "wmplanning"):
            span = row[f"span_{task}"]
            accuracy = row[f"accuracy_{task}"]
            if span != "" and not 0 <= int(span) <= MAX_LEVEL:
                raise ValueError(f"Invalid span for {row['participant_id']}: {span}")
            if accuracy != "" and not 0 <= int(accuracy) <= 16:
                raise ValueError(
                    f"Invalid accuracy for {row['participant_id']}: {accuracy}"
                )
            total = row[f"{task}_mean_total_time"]
            planning = row[f"{task}_mean_planning_time"]
            execution = row[f"{task}_mean_execution_time"]
            if all(value != "" for value in (total, planning, execution)):
                if abs(float(total) - float(planning) - float(execution)) > 1e-5:
                    raise ValueError(
                        f"Timing identity failed for {row['participant_id']} / {task}"
                    )
                if float(execution) < 0:
                    raise ValueError(
                        f"Negative mean execution time for {row['participant_id']} / {task}"
                    )
        for task in ("planning", "wmplanning"):
            value = row[f"route_efficiency_pct_{task}"]
            if value != "" and not 0 <= float(value) <= 100:
                raise ValueError(
                    f"Invalid route efficiency percentage for {row['participant_id']} / {task}"
                )


PROJECT_ID_RE = re.compile(r"^(EXP|GC|SPR|PPR|AC)0*(\d+)$")
PROJECT_TASKS = ("wm", "planning", "wmplanning")
PROJECT_OCCASIONS = ("PRE", "POST")
SINGLE_OCCASION_ASSIGNMENTS = {
    "EXP02": "PRE",
    "EXP04": "PRE",
    "EXP06": "PRE",
    "EXP07": "PRE",
    "EXP08": "PRE",
    "EXP10": "PRE",
    "EXP12": "PRE",
    "EXP15": "PRE",
    "EXP16": "PRE",
    "SPR3": "POST",
}
EXCLUDED_PROJECT_ID_BIRTH_DATES = {
    (
        "SPR2",
        "29/03/2017",
    ): "ID SPR2 reused for a different child; historical scores and DOB identify 27/06/2017 as the project participant",
}
SESSION_SELECTION_FIELDS = [
    "participant_id",
    "task",
    "test_date",
    "occasion",
    "id_test",
    "test_timestamp",
    "source_files",
    "rows",
    "unique_descriptions",
    "selected",
    "selection_reason",
]
COVERAGE_FIELDS = [
    "participant_id",
    "status",
    "pre_tasks",
    "post_tasks",
    "pre_date",
    "post_date",
    "source_files",
]
ID_ISSUE_FIELDS = [
    "participant_id",
    "birth_date",
    "action",
    "reason",
    "rows",
    "source_files",
]


@dataclass
class ProjectSession:
    participant_id: str
    task: str
    test_date: str
    id_test: str
    rows: list[RawTrial]
    occasion: str = ""
    selected: bool = False


def canonical_project_id(value) -> str:
    match = PROJECT_ID_RE.fullmatch(normalize_id(value))
    if match is None:
        return ""
    prefix, number_text = match.groups()
    number = int(number_text)
    if prefix == "SPR":
        return f"{prefix}{number}"
    return f"{prefix}{number:02d}"


def birth_date_key(value) -> str:
    return clean(value).split(" ", 1)[0]


def load_official_database_records() -> tuple[list[str], list[dict]]:
    workbook = load_workbook(OFFICIAL_DATABASE, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    values = list(sheet.iter_rows(values_only=True))
    workbook.close()
    if not values:
        raise ValueError(f"Empty official database: {rel(OFFICIAL_DATABASE)}")

    headers = [str(value) if value is not None else "" for value in values[0]]
    if len(headers) != len(set(headers)):
        raise ValueError("The official database contains duplicate column names")
    if "ID" not in headers:
        raise ValueError("The official database has no ID column")

    records = []
    seen_ids = set()
    for raw_row in values[1:]:
        record = dict(zip(headers, raw_row))
        participant_id = canonical_project_id(record.get("ID"))
        if not participant_id:
            continue
        if participant_id in seen_ids:
            raise ValueError(f"Duplicate official participant ID: {participant_id}")
        record["ID"] = participant_id
        records.append(record)
        seen_ids.add(participant_id)
    return headers, records


def project_id_from_trial(trial: RawTrial, official_ids: set[str]) -> str:
    candidates = {
        canonical_project_id(value)
        for value in (trial.surname_raw, trial.name_raw, trial.participant_id)
    }
    candidates.discard("")
    official_candidates = candidates & official_ids
    if len(official_candidates) > 1:
        raise ValueError(
            f"Ambiguous project ID in {trial.source_file}:{trial.source_row}: "
            f"{sorted(official_candidates)}"
        )
    if official_candidates:
        return next(iter(official_candidates))
    return ""


def prepare_project_trials(
    trials: list[RawTrial],
    official_ids: set[str],
    counters: Counter,
) -> tuple[list[RawTrial], list[dict]]:
    accepted = []
    issue_groups: dict[tuple[str, str, str], dict] = {}
    unmatched_ids = Counter()

    for trial in trials:
        participant_id = project_id_from_trial(trial, official_ids)
        if not participant_id:
            raw_id = canonical_project_id(trial.participant_id) or normalize_id(
                trial.participant_id
            )
            unmatched_ids[raw_id or "BLANK"] += 1
            counters["project_rows_excluded_nonofficial_id"] += 1
            continue

        birth_date = birth_date_key(trial.birth_date_raw)
        exclusion_reason = EXCLUDED_PROJECT_ID_BIRTH_DATES.get(
            (participant_id, birth_date)
        )
        if exclusion_reason:
            key = (participant_id, birth_date, exclusion_reason)
            group = issue_groups.setdefault(
                key,
                {
                    "participant_id": participant_id,
                    "birth_date": birth_date,
                    "action": "excluded",
                    "reason": exclusion_reason,
                    "rows": 0,
                    "source_files": set(),
                },
            )
            group["rows"] += 1
            group["source_files"].add(trial.source_file)
            counters["project_rows_excluded_reused_id"] += 1
            continue

        trial.participant_id = participant_id
        accepted.append(trial)
        counters["project_candidate_rows"] += 1

    for raw_id, count in unmatched_ids.items():
        counters[f"nonofficial_id:{raw_id}"] = count

    issue_rows = []
    for group in issue_groups.values():
        output = dict(group)
        output["source_files"] = ", ".join(sorted(group["source_files"]))
        issue_rows.append(output)
    return accepted, sorted(
        issue_rows,
        key=lambda row: (row["participant_id"], row["birth_date"]),
    )


def session_test_date(rows: list[RawTrial]) -> str:
    timestamps = [row.test_timestamp for row in rows if row.test_timestamp is not None]
    if not timestamps:
        raise ValueError(
            f"Session without timestamp: {rows[0].participant_id} / {rows[0].task}"
        )
    dates = {timestamp.date().isoformat() for timestamp in timestamps}
    if len(dates) != 1:
        raise ValueError(
            f"Session spans multiple dates: {rows[0].participant_id} / {rows[0].task}"
        )
    return next(iter(dates))


def session_completeness_key(session: ProjectSession) -> tuple:
    descriptions = {row.description for row in session.rows}
    levels = {row.level for row in session.rows}
    return (
        len(descriptions),
        len(levels),
        max(levels, default=-1),
        session_sort_key(session.rows),
    )


def select_project_sessions(
    trials: list[RawTrial],
    counters: Counter,
) -> tuple[list[ProjectSession], list[dict]]:
    grouped: dict[tuple[str, str, str], list[RawTrial]] = defaultdict(list)
    for trial in trials:
        grouped[(trial.participant_id, trial.task, trial.id_test)].append(trial)

    sessions = [
        ProjectSession(
            participant_id=participant_id,
            task=task,
            id_test=id_test,
            test_date=session_test_date(rows),
            rows=rows,
        )
        for (participant_id, task, id_test), rows in grouped.items()
    ]

    by_day: dict[tuple[str, str, str], list[ProjectSession]] = defaultdict(list)
    for session in sessions:
        by_day[(session.participant_id, session.task, session.test_date)].append(
            session
        )

    selected = []
    for same_day_sessions in by_day.values():
        chosen = max(same_day_sessions, key=session_completeness_key)
        chosen.selected = True
        selected.append(chosen)
        if len(same_day_sessions) > 1:
            counters["same_day_restarts_resolved"] += len(same_day_sessions) - 1

    dates_by_participant: dict[str, set[str]] = defaultdict(set)
    for session in selected:
        dates_by_participant[session.participant_id].add(session.test_date)

    occasion_by_participant_date: dict[tuple[str, str], str] = {}
    for participant_id, date_values in dates_by_participant.items():
        dates = sorted(date_values)
        if len(dates) > 2:
            raise ValueError(
                f"More than two assessment dates for {participant_id}: {dates}"
            )
        if len(dates) == 2:
            occasion_by_participant_date[(participant_id, dates[0])] = "PRE"
            occasion_by_participant_date[(participant_id, dates[1])] = "POST"
        else:
            occasion = SINGLE_OCCASION_ASSIGNMENTS.get(participant_id)
            if occasion is None:
                raise ValueError(
                    f"Single assessment date needs explicit PRE/POST assignment: "
                    f"{participant_id} / {dates[0]}"
                )
            occasion_by_participant_date[(participant_id, dates[0])] = occasion
            counters[f"single_occasion_assigned_{occasion.lower()}"] += 1

    for session in sessions:
        session.occasion = occasion_by_participant_date[
            (session.participant_id, session.test_date)
        ]

    selection_rows = []
    for session in sorted(
        sessions,
        key=lambda item: (
            item.participant_id,
            item.test_date,
            item.task,
            id_test_sort_value(item.id_test),
        ),
    ):
        selection_rows.append(
            {
                "participant_id": session.participant_id,
                "task": session.task,
                "test_date": session.test_date,
                "occasion": session.occasion,
                "id_test": session.id_test,
                "test_timestamp": max(
                    (
                        row.test_timestamp_raw
                        for row in session.rows
                        if row.test_timestamp is not None
                    ),
                    default="",
                ),
                "source_files": ", ".join(
                    sorted({row.source_file for row in session.rows})
                ),
                "rows": len(session.rows),
                "unique_descriptions": len(
                    {row.description for row in session.rows}
                ),
                "selected": "yes" if session.selected else "no",
                "selection_reason": (
                    "most_complete_then_latest_on_date"
                    if session.selected
                    else "same_day_restart_superseded"
                ),
            }
        )
    counters["project_sessions_selected"] = len(selected)
    return selected, selection_rows


def score_project_sessions(
    selected_sessions: list[ProjectSession],
    official_order: list[str],
    counters: Counter,
) -> tuple[list[dict], list[dict]]:
    grouped: dict[tuple[str, str], list[RawTrial]] = defaultdict(list)
    for session in selected_sessions:
        grouped[(session.participant_id, session.occasion)].extend(session.rows)

    calculated_rows: dict[tuple[str, str], dict] = {}
    all_audit_rows = []
    for (participant_id, occasion), rows in sorted(grouped.items()):
        audit_rows = score_selected_trials(rows, counters)
        participant_rows = build_participant_outputs(audit_rows)
        if len(participant_rows) != 1:
            raise ValueError(
                f"Expected one score row for {participant_id} / {occasion}"
            )
        score_row = participant_rows[0]
        score_row["occasion"] = occasion
        calculated_rows[(participant_id, occasion)] = score_row
        for audit_row in audit_rows:
            audit_row["occasion"] = occasion
        all_audit_rows.extend(audit_rows)

    long_rows = []
    for participant_id in official_order:
        for occasion in PROJECT_OCCASIONS:
            row = {
                "participant_id": participant_id,
                "occasion": occasion,
                **{field: "" for field in SCORE_FIELDS[1:]},
            }
            calculated = calculated_rows.get((participant_id, occasion))
            if calculated:
                for field in SCORE_FIELDS[1:]:
                    row[field] = calculated.get(field, "")
            long_rows.append(row)

    validation_rows = [
        {field: row.get(field, "") for field in SCORE_FIELDS}
        for row in long_rows
    ]
    validate_outputs(
        load_configurations(),
        all_audit_rows,
        validation_rows,
    )
    all_audit_rows.sort(
        key=lambda row: (
            row["participant_id"],
            row["occasion"],
            row["task"],
            id_test_sort_value(row["id_test"]),
            int(row["level"]),
            row["attempt"],
        )
    )
    return long_rows, all_audit_rows


def build_coverage_rows(
    official_order: list[str],
    selected_sessions: list[ProjectSession],
) -> list[dict]:
    tasks: dict[tuple[str, str], set[str]] = defaultdict(set)
    dates: dict[tuple[str, str], set[str]] = defaultdict(set)
    sources: dict[str, set[str]] = defaultdict(set)
    for session in selected_sessions:
        key = (session.participant_id, session.occasion)
        tasks[key].add(session.task)
        dates[key].add(session.test_date)
        sources[session.participant_id].update(
            row.source_file for row in session.rows
        )

    coverage_rows = []
    expected_tasks = set(PROJECT_TASKS)
    for participant_id in official_order:
        pre_tasks = tasks[(participant_id, "PRE")]
        post_tasks = tasks[(participant_id, "POST")]
        if not pre_tasks and not post_tasks:
            status = "missing_raw"
        elif pre_tasks == expected_tasks and post_tasks == expected_tasks:
            status = "complete_pre_post"
        elif pre_tasks == expected_tasks and not post_tasks:
            status = "pre_only"
        elif post_tasks == expected_tasks and not pre_tasks:
            status = "post_only"
        else:
            status = "partial_pre_post"
        coverage_rows.append(
            {
                "participant_id": participant_id,
                "status": status,
                "pre_tasks": ", ".join(sorted(pre_tasks)),
                "post_tasks": ", ".join(sorted(post_tasks)),
                "pre_date": ", ".join(sorted(dates[(participant_id, "PRE")])),
                "post_date": ", ".join(
                    sorted(dates[(participant_id, "POST")])
                ),
                "source_files": ", ".join(sorted(sources[participant_id])),
            }
        )
    return coverage_rows


def is_legacy_minefield_column(header: str) -> bool:
    normalized = clean(header).upper()
    return normalized.startswith(("WM_", "PLANNING_", "PWM_"))


def database_score_field(score_field: str, occasion: str) -> str:
    return f"MF_{score_field.upper()}_{occasion}"


def database_best_available_field(score_field: str, occasion: str) -> str:
    return f"MF_BEST_AVAILABLE_{score_field.upper()}_{occasion}"


HISTORICAL_FALLBACK_COLUMNS = {
    "span_wm": {"PRE": "WM_SPAN_PRE", "POST": "WM_SPAN_POST"},
    "accuracy_wm": {"PRE": "WM_ACC_PRE", "POST": "WM_ACC_POST"},
    "span_planning": {"PRE": "PLANNING_SPAN_PRE", "POST": "PLANNING_SPAN_POST"},
    "accuracy_planning": {"PRE": "PLANNING_ACC_PRE", "POST": "PLANNING_ACC_POST"},
    "route_efficiency_pct_planning": {"PRE": "PLANNING_PERC_PRE", "POST": "PLANNING_PERC_POST"},
    "planning_mean_total_time": {"PRE": "PLANNING_TR_PRE", "POST": "PLANNING_TOT_POST"},
    "planning_mean_planning_time": {"PRE": "PLANNING_TP_PRE", "POST": "PLANNING_TP_POST"},
    "planning_mean_execution_time": {"PRE": "PLANNING_TE_PRE", "POST": "PLANNING_TE_POST"},
    "accuracy_wmplanning": {"PRE": "PWM_ACC_PRE", "POST": "PWM_ACC_POST"},
    "route_efficiency_pct_wmplanning": {"PRE": "PWM_PERC_PRE", "POST": "PWM_PERC_POST"},
    "wmplanning_mean_total_time": {"PRE": "PWM_TR_PRE", "POST": "PWM_T_POST"},
    "wmplanning_mean_planning_time": {"PRE": "PWM_TP_PRE", "POST": "PWM_TP_POST"},
    "wmplanning_mean_execution_time": {"PRE": "PWM_TE_PRE", "POST": "PWM_TE_POST"},
}

PROVENANCE_FIELDS = [
    "participant_id",
    "occasion",
    "score_field",
    "database_column",
    "source",
    "historical_source_column",
    "value",
]


def normalized_header(value: str) -> str:
    return re.sub(r"\s+", "", clean(value).upper())


def complete_database_rows(
    original_headers: list[str],
    original_records: list[dict],
    long_rows: list[dict],
) -> tuple[list[str], list[dict], list[dict], list[dict]]:
    score_headers = [
        database_score_field(field, occasion)
        for occasion in PROJECT_OCCASIONS
        for field in SCORE_FIELDS[1:]
    ]
    output_headers = [*original_headers, *score_headers]
    original_header_lookup = {
        normalized_header(header): header for header in original_headers
    }
    original_by_id = {record["ID"]: record for record in original_records}

    complete_long_rows = []
    provenance_rows = []
    for raw_row in long_rows:
        row = dict(raw_row)
        participant_id = row["participant_id"]
        occasion = row["occasion"]
        historical_record = original_by_id[participant_id]
        for field in SCORE_FIELDS[1:]:
            value = row.get(field, "")
            source = "recalculated_raw" if value not in {None, ""} else "missing"
            historical_source_column = ""
            fallback_name = HISTORICAL_FALLBACK_COLUMNS.get(field, {}).get(occasion)
            if fallback_name:
                actual_header = original_header_lookup.get(normalized_header(fallback_name))
                if actual_header:
                    historical_source_column = actual_header
                    historical_value = historical_record.get(actual_header)
                    if value in {None, ""} and historical_value not in {None, ""}:
                        value = historical_value
                        row[field] = value
                        source = "historical_fallback"
            provenance_rows.append(
                {
                    "participant_id": participant_id,
                    "occasion": occasion,
                    "score_field": field,
                    "database_column": database_best_available_field(field, occasion),
                    "source": source,
                    "historical_source_column": historical_source_column,
                    "value": value,
                }
            )
        complete_long_rows.append(row)

    complete_score_index = {
        (row["participant_id"], row["occasion"]): row
        for row in complete_long_rows
    }
    raw_score_index = {
        (row["participant_id"], row["occasion"]): row
        for row in long_rows
    }
    best_available_headers = [
        database_best_available_field(field, occasion)
        for occasion in PROJECT_OCCASIONS
        for field in SCORE_FIELDS[1:]
    ]
    output_headers = [*original_headers, *score_headers, *best_available_headers]

    output_rows = []
    for record in original_records:
        participant_id = record["ID"]
        output = {header: record.get(header) for header in original_headers}
        for occasion in PROJECT_OCCASIONS:
            raw_score_row = raw_score_index[(participant_id, occasion)]
            complete_score_row = complete_score_index[(participant_id, occasion)]
            for field in SCORE_FIELDS[1:]:
                output[database_score_field(field, occasion)] = raw_score_row[field]
                output[database_best_available_field(field, occasion)] = complete_score_row[field]
        output_rows.append(output)
    return output_headers, output_rows, complete_long_rows, provenance_rows


def method_rows(
    original_headers: list[str],
    provenance_rows: list[dict],
) -> list[dict]:
    source_counts = Counter(row["source"] for row in provenance_rows)
    return [
        {"item": "source_database", "value": rel(OFFICIAL_DATABASE)},
        {"item": "raw_directory", "value": rel(RAW_DIR)},
        {"item": "configuration_file", "value": rel(CONFIGURATION_FILE)},
        {"item": "participant_key", "value": "ID"},
        {
            "item": "pre_post_assignment",
            "value": "earliest assessment date = PRE; latest assessment date = POST",
        },
        {
            "item": "same_day_restarts",
            "value": "most complete session retained; latest timestamp breaks ties",
        },
        {
            "item": "minimum_path",
            "value": "BFS on 8x8 grid with orthogonal movement and bombs blocked",
        },
        {
            "item": "historical_minefield_columns",
            "value": "all 48 historical Minefield columns retained unchanged alongside 34 recalculated PRE/POST fields",
        },
        {
            "item": "fallback_policy",
            "value": "MF_ columns remain recalculated-only; MF_BEST_AVAILABLE_ columns use the historically equivalent value when recalculation is blank",
        },
        {
            "item": "recalculated_raw_values",
            "value": source_counts["recalculated_raw"],
        },
        {
            "item": "historical_fallback_values",
            "value": source_counts["historical_fallback"],
        },
        {
            "item": "values_still_missing",
            "value": source_counts["missing"],
        },
        {
            "item": "formula_policy",
            "value": "cached values retained; source workbook remains unchanged",
        },
    ]


def write_project_scoring_workbook(
    path: Path,
    long_rows: list[dict],
    audit_rows: list[dict],
    configurations: dict[str, TrialConfiguration],
    selection_rows: list[dict],
    coverage_rows: list[dict],
    issue_rows: list[dict],
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    write_table_sheet(
        workbook,
        "minefield_scores_long",
        long_rows,
        LONG_SCORE_FIELDS,
    )
    write_table_sheet(
        workbook,
        "trial_audit",
        audit_rows,
        PROJECT_AUDIT_FIELDS,
        highlight_negative=True,
    )
    write_table_sheet(
        workbook,
        "session_selection",
        selection_rows,
        SESSION_SELECTION_FIELDS,
    )
    write_table_sheet(
        workbook,
        "id_coverage",
        coverage_rows,
        COVERAGE_FIELDS,
    )
    write_table_sheet(
        workbook,
        "id_issues",
        issue_rows,
        ID_ISSUE_FIELDS,
    )
    config_rows = configuration_rows(configurations)
    write_table_sheet(
        workbook,
        "configurations",
        config_rows,
        [
            "description",
            "test_type",
            "bombs",
            "start",
            "stop",
            "minimum_moves",
            "minimum_intermediate_cells",
            "minimum_path",
        ],
    )
    write_table_sheet(
        workbook,
        "source_files",
        source_file_rows(),
        ["source_file", "sha256", "size_bytes"],
    )
    workbook.save(path)


def write_complete_database(
    path: Path,
    output_headers: list[str],
    output_rows: list[dict],
    raw_long_rows: list[dict],
    complete_long_rows: list[dict],
    coverage_rows: list[dict],
    methods: list[dict],
    provenance_rows: list[dict],
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    write_table_sheet(
        workbook,
        "complete_database",
        output_rows,
        output_headers,
    )
    write_table_sheet(
        workbook,
        "minefield_scores_long",
        raw_long_rows,
        LONG_SCORE_FIELDS,
    )
    write_table_sheet(
        workbook,
        "minefield_best_available",
        complete_long_rows,
        LONG_SCORE_FIELDS,
    )
    write_table_sheet(
        workbook,
        "minefield_provenance",
        provenance_rows,
        PROVENANCE_FIELDS,
    )
    write_table_sheet(
        workbook,
        "id_coverage",
        coverage_rows,
        COVERAGE_FIELDS,
    )
    write_table_sheet(
        workbook,
        "method",
        methods,
        ["item", "value"],
    )
    workbook.save(path)


def validate_complete_database(
    path: Path,
    original_records: list[dict],
    original_headers: list[str],
    output_headers: list[str],
    provenance_rows: list[dict],
) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["complete_database"]
    values = list(sheet.iter_rows(values_only=True))
    workbook.close()
    headers = [str(value) if value is not None else "" for value in values[0]]
    if headers != output_headers:
        raise ValueError("Complete database headers do not match the expected contract")

    output_records = [dict(zip(headers, row)) for row in values[1:]]
    if len(output_records) != len(original_records):
        raise ValueError("Complete database participant count changed")
    output_by_id = {row["ID"]: row for row in output_records}
    if len(output_by_id) != len(output_records):
        raise ValueError("Complete database has duplicate IDs")

    for original in original_records:
        output = output_by_id.get(original["ID"])
        if output is None:
            raise ValueError(f"Participant lost from complete database: {original['ID']}")
        for header in original_headers:
            if output.get(header) != original.get(header):
                raise ValueError(
                    f"Historical database value changed: {original['ID']} / {header}"
                )

    score_headers = [
        header
        for header in headers
        if header.startswith("MF_")
        and not header.startswith("MF_BEST_AVAILABLE_")
    ]
    if len(score_headers) != 34:
        raise ValueError(f"Expected 34 Minefield PRE/POST fields, found {len(score_headers)}")
    best_available_headers = [
        header for header in headers if header.startswith("MF_BEST_AVAILABLE_")
    ]
    if len(best_available_headers) != 34:
        raise ValueError(
            f"Expected 34 best-available Minefield fields, found {len(best_available_headers)}"
        )

    if len(provenance_rows) != len(original_records) * 2 * len(SCORE_FIELDS[1:]):
        raise ValueError("Minefield provenance row count is incomplete")
    provenance_index = {
        (row["participant_id"], row["database_column"]): row
        for row in provenance_rows
    }
    for participant_id, output in output_by_id.items():
        for header in best_available_headers:
            provenance = provenance_index[(participant_id, header)]
            output_value = output.get(header)
            provenance_value = provenance.get("value")
            if output_value in {None, ""} and provenance_value in {None, ""}:
                continue
            if output_value != provenance_value:
                raise ValueError(
                    f"Minefield provenance mismatch: {participant_id} / {header}"
                )


def write_full_best_available_summary(
    counters: Counter,
    configurations: dict[str, TrialConfiguration],
    long_rows: list[dict],
    audit_rows: list[dict],
    coverage_rows: list[dict],
    original_headers: list[str],
    provenance_rows: list[dict],
) -> None:
    coverage_counts = Counter(row["status"] for row in coverage_rows)
    missing_ids = [
        row["participant_id"]
        for row in coverage_rows
        if row["status"] == "missing_raw"
    ]
    negative_efficiency_rows = sum(
        row["exclusion_reason"] == "negative_efficency"
        for row in audit_rows
    )
    source_counts = Counter(row["source"] for row in provenance_rows)
    source_counts_by_field = Counter(
        (row["score_field"], row["source"]) for row in provenance_rows
    )
    lines = [
        "# Chess project full best-available Minefield database",
        "",
        "## Method",
        "",
        "- Trial minimum paths were recalculated from `configuration_trial.xlsx` with BFS on an 8x8 grid.",
        "- Only orthogonal movement is allowed and bomb cells are blocked.",
        "- Historical raw `percorso_minimo` values were retained only for audit.",
        "- Repeated sessions on the same date were resolved by completeness, then recency.",
        "- The earliest assessment date is PRE and the latest is POST.",
        "- The source `FINAL_DATABASE.xlsx` was not modified.",
        "- All historical columns are retained unchanged; `MF_` columns remain recalculated-only and `MF_BEST_AVAILABLE_` columns add historical fallback values.",
        "",
        "## Counts",
        "",
        f"- Trial configurations: {len(configurations)}",
        f"- Raw CSV rows read: {counters['raw_rows_seen']}",
        f"- Official-project trial rows retained: {counters['project_candidate_rows']}",
        f"- Rows excluded because ID is not in the official database: {counters['project_rows_excluded_nonofficial_id']}",
        f"- Rows excluded for reused SPR2 ID: {counters['project_rows_excluded_reused_id']}",
        f"- Same-day restarted sessions superseded: {counters['same_day_restarts_resolved']}",
        f"- Selected participant/task sessions: {counters['project_sessions_selected']}",
        f"- Trial audit rows: {len(audit_rows)}",
        f"- Successful trial rows included in timing means: {counters['trial_rows_included_in_timing']}",
        f"- Negative execution-time rows excluded: {counters['trial_rows_negative_execution_time']}",
        f"- Negative path-efficiency rows excluded: {negative_efficiency_rows}",
        f"- Long score rows: {len(long_rows)} (101 participants x PRE/POST)",
        f"- Original database columns: {len(original_headers)}",
        f"- Original database columns retained: {len(original_headers)}",
        "- Recalculated Minefield columns added: 34",
        "- Best-available Minefield columns added: 34",
        f"- Recalculated raw values: {source_counts['recalculated_raw']}",
        f"- Historical fallback values: {source_counts['historical_fallback']}",
        f"- Values still missing: {source_counts['missing']}",
        "",
        "## Source Counts by Measure",
        "",
        "| Measure | Recalculated raw | Historical fallback | Missing |",
        "|---|---:|---:|---:|",
    ]
    for field in SCORE_FIELDS[1:]:
        lines.append(
            "| "
            f"`{field}` | "
            f"{source_counts_by_field[(field, 'recalculated_raw')]} | "
            f"{source_counts_by_field[(field, 'historical_fallback')]} | "
            f"{source_counts_by_field[(field, 'missing')]} |"
        )
    lines.extend(
        [
            "",
            "## ID coverage",
            "",
        ]
    )
    for status, count in sorted(coverage_counts.items()):
        lines.append(f"- `{status}`: {count}")
    lines.extend(
        [
            f"- Missing raw IDs: {', '.join(missing_ids) if missing_ids else 'none'}",
            "",
            "## Outputs",
            "",
            f"- `{rel(FULL_BEST_AVAILABLE_DATABASE)}`",
            f"- `{rel(OUTPUT_DIR / 'minefield_scores_long.csv')}`",
            f"- `{rel(OUTPUT_DIR / 'minefield_scoring.xlsx')}`",
        ]
    )
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    (OUTPUT_DIR / "full_best_available_summary.md").write_text(
        "\n".join(lines) + "\n",
        encoding="utf-8",
    )


EFFICIENCY_DATABASE_COLUMNS = {
    ("PRE", "route_efficiency_pct_planning"): "PLANNING_PERC_PRE",
    ("POST", "route_efficiency_pct_planning"): "PLANNING_PERC_POST",
    ("PRE", "route_efficiency_pct_wmplanning"): "PWM_PERC_PRE",
    ("POST", "route_efficiency_pct_wmplanning"): "PWM_PERC_POST",
}
EFFICIENCY_UPDATE_FIELDS = [
    "participant_id",
    "occasion",
    "score_field",
    "database_column",
    "source",
    "original_value",
    "output_value",
]


def build_efficiency_only_database(
    original_headers: list[str],
    original_records: list[dict],
    long_rows: list[dict],
) -> tuple[list[dict], list[dict]]:
    header_lookup = {
        normalized_header(header): header for header in original_headers
    }
    score_index = {
        (row["participant_id"], row["occasion"]): row
        for row in long_rows
    }
    output_rows = []
    update_rows = []
    for original in original_records:
        participant_id = original["ID"]
        output = dict(original)
        for (occasion, score_field), requested_header in EFFICIENCY_DATABASE_COLUMNS.items():
            actual_header = header_lookup.get(normalized_header(requested_header))
            if actual_header is None:
                raise KeyError(f"Historical route-efficiency column not found: {requested_header}")
            original_value = original.get(actual_header)
            recalculated_value = score_index[(participant_id, occasion)].get(score_field, "")
            if recalculated_value in {None, ""}:
                output_value = original_value
                source = "historical_original"
            else:
                output_value = recalculated_value
                source = "recalculated_raw"
            output[actual_header] = output_value
            update_rows.append(
                {
                    "participant_id": participant_id,
                    "occasion": occasion,
                    "score_field": score_field,
                    "database_column": actual_header,
                    "source": source,
                    "original_value": original_value,
                    "output_value": output_value,
                }
            )
        output_rows.append(output)
    return output_rows, update_rows


def write_efficiency_only_database(
    path: Path,
    headers: list[str],
    rows: list[dict],
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    write_table_sheet(workbook, "Foglio1", rows, headers)
    workbook.save(path)


def validate_efficiency_only_database(
    original_headers: list[str],
    original_records: list[dict],
    update_rows: list[dict],
) -> None:
    workbook = load_workbook(COMPLETE_DATABASE, read_only=True, data_only=True)
    if workbook.sheetnames != ["Foglio1"]:
        raise ValueError(f"Unexpected recalculated database sheets: {workbook.sheetnames}")
    sheet = workbook["Foglio1"]
    values = list(sheet.iter_rows(values_only=True))
    workbook.close()
    headers = [str(value) if value is not None else "" for value in values[0]]
    if headers != original_headers:
        raise ValueError("Recalculated database columns differ from the original database")
    output_records = [dict(zip(headers, row)) for row in values[1:]]
    output_by_id = {row["ID"]: row for row in output_records}
    original_by_id = {row["ID"]: row for row in original_records}
    if len(output_records) != len(original_records) or len(output_by_id) != len(output_records):
        raise ValueError("Participant rows changed in the recalculated database")

    target_columns = {row["database_column"] for row in update_rows}
    for participant_id, original in original_by_id.items():
        output = output_by_id[participant_id]
        for header in original_headers:
            if header in target_columns:
                continue
            if output.get(header) != original.get(header):
                raise ValueError(
                    f"Non-efficiency value changed: {participant_id} / {header}"
                )

    update_index = {
        (row["participant_id"], row["database_column"]): row
        for row in update_rows
    }
    for participant_id, output in output_by_id.items():
        for column in target_columns:
            expected = update_index[(participant_id, column)]["output_value"]
            actual = output.get(column)
            if actual != expected:
                raise ValueError(
                    f"Route-efficiency update mismatch: {participant_id} / {column}"
                )


def write_efficiency_only_summary(
    counters: Counter,
    configurations: dict[str, TrialConfiguration],
    audit_rows: list[dict],
    update_rows: list[dict],
) -> None:
    sources = Counter(row["source"] for row in update_rows)
    lines = [
        "# Chess project route-efficiency recalculation",
        "",
        "- The output database retains the original 104-column schema and 101 participant rows.",
        "- Only `PLANNING_PERC_PRE`, `PLANNING_PERC_POST`, `PWM_PERC_PRE`, and `PWM_PERC_POST` may change.",
        "- Recalculated route efficiency is used when raw data are available; otherwise the original value is retained.",
        "- `PLANNING_EFF_*` and `PWM_EFF_*` remain unchanged because they are tile-difference measures.",
        "",
        "## Counts",
        "",
        f"- Configurations: {len(configurations)}",
        f"- Trial audit rows: {len(audit_rows)}",
        f"- Efficiency cells considered: {len(update_rows)}",
        f"- Recalculated cells used: {sources['recalculated_raw']}",
        f"- Original cells retained: {sources['historical_original']}",
        "",
        "## Outputs",
        "",
        f"- `{rel(COMPLETE_DATABASE)}`",
        f"- `{rel(OUTPUT_DIR / 'database_route_efficiency_updates.csv')}`",
        f"- `{rel(OUTPUT_DIR / 'minefield_scoring.xlsx')}`",
    ]
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    (OUTPUT_DIR / "summary.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    configurations = load_configurations()
    raw_trials, counters = load_raw_trials(configurations)
    original_headers, original_records = load_official_database_records()
    official_order = [record["ID"] for record in original_records]
    official_ids = set(official_order)
    project_trials, issue_rows = prepare_project_trials(
        raw_trials,
        official_ids,
        counters,
    )
    selected_sessions, selection_rows = select_project_sessions(
        project_trials,
        counters,
    )
    long_rows, audit_rows = score_project_sessions(
        selected_sessions,
        official_order,
        counters,
    )
    coverage_rows = build_coverage_rows(official_order, selected_sessions)
    (
        full_database_headers,
        full_database_rows,
        complete_long_rows,
        provenance_rows,
    ) = complete_database_rows(
        original_headers,
        original_records,
        long_rows,
    )
    method_sheet_rows = method_rows(original_headers, provenance_rows)
    write_csv(
        OUTPUT_DIR / "minefield_scores_long.csv",
        long_rows,
        LONG_SCORE_FIELDS,
    )
    write_project_scoring_workbook(
        OUTPUT_DIR / "minefield_scoring.xlsx",
        long_rows,
        audit_rows,
        configurations,
        selection_rows,
        coverage_rows,
        issue_rows,
    )
    write_complete_database(
        FULL_BEST_AVAILABLE_DATABASE,
        full_database_headers,
        full_database_rows,
        long_rows,
        complete_long_rows,
        coverage_rows,
        method_sheet_rows,
        provenance_rows,
    )
    validate_complete_database(
        FULL_BEST_AVAILABLE_DATABASE,
        original_records,
        original_headers,
        full_database_headers,
        provenance_rows,
    )
    write_full_best_available_summary(
        counters,
        configurations,
        long_rows,
        audit_rows,
        coverage_rows,
        original_headers,
        provenance_rows,
    )

    print(
        f"Wrote {len(long_rows)} participant/occasion rows to "
        f"{rel(OUTPUT_DIR / 'minefield_scores_long.csv')}"
    )
    print(
        f"Wrote {len(audit_rows)} trial rows to "
        f"{rel(OUTPUT_DIR / 'minefield_scoring.xlsx')}"
    )
    print(
        "Wrote full best-available Minefield database to "
        f"{rel(FULL_BEST_AVAILABLE_DATABASE)}"
    )
    print(f"Left route-efficiency database unchanged at {rel(COMPLETE_DATABASE)}")


if __name__ == "__main__":
    main()

import pandas as pd
from scipy.stats import pearsonr
from statsmodels.stats.multitest import multipletests
import numpy as np
import openpyxl
from openpyxl.styles import PatternFill
from pathlib import Path

# Excel file path
ROOT = Path(__file__).resolve().parents[2]
OUTPUT_DIR = ROOT / "output" / "correlations"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
excel_path = ROOT / "output" / "descriptives" / "final_chess_analysis.xlsx"

# Load the Excel file
# If there are multiple sheets, adjust sheet_name accordingly
try:
    df = pd.read_excel(excel_path, sheet_name=0)
except Exception as e:
    print(f'Error loading file: {e}')
    exit(1)

# Select numeric columns only
numeric_df = df.select_dtypes(include=['number'])

# Compute the correlation and p-value matrices
cols = numeric_df.columns
n = len(cols)
correlations = np.zeros((n, n))
p_values = np.zeros((n, n))

for i in range(n):
    for j in range(n):
        if i == j:
            correlations[i, j] = 1.0
            p_values[i, j] = np.nan
        else:
            x = numeric_df.iloc[:, i].dropna()
            y = numeric_df.iloc[:, j].dropna()
            # Index intersection
            common_idx = x.index.intersection(y.index)
            x_common = x.loc[common_idx]
            y_common = y.loc[common_idx]
            # Require at least two values and non-constant inputs
            if len(x_common) < 2 or len(y_common) < 2 or x_common.nunique() < 2 or y_common.nunique() < 2:
                correlations[i, j] = np.nan
                p_values[i, j] = np.nan
                print(f'Warning: correlation not computed between {cols[i]} e {cols[j]} (insufficient or constant data)')
            else:
                corr, p = pearsonr(x_common, y_common)
                correlations[i, j] = corr
                p_values[i, j] = p

# Apply Holm correction to p-values in the upper triangle
mask = np.triu(np.ones((n, n), dtype=bool), k=1)
p_flat = p_values[mask]
_, p_corrected, _, _ = multipletests(p_flat, method='holm')
p_values_corrected = p_values.copy()
p_values_corrected[mask] = p_corrected
p_values_corrected = p_values_corrected.T
p_values_corrected[mask] = p_corrected
p_values_corrected = p_values_corrected.T

# Create an Excel file with correlations and p-values
output_path = OUTPUT_DIR / "all_pearson_correlations.xlsx"
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Correlazioni'

# Header
ws.cell(row=1, column=1, value='')
for idx, col in enumerate(cols):
    ws.cell(row=1, column=idx+2, value=col)
    ws.cell(row=idx+2, column=1, value=col)

# Write correlations and p-values
for i in range(n):
    for j in range(n):
        corr = correlations[i, j]
        pval = p_values_corrected[i, j]
        cell = ws.cell(row=i+2, column=j+2)
        if i == j:
            cell.value = 1.0
        else:
            cell.value = f'{corr:.3f}\np={pval:.3g}'
            # Highlight green when p < 0.05
            if not np.isnan(pval) and pval < 0.05:
                cell.fill = PatternFill(start_color='90EE90', end_color='90EE90', fill_type='solid')

wb.save(output_path)
print(f'Pearson correlation with p-values and Holm correction saved to {output_path}')

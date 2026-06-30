from pathlib import Path
import sys
import warnings

import numpy as np
import pandas as pd
import statsmodels.api as sm
from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill


ROOT = Path(__file__).resolve().parents[2]
sys.path.append(str(ROOT / "analysis" / "descriptives"))
sys.path.append(str(ROOT / "analysis" / "training_effect"))

from add_composites import build_composite_dataframe, to_numeric  # noqa: E402
from treatment_effectiveness import (  # noqa: E402
    ALPHA,
    bh_fdr,
    build_individual_effectiveness,
    load_metadata,
)


OUT_DIR = ROOT / "output" / "predictors" / "treatment_effectiveness"
OUT_PATH = OUT_DIR / "treatment_effectiveness_predictors.xlsx"
MIN_N = 20


def candidate_predictors(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    seen = set()
    for col in df.columns:
        if col in seen:
            continue
        col_text = str(col)
        predictor_type = None
        if col_text.startswith(("ABAS_", "BRIEF_")):
            predictor_type = "questionnaire"
        elif col_text.upper().endswith("_PRE") or col_text.endswith("_pre"):
            predictor_type = "baseline_pre"

        if not predictor_type:
            continue

        values = to_numeric(df[col])
        if values.notna().sum() < MIN_N or values.nunique(dropna=True) < 2:
            continue

        rows.append(
            {
                "predictor": col,
                "predictor_type": predictor_type,
                "n_nonmissing": int(values.notna().sum()),
                "mean": values.mean(),
                "sd": values.std(ddof=1),
                "min": values.min(),
                "max": values.max(),
            }
        )
        seen.add(col)
    return pd.DataFrame(rows).sort_values(["predictor_type", "predictor"]).reset_index(drop=True)


def zscore(series: pd.Series) -> pd.Series:
    values = to_numeric(series)
    sd = values.std(ddof=1)
    if pd.isna(sd) or sd == 0:
        return pd.Series(np.nan, index=series.index, dtype=float)
    return (values - values.mean()) / sd


def build_model_frame(effectiveness: pd.DataFrame, source_df: pd.DataFrame, predictor: str) -> pd.DataFrame:
    model_df = effectiveness[["ID", "variable", "effectiveness_pct", "group", "age", "sex", "class"]].copy()
    source = source_df[["ID", predictor]].copy()
    source[predictor] = to_numeric(source[predictor])
    model_df = model_df.merge(source, on="ID", how="left")

    model_df["effectiveness_pct"] = to_numeric(model_df["effectiveness_pct"])
    model_df["group"] = pd.to_numeric(model_df["group"], errors="coerce")
    model_df["group_experimental"] = (model_df["group"] == 1).astype(float)
    model_df["age_z"] = zscore(model_df["age"])
    model_df["predictor_z"] = zscore(model_df[predictor])
    model_df["sex"] = model_df["sex"].astype(str).str.strip().replace({"nan": np.nan, "None": np.nan, "": np.nan})
    model_df["class_cov"] = model_df["class"].astype(str).str.strip().replace({"nan": np.nan, "None": np.nan, "": np.nan})
    return model_df.dropna(
        subset=[
            "effectiveness_pct",
            "group_experimental",
            "age_z",
            "predictor_z",
            "sex",
            "class_cov",
        ]
    )


def fit_ols(model_df: pd.DataFrame, include_interaction: bool):
    x = pd.DataFrame(index=model_df.index)
    x["const"] = 1.0
    x["predictor_z"] = model_df["predictor_z"].astype(float)
    x["group_experimental"] = model_df["group_experimental"].astype(float)
    x["age_z"] = model_df["age_z"].astype(float)
    if include_interaction:
        x["predictor_x_group"] = x["predictor_z"] * x["group_experimental"]

    dummies = pd.get_dummies(model_df[["sex", "class_cov"]], columns=["sex", "class_cov"], drop_first=True)
    x = pd.concat([x, dummies.astype(float)], axis=1)
    y = model_df["effectiveness_pct"].astype(float)
    return sm.OLS(y, x).fit(cov_type="HC3")


def extract_term(fit, term: str) -> dict:
    conf = fit.conf_int()
    if term not in fit.params.index:
        return {
            "term": term,
            "estimate": np.nan,
            "se": np.nan,
            "t": np.nan,
            "p_value": np.nan,
            "ci_low": np.nan,
            "ci_high": np.nan,
        }
    return {
        "term": term,
        "estimate": fit.params[term],
        "se": fit.bse[term],
        "t": fit.tvalues[term],
        "p_value": fit.pvalues[term],
        "ci_low": conf.loc[term, 0],
        "ci_high": conf.loc[term, 1],
    }


def run_models(effectiveness: pd.DataFrame, source_df: pd.DataFrame, predictors: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    main_rows = []
    moderation_rows = []
    outcomes = sorted(effectiveness["variable"].dropna().unique())

    for _, predictor_info in predictors.iterrows():
        predictor = predictor_info["predictor"]
        model_frame_all = build_model_frame(effectiveness, source_df, predictor)
        for outcome in outcomes:
            model_df = model_frame_all[model_frame_all["variable"].eq(outcome)].copy()
            base = {
                "outcome": outcome,
                "predictor": predictor,
                "predictor_type": predictor_info["predictor_type"],
                "n": int(len(model_df)),
                "n_ids": int(model_df["ID"].nunique()) if not model_df.empty else 0,
            }

            if len(model_df) < MIN_N or model_df["group_experimental"].nunique() < 2:
                main_rows.append({**base, "status": "skipped_insufficient_data"})
                moderation_rows.append({**base, "status": "skipped_insufficient_data"})
                continue

            try:
                with warnings.catch_warnings():
                    warnings.filterwarnings("ignore", category=RuntimeWarning)
                    main_fit = fit_ols(model_df, include_interaction=False)
                main_rows.append(
                    {
                        **base,
                        "status": "ok",
                        "model": "effectiveness_pct ~ predictor_z + group_experimental + age_z + sex + class",
                        **extract_term(main_fit, "predictor_z"),
                        "r_squared": main_fit.rsquared,
                        "adj_r_squared": main_fit.rsquared_adj,
                    }
                )
            except Exception as exc:
                main_rows.append({**base, "status": f"error: {type(exc).__name__}"})

            try:
                with warnings.catch_warnings():
                    warnings.filterwarnings("ignore", category=RuntimeWarning)
                    moderation_fit = fit_ols(model_df, include_interaction=True)
                moderation_rows.append(
                    {
                        **base,
                        "status": "ok",
                        "model": "effectiveness_pct ~ predictor_z * group_experimental + age_z + sex + class",
                        **extract_term(moderation_fit, "predictor_x_group"),
                        "r_squared": moderation_fit.rsquared,
                        "adj_r_squared": moderation_fit.rsquared_adj,
                    }
                )
            except Exception as exc:
                moderation_rows.append({**base, "status": f"error: {type(exc).__name__}"})

    main = pd.DataFrame(main_rows)
    moderation = pd.DataFrame(moderation_rows)
    for frame in [main, moderation]:
        if not frame.empty and "p_value" in frame.columns:
            ok = frame["status"].eq("ok") & frame["p_value"].notna()
            frame["q_fdr_bh"] = np.nan
            if ok.any():
                frame.loc[ok, "q_fdr_bh"] = bh_fdr(frame.loc[ok, "p_value"])
            frame["significant_p05"] = frame["p_value"] < ALPHA
            frame["significant_fdr05"] = frame["q_fdr_bh"] < ALPHA
    return main, moderation


def significant_table(frame: pd.DataFrame) -> pd.DataFrame:
    if frame.empty or "significant_p05" not in frame.columns:
        return pd.DataFrame()
    return frame[frame["significant_p05"].fillna(False)].sort_values("p_value").reset_index(drop=True)


def write_workbook(
    predictors: pd.DataFrame,
    main_models: pd.DataFrame,
    moderation_models: pd.DataFrame,
    individual_effectiveness: pd.DataFrame,
) -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    readme = pd.DataFrame(
        [
            {"item": "input_database", "value": "data/FINAL_DATABASE.xlsx"},
            {"item": "effectiveness_source", "value": "recomputed in memory from analysis/shared/variable_metadata.csv"},
            {"item": "main_model", "value": "effectiveness_pct ~ predictor_z + group_experimental + age_z + sex + class"},
            {
                "item": "moderation_model",
                "value": "effectiveness_pct ~ predictor_z * group_experimental + age_z + sex + class",
            },
            {"item": "standard_errors", "value": "HC3 robust standard errors"},
            {"item": "predictor_scope", "value": "all numeric ABAS/BRIEF variables plus all numeric PRE variables"},
        ]
    )
    with pd.ExcelWriter(OUT_PATH, engine="openpyxl") as writer:
        readme.to_excel(writer, sheet_name="readme", index=False)
        predictors.to_excel(writer, sheet_name="predictor_dictionary", index=False)
        main_models.to_excel(writer, sheet_name="main_predictor_models", index=False)
        moderation_models.to_excel(writer, sheet_name="moderation_models", index=False)
        significant_table(main_models).to_excel(writer, sheet_name="significant_main", index=False)
        significant_table(moderation_models).to_excel(writer, sheet_name="significant_moderation", index=False)
        individual_effectiveness.to_excel(writer, sheet_name="effectiveness_used", index=False)
    style_workbook(OUT_PATH)


def style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    green_fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        headers = [cell.value for cell in ws[1]]
        for col_idx, header in enumerate(headers, start=1):
            col_letter = ws.cell(row=1, column=col_idx).column_letter
            max_len = len(str(header or ""))
            for cell in ws[col_letter][1:]:
                if cell.value is not None:
                    max_len = max(max_len, min(len(str(cell.value)), 45))
            ws.column_dimensions[col_letter].width = min(max_len + 2, 48)
            header_text = str(header or "").lower()
            if header_text in {"p_value", "q_fdr_bh"} or header_text.endswith("_p"):
                ws.conditional_formatting.add(
                    f"{col_letter}2:{col_letter}{ws.max_row}",
                    CellIsRule(operator="lessThan", formula=["0.05"], fill=green_fill),
                )
            if header_text in {"significant_p05", "significant_fdr05"}:
                ws.conditional_formatting.add(
                    f"{col_letter}2:{col_letter}{ws.max_row}",
                    CellIsRule(operator="equal", formula=["TRUE"], fill=yellow_fill),
                )
    wb.save(path)


def main() -> None:
    source_df = build_composite_dataframe()
    metadata, validation = load_metadata(source_df)
    valid_metadata = metadata[metadata["variable"].isin(validation.loc[validation["status"].eq("ok"), "variable"])]
    individual = build_individual_effectiveness(source_df, valid_metadata)
    individual = individual[individual["effectiveness_status"].eq("ok")].copy()

    predictors = candidate_predictors(source_df)
    main_models, moderation_models = run_models(individual, source_df, predictors)
    write_workbook(predictors, main_models, moderation_models, individual)

    main_sig = int(main_models["significant_p05"].sum()) if "significant_p05" in main_models.columns else 0
    main_fdr = int(main_models["significant_fdr05"].sum()) if "significant_fdr05" in main_models.columns else 0
    mod_sig = int(moderation_models["significant_p05"].sum()) if "significant_p05" in moderation_models.columns else 0
    mod_fdr = int(moderation_models["significant_fdr05"].sum()) if "significant_fdr05" in moderation_models.columns else 0
    print("Treatment-effectiveness predictor analysis complete")
    print(f"Predictors tested: {len(predictors)}")
    print(f"Main predictor models: {len(main_models)}; p < .05: {main_sig}; FDR < .05: {main_fdr}")
    print(f"Moderation models: {len(moderation_models)}; p < .05: {mod_sig}; FDR < .05: {mod_fdr}")
    print(f"Results workbook: {OUT_PATH}")


if __name__ == "__main__":
    main()

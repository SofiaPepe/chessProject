import pandas as pd
from scipy.stats import pearsonr
from statsmodels.stats.multitest import multipletests
import numpy as np
import openpyxl
from openpyxl.styles import PatternFill

# Percorso del file Excel
excel_path = r'C:\Users\spepe\Desktop\progetto_scacchi\descrittive\analisi_scacchi_definitive.xlsx'

# Carica il file Excel
# Se ci sono più fogli, modifica sheet_name di conseguenza
try:
    df = pd.read_excel(excel_path, sheet_name=0)
except Exception as e:
    print(f'Errore nel caricamento del file: {e}')
    exit(1)

# Seleziona solo le colonne numeriche
numeric_df = df.select_dtypes(include=['number'])

# Calcola la matrice di correlazione e p-value
cols = numeric_df.columns
n = len(cols)
correlations = np.zeros((n, n))
p_values = np.zeros((n, n))

for i in range(n):
    for j in range(n):
        if i == j:
            correlations[i, j] = 1.0
            p_values[i, j] = np.nan
        else:
            x = numeric_df.iloc[:, i].dropna()
            y = numeric_df.iloc[:, j].dropna()
            # Intersezione degli indici
            common_idx = x.index.intersection(y.index)
            x_common = x.loc[common_idx]
            y_common = y.loc[common_idx]
            # Controllo almeno 2 valori e non costanti
            if len(x_common) < 2 or len(y_common) < 2 or x_common.nunique() < 2 or y_common.nunique() < 2:
                correlations[i, j] = np.nan
                p_values[i, j] = np.nan
                print(f'Attenzione: correlazione non calcolata tra {cols[i]} e {cols[j]} (dati insufficienti o costanti)')
            else:
                corr, p = pearsonr(x_common, y_common)
                correlations[i, j] = corr
                p_values[i, j] = p

# Applica la correzione di Holm ai p-value (solo triangolo superiore)
mask = np.triu(np.ones((n, n), dtype=bool), k=1)
p_flat = p_values[mask]
_, p_corrected, _, _ = multipletests(p_flat, method='holm')
p_values_corrected = p_values.copy()
p_values_corrected[mask] = p_corrected
p_values_corrected = p_values_corrected.T
p_values_corrected[mask] = p_corrected
p_values_corrected = p_values_corrected.T

# Crea un file Excel con correlazioni e p-value
output_path = r'C:\Users\spepe\Desktop\progetto_scacchi\correlazioni\correlazione_pearson_tutte.xlsx'
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Correlazioni'

# Header
ws.cell(row=1, column=1, value='')
for idx, col in enumerate(cols):
    ws.cell(row=1, column=idx+2, value=col)
    ws.cell(row=idx+2, column=1, value=col)

# Scrivi correlazioni e p-value
for i in range(n):
    for j in range(n):
        corr = correlations[i, j]
        pval = p_values_corrected[i, j]
        cell = ws.cell(row=i+2, column=j+2)
        if i == j:
            cell.value = 1.0
        else:
            cell.value = f'{corr:.3f}\np={pval:.3g}'
            # Evidenzia in verde se p < 0.05
            if not np.isnan(pval) and pval < 0.05:
                cell.fill = PatternFill(start_color='90EE90', end_color='90EE90', fill_type='solid')

wb.save(output_path)
print(f'Correlazione di Pearson (con p-value e correzione Holm) salvata in {output_path}')

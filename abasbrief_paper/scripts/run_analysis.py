from __future__ import annotations

import hashlib
import platform
import shutil
import sys
import warnings
from datetime import datetime, timezone
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd
import scipy
import seaborn as sns
import statsmodels
import statsmodels.api as sm
from openpyxl.styles import Font, PatternFill
from scipy import stats
from statsmodels.stats.diagnostic import het_breuschpagan
from statsmodels.stats.multitest import multipletests
from statsmodels.stats.outliers_influence import variance_inflation_factor

import config as cfg


def to_numeric(series: pd.Series) -> pd.Series:
    if pd.api.types.is_numeric_dtype(series):
        return pd.to_numeric(series, errors="coerce")
    cleaned = series.astype(str).str.replace(",", ".", regex=False)
    cleaned = cleaned.replace({"nan": np.nan, "None": np.nan, "": np.nan})
    return pd.to_numeric(cleaned, errors="coerce")


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def clean_generated_outputs() -> None:
    for target, expected_parent in [
        (cfg.OUTPUT_ROOT.resolve(), cfg.SUBPROJECT_ROOT.resolve()),
        (cfg.FIGURE_ROOT.resolve(), cfg.SUBPROJECT_ROOT.resolve()),
    ]:
        if target.parent != expected_parent:
            raise RuntimeError(f"Refusing to clean unexpected directory: {target}")
        if target.exists():
            shutil.rmtree(target)
        target.mkdir(parents=True, exist_ok=True)
    cfg.TABLE_ROOT.mkdir(parents=True, exist_ok=True)


def required_source_columns() -> list[str]:
    columns = [*cfg.DEMOGRAPHIC_COLUMNS, "Raven_ACC_PRE"]
    columns.extend(pre_col for _, pre_col, _ in cfg.OUTCOMES)
    columns.extend(cfg.ABAS_STANDARD_SCALES)
    columns.extend(cfg.BRIEF_STANDARD_SCALES)
    columns.extend(cfg.BRIEF_INDICES)
    return list(dict.fromkeys(columns))


def load_and_prepare() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    if not cfg.INPUT_DATABASE.exists():
        raise FileNotFoundError(f"Input database not found: {cfg.INPUT_DATABASE}")
    source = pd.read_excel(cfg.INPUT_DATABASE, sheet_name=cfg.INPUT_SHEET)
    source.columns = source.columns.astype(str).str.strip()
    if source["ID"].duplicated().any():
        raise ValueError("Duplicate participant IDs found in source database")

    required = required_source_columns()
    missing = [column for column in required if column not in source]
    if missing:
        raise KeyError(f"Required source columns missing: {missing}")
    selected = source[required].copy()
    forbidden = [
        column
        for column in selected
        if str(column).upper().endswith("_POST") or str(column).lower() == "group"
    ]
    if forbidden:
        raise RuntimeError(f"Forbidden columns entered PRE-only analysis: {forbidden}")

    age = to_numeric(selected["age"])
    raven = to_numeric(selected["Raven_ACC_PRE"])
    cutoff = age.map(cfg.RAVEN_CUTOFFS)
    selected["raven_cutoff"] = cutoff
    selected["raven_flag"] = raven.notna() & cutoff.notna() & raven.le(cutoff)
    if len(selected) != 101:
        raise RuntimeError(f"Expected 101 source participants, found {len(selected)}")
    if int(selected["raven_flag"].sum()) != 2:
        raise RuntimeError("Raven screening did not identify exactly two participants")

    for total, components in cfg.COMPLETE_TOTALS.items():
        numeric = selected[components].apply(to_numeric)
        selected[total] = numeric.sum(axis=1, min_count=len(components))
        selected[f"{total}_components_observed"] = numeric.notna().sum(axis=1)

    primary_with_screening = selected.loc[~selected["raven_flag"]].copy()
    if len(primary_with_screening) != 99:
        raise RuntimeError(f"Expected primary n=99, found {len(primary_with_screening)}")

    screening_by_age = []
    for age_value, subset in selected.groupby("age", dropna=False):
        screening_by_age.append(
            {
                "age": age_value,
                "raven_cutoff": cfg.RAVEN_CUTOFFS.get(int(age_value)) if pd.notna(age_value) else np.nan,
                "source_n": len(subset),
                "excluded_n": int(subset["raven_flag"].sum()),
                "included_n": int((~subset["raven_flag"]).sum()),
            }
        )
    screening_summary = pd.DataFrame(screening_by_age)

    analysis = primary_with_screening.drop(
        columns=["Raven_ACC_PRE", "raven_cutoff", "raven_flag", "ID"]
    )
    forbidden_analysis = [
        column
        for column in analysis
        if "RAVEN" in str(column).upper()
        or str(column).upper().endswith("_POST")
        or str(column).lower() in {"group", "id"}
    ]
    if forbidden_analysis:
        raise RuntimeError(f"Forbidden columns in model dataset: {forbidden_analysis}")
    return selected, analysis, screening_summary


def numeric_summary(series: pd.Series) -> dict:
    values = to_numeric(series)
    observed = values.dropna()
    return {
        "n": int(observed.size),
        "missing": int(values.isna().sum()),
        "mean": observed.mean(),
        "sd": observed.std(ddof=1),
        "median": observed.median(),
        "q1": observed.quantile(0.25),
        "q3": observed.quantile(0.75),
        "minimum": observed.min(),
        "maximum": observed.max(),
        "skewness": observed.skew(),
        "excess_kurtosis": observed.kurt(),
        "n_negative": int(observed.lt(0).sum()),
    }


def build_descriptives(
    full: pd.DataFrame,
    analysis: pd.DataFrame,
    screening_summary: pd.DataFrame,
) -> dict[str, pd.DataFrame]:
    age = to_numeric(analysis["age"])
    sex = analysis["sex"].astype(str).str.upper()
    participants = pd.DataFrame(
        [
            {"characteristic": "Source database participants", "value": len(full)},
            {"characteristic": "Excluded by Raven screening", "value": len(full) - len(analysis)},
            {"characteristic": "Primary analytic sample", "value": len(analysis)},
            {"characteristic": "Age, mean", "value": age.mean()},
            {"characteristic": "Age, SD", "value": age.std(ddof=1)},
            {"characteristic": "Age, minimum", "value": age.min()},
            {"characteristic": "Age, maximum", "value": age.max()},
            {"characteristic": "Male", "value": int(sex.eq("M").sum())},
            {"characteristic": "Female", "value": int(sex.eq("F").sum())},
        ]
    )

    outcome_rows = []
    for outcome, pre_col, family in cfg.OUTCOMES:
        outcome_rows.append(
            {
                "variable": outcome,
                "label": cfg.OUTCOME_LABELS[outcome],
                "family": family,
                **numeric_summary(analysis[pre_col]),
            }
        )
    outcome_descriptives = pd.DataFrame(outcome_rows)

    predictor_columns = [
        *cfg.COMPLETE_TOTALS,
        *cfg.ABAS_STANDARD_SCALES,
        *cfg.BRIEF_STANDARD_SCALES,
        *cfg.BRIEF_INDICES,
    ]
    predictor_rows = []
    for predictor in predictor_columns:
        predictor_rows.append(
            {
                "variable": predictor,
                "label": cfg.PREDICTOR_LABELS[predictor],
                "family": "ABAS-II" if predictor.startswith("ABAS") else "BRIEF-2",
                **numeric_summary(analysis[predictor]),
            }
        )
    predictor_descriptives = pd.DataFrame(predictor_rows)

    variable_dictionary = pd.DataFrame(
        [
            {
                "role": "outcome",
                "variable": outcome,
                "source_column": pre_col,
                "label": cfg.OUTCOME_LABELS[outcome],
                "direction": (
                    "lower_better"
                    if outcome in {"TOL_VIO_REG", "PLANNING_EFF", "PWM_EFF"}
                    else "higher_better"
                ),
            }
            for outcome, pre_col, _ in cfg.OUTCOMES
        ]
        + [
            {
                "role": "predictor",
                "variable": predictor,
                "source_column": predictor,
                "label": cfg.PREDICTOR_LABELS[predictor],
                "direction": "higher_better" if predictor.startswith("ABAS") else "higher_more_difficulties",
            }
            for predictor in predictor_columns
        ]
    )

    correlation_columns = [
        "age",
        *(pre_col for _, pre_col, _ in cfg.OUTCOMES),
        *predictor_columns,
    ]
    correlation_input = analysis[correlation_columns].apply(to_numeric)
    correlations = correlation_input.corr(method="pearson")
    correlations.insert(0, "variable", correlations.index)
    correlations = correlations.reset_index(drop=True)

    return {
        "participants": participants,
        "screening_by_age": screening_summary,
        "outcome_descriptives": outcome_descriptives,
        "predictor_descriptives": predictor_descriptives,
        "variable_dictionary": variable_dictionary,
        "correlations": correlations,
    }


def zscore(series: pd.Series) -> pd.Series:
    values = to_numeric(series)
    sd = values.std(ddof=1)
    if pd.isna(sd) or sd <= 0:
        return pd.Series(np.nan, index=series.index, dtype=float)
    return (values - values.mean()) / sd


def model_diagnostics(fit, design: pd.DataFrame) -> dict:
    influence = fit.get_influence()
    cooks = influence.cooks_distance[0]
    leverage = influence.hat_matrix_diag
    studentized = influence.resid_studentized_external
    try:
        bp_p = float(het_breuschpagan(fit.resid, design)[1])
    except Exception:
        bp_p = np.nan
    threshold = 4 / fit.nobs
    return {
        "residual_skewness": stats.skew(fit.resid, bias=False),
        "residual_excess_kurtosis": stats.kurtosis(fit.resid, bias=False),
        "max_abs_studentized_residual": np.nanmax(np.abs(studentized)),
        "max_leverage": np.nanmax(leverage),
        "max_cooks_distance": np.nanmax(cooks),
        "cooks_threshold_4_over_n": threshold,
        "n_cooks_above_threshold": int(np.sum(cooks > threshold)),
        "breusch_pagan_p": bp_p,
    }


def focal_vifs(frame: pd.DataFrame, predictors: list[str], features: list[str]) -> dict[str, float]:
    if len(predictors) < 2:
        return {predictor: np.nan for predictor in predictors}
    design = sm.add_constant(frame[features], has_constant="add")
    values = {}
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        for predictor in predictors:
            position = list(design.columns).index(predictor)
            values[predictor] = float(
                variance_inflation_factor(design.to_numpy(), position)
            )
    return values


def fit_model(
    analysis: pd.DataFrame,
    *,
    block: str,
    model_type: str,
    outcome: str,
    pre_col: str,
    outcome_family: str,
    predictors: list[str],
) -> list[dict]:
    required = [pre_col, *predictors, "age"]
    complete = analysis[required].apply(to_numeric).dropna()
    base = {
        "block": block,
        "model_type": model_type,
        "outcome": outcome,
        "outcome_label": cfg.OUTCOME_LABELS[outcome],
        "pre_col": pre_col,
        "outcome_family": outcome_family,
        "n": len(complete),
    }
    if len(complete) < cfg.MIN_N:
        return [
            {**base, "predictor": predictor, "status": "insufficient_data"}
            for predictor in predictors
        ]

    standardized = complete.apply(zscore)
    if standardized.isna().any().any():
        return [
            {**base, "predictor": predictor, "status": "zero_variance"}
            for predictor in predictors
        ]

    features = [*predictors, "age"]
    design = sm.add_constant(standardized[features], has_constant="add")
    age_design = sm.add_constant(standardized[["age"]], has_constant="add")
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            fit = sm.OLS(standardized[pre_col], design).fit(cov_type="HC3")
            age_fit = sm.OLS(standardized[pre_col], age_design).fit(cov_type="HC3")
        diagnostics = model_diagnostics(fit, design)
        vifs = focal_vifs(standardized, predictors, features)
        rows = []
        for predictor in predictors:
            reduced_features = [feature for feature in features if feature != predictor]
            reduced_design = sm.add_constant(
                standardized[reduced_features], has_constant="add"
            )
            reduced_fit = sm.OLS(standardized[pre_col], reduced_design).fit()
            partial_r2 = (
                (reduced_fit.ssr - fit.ssr) / reduced_fit.ssr
                if reduced_fit.ssr > 0
                else np.nan
            )
            ci = fit.conf_int().loc[predictor]
            rows.append(
                {
                    **base,
                    "predictor": predictor,
                    "predictor_label": cfg.PREDICTOR_LABELS[predictor],
                    "predictor_family": "ABAS-II" if predictor.startswith("ABAS") else "BRIEF-2",
                    "status": "ok",
                    "beta_standardized": fit.params[predictor],
                    "se_hc3": fit.bse[predictor],
                    "ci_low": ci.iloc[0],
                    "ci_high": ci.iloc[1],
                    "p_value": fit.pvalues[predictor],
                    "r_squared": fit.rsquared,
                    "adj_r_squared": fit.rsquared_adj,
                    "age_only_r_squared": age_fit.rsquared,
                    "delta_r_squared_vs_age": fit.rsquared - age_fit.rsquared,
                    "partial_r_squared": partial_r2,
                    "vif": vifs[predictor],
                    "formula": "outcome_z ~ "
                    + " + ".join(f"{feature}_z" for feature in features),
                    **diagnostics,
                }
            )
        return rows
    except Exception as exc:
        return [
            {
                **base,
                "predictor": predictor,
                "predictor_label": cfg.PREDICTOR_LABELS[predictor],
                "status": f"{type(exc).__name__}: {exc}",
            }
            for predictor in predictors
        ]


def run_primary_models(analysis: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict] = []
    for outcome, pre_col, outcome_family in cfg.OUTCOMES:
        for block, predictors in cfg.PREDICTOR_BLOCKS.items():
            for predictor in predictors:
                rows.extend(
                    fit_model(
                        analysis,
                        block=block,
                        model_type="separate",
                        outcome=outcome,
                        pre_col=pre_col,
                        outcome_family=outcome_family,
                        predictors=[predictor],
                    )
                )
        rows.extend(
            fit_model(
                analysis,
                block="totals_joint",
                model_type="joint",
                outcome=outcome,
                pre_col=pre_col,
                outcome_family=outcome_family,
                predictors=list(cfg.COMPLETE_TOTALS),
            )
        )

    result = pd.DataFrame(rows)
    result["q_fdr_bh"] = np.nan
    valid = result["status"].eq("ok") & result["p_value"].notna()
    for _, index in result.loc[valid].groupby("block", sort=False).groups.items():
        result.loc[index, "q_fdr_bh"] = multipletests(
            result.loc[index, "p_value"], method="fdr_bh"
        )[1]
    result["significant_raw_p05"] = result["p_value"] < cfg.ALPHA
    result["significant_fdr05"] = result["q_fdr_bh"] < cfg.ALPHA
    result = result.sort_values(
        ["block", "q_fdr_bh", "p_value"], na_position="last"
    ).reset_index(drop=True)

    ok = result.loc[result["status"].eq("ok")]
    if len(ok) != 216:
        raise RuntimeError(f"Expected 216 model terms, found {len(ok)}")
    observed_sizes = ok.groupby("block").size().to_dict()
    if observed_sizes != cfg.EXPECTED_BLOCK_TERMS:
        raise RuntimeError(
            f"Unexpected block sizes: {observed_sizes}; expected {cfg.EXPECTED_BLOCK_TERMS}"
        )
    if int(ok["significant_fdr05"].sum()) != 22:
        raise RuntimeError(
            f"Expected 22 FDR-significant terms, found {int(ok['significant_fdr05'].sum())}"
        )
    if ok["formula"].str.contains("Raven|POST|group", case=False, regex=True).any():
        raise RuntimeError("Forbidden term found in primary model formula")
    return result


def build_block_summary(models: pd.DataFrame) -> pd.DataFrame:
    ok = models.loc[models["status"].eq("ok")]
    return (
        ok.groupby("block", sort=False)
        .agg(
            terms_tested=("p_value", "size"),
            raw_p05=("significant_raw_p05", "sum"),
            fdr_q05=("significant_fdr05", "sum"),
            median_n=("n", "median"),
            minimum_n=("n", "min"),
            maximum_n=("n", "max"),
        )
        .reset_index()
    )


def style_workbook(path: Path) -> None:
    workbook = openpyxl.load_workbook(path)
    fill = PatternFill("solid", fgColor="D9EAF7")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
        for column_cells in sheet.columns:
            values = [str(cell.value) if cell.value is not None else "" for cell in column_cells[:200]]
            width = min(max(max(map(len, values), default=0) + 2, 10), 45)
            sheet.column_dimensions[column_cells[0].column_letter].width = width
    workbook.save(path)


def write_workbook(path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            pd.DataFrame(frame).to_excel(writer, sheet_name=name[:31], index=False)
    style_workbook(path)


def format_p(value: float) -> str:
    if pd.isna(value):
        return ""
    if value < 0.001:
        return "<.001"
    return f"{value:.3f}".lstrip("0")


def markdown_table(frame: pd.DataFrame, columns: list[tuple[str, str]]) -> str:
    if frame.empty:
        return "No results met the specified criterion."
    lines = [
        "| " + " | ".join(label for _, label in columns) + " |",
        "|" + "|".join("---" for _ in columns) + "|",
    ]
    for _, row in frame.iterrows():
        lines.append(
            "| "
            + " | ".join(str(row.get(column, "")) for column, _ in columns)
            + " |"
        )
    return "\n".join(lines)


def paper_result_table(models: pd.DataFrame) -> pd.DataFrame:
    table = models.loc[models["significant_fdr05"].fillna(False)].copy()
    table["n_fmt"] = table["n"].astype(int).astype(str)
    table["beta_fmt"] = table["beta_standardized"].map(lambda x: f"{x:.3f}")
    table["ci_fmt"] = table.apply(
        lambda row: f"[{row['ci_low']:.3f}, {row['ci_high']:.3f}]", axis=1
    )
    table["p_fmt"] = table["p_value"].map(format_p)
    table["q_fmt"] = table["q_fdr_bh"].map(format_p)
    table["partial_r2_fmt"] = table["partial_r_squared"].map(lambda x: f"{x:.3f}")
    table["delta_r2_fmt"] = table["delta_r_squared_vs_age"].map(lambda x: f"{x:.3f}")
    return table


def write_paper_tables(
    descriptives: dict[str, pd.DataFrame],
    models: pd.DataFrame,
    block_summary: pd.DataFrame,
) -> list[Path]:
    cfg.TABLE_ROOT.mkdir(parents=True, exist_ok=True)
    files = []

    participants = descriptives["participants"].copy()
    participants["value_fmt"] = participants["value"].map(
        lambda value: f"{value:.2f}" if isinstance(value, (float, np.floating)) else str(value)
    )
    participant_md = markdown_table(
        participants, [("characteristic", "Characteristic"), ("value_fmt", "Value")]
    )

    outcomes = descriptives["outcome_descriptives"].copy()
    for column in ["mean", "sd", "median", "minimum", "maximum"]:
        outcomes[f"{column}_fmt"] = outcomes[column].map(lambda x: f"{x:.2f}")
    outcome_md = markdown_table(
        outcomes,
        [
            ("label", "Outcome"),
            ("n", "n"),
            ("missing", "Missing"),
            ("mean_fmt", "Mean"),
            ("sd_fmt", "SD"),
            ("median_fmt", "Median"),
            ("minimum_fmt", "Min"),
            ("maximum_fmt", "Max"),
        ],
    )

    summary = block_summary.copy()
    summary["median_n_fmt"] = summary["median_n"].map(lambda x: f"{x:.0f}")
    block_md = markdown_table(
        summary,
        [
            ("block", "FDR family"),
            ("terms_tested", "Terms"),
            ("raw_p05", "Raw p<.05"),
            ("fdr_q05", "FDR q<.05"),
            ("median_n_fmt", "Median n"),
        ],
    )

    paper_results = paper_result_table(models)
    totals = paper_results.loc[
        paper_results["block"].isin(["totals_separate", "totals_joint"])
    ]
    secondary = paper_results.loc[
        paper_results["block"].isin(
            ["abas_subscales", "brief_subscales", "brief_indices"]
        )
    ]
    result_columns = [
        ("predictor_label", "Predictor"),
        ("outcome_label", "Outcome"),
        ("n_fmt", "n"),
        ("beta_fmt", "β"),
        ("ci_fmt", "95% CI"),
        ("p_fmt", "p"),
        ("q_fmt", "q"),
        ("delta_r2_fmt", "ΔR²"),
        ("partial_r2_fmt", "Partial R²"),
    ]
    totals_md = markdown_table(totals, result_columns)
    secondary_md = markdown_table(secondary, result_columns)

    planning = paper_results.loc[paper_results["outcome"].eq("PLANNING_ACC")]
    narrative = [
        "The screened analytic sample comprised 99 children. Across the five prespecified ",
        "FDR families, 216 questionnaire coefficient tests were estimated and 22 survived ",
        "Benjamini–Hochberg correction. These coefficients should not be interpreted as ",
        "22 independent phenomena because correlated questionnaire scales frequently mapped ",
        "onto the same cognitive outcomes.",
        "",
        f"Five FDR-significant associations involved questionnaire total scores, and {len(secondary)} involved subscales or BRIEF-2 indices.",
        "",
        f"Planning accuracy was associated with {len(planning)} questionnaire predictors after FDR correction. No FDR-significant association was observed for either excess-tile outcome, PWM accuracy, or Corsi forward span.",
    ]

    content = {
        "table1_participants.md": participant_md + "\n",
        "table2_outcomes.md": outcome_md + "\n",
        "table3_block_summary.md": block_md + "\n",
        "table4_totals.md": totals_md + "\n",
        "table5_subscales_indices.md": secondary_md + "\n",
        "results_narrative.md": "\n".join(narrative) + "\n",
    }
    for name, text in content.items():
        path = cfg.TABLE_ROOT / name
        path.write_text(text, encoding="utf-8")
        files.append(path)
    return files


def create_figures(models: pd.DataFrame) -> list[Path]:
    cfg.FIGURE_ROOT.mkdir(parents=True, exist_ok=True)
    significant = models.loc[models["significant_fdr05"].fillna(False)].copy()
    significant["association"] = (
        significant["predictor_label"] + " → " + significant["outcome_label"]
    )
    significant = significant.sort_values(
        ["outcome_label", "beta_standardized"], ascending=[True, True]
    ).reset_index(drop=True)

    sns.set_theme(style="whitegrid")
    height = max(8, len(significant) * 0.42)
    fig, axis = plt.subplots(figsize=(10, height))
    y = np.arange(len(significant))
    colors = significant["predictor_family"].map(
        {"ABAS-II": "#2878B5", "BRIEF-2": "#D95F02"}
    )
    lower = significant["beta_standardized"] - significant["ci_low"]
    upper = significant["ci_high"] - significant["beta_standardized"]
    for index, row in significant.iterrows():
        axis.errorbar(
            row["beta_standardized"],
            index,
            xerr=[[lower.iloc[index]], [upper.iloc[index]]],
            fmt="o",
            color=colors.iloc[index],
            capsize=3,
        )
    axis.axvline(0, color="black", linewidth=1)
    axis.set_yticks(y)
    axis.set_yticklabels(significant["association"], fontsize=8)
    axis.set_xlabel("Standardized coefficient β (95% CI)")
    axis.set_ylabel("")
    axis.set_title("FDR-significant age-adjusted associations")
    axis.invert_yaxis()
    fig.tight_layout()
    forest_png = cfg.FIGURE_ROOT / "primary_forest_plot.png"
    forest_svg = cfg.FIGURE_ROOT / "primary_forest_plot.svg"
    fig.savefig(forest_png, dpi=300, bbox_inches="tight")
    fig.savefig(forest_svg, bbox_inches="tight")
    plt.close(fig)

    heat = significant.pivot_table(
        index="predictor_label",
        columns="outcome_label",
        values="beta_standardized",
        aggfunc="first",
    )
    fig, axis = plt.subplots(
        figsize=(max(8, heat.shape[1] * 1.7), max(6, heat.shape[0] * 0.45))
    )
    sns.heatmap(
        heat,
        annot=True,
        fmt=".2f",
        cmap="vlag",
        center=0,
        linewidths=0.5,
        mask=heat.isna(),
        cbar_kws={"label": "Standardized β"},
        ax=axis,
    )
    axis.set_xlabel("")
    axis.set_ylabel("")
    axis.set_title("FDR-significant predictor–outcome map")
    fig.tight_layout()
    heat_png = cfg.FIGURE_ROOT / "significant_beta_heatmap.png"
    heat_svg = cfg.FIGURE_ROOT / "significant_beta_heatmap.svg"
    fig.savefig(heat_png, dpi=300, bbox_inches="tight")
    fig.savefig(heat_svg, bbox_inches="tight")
    plt.close(fig)
    return [forest_png, forest_svg, heat_png, heat_svg]


def write_summary(
    descriptives: dict[str, pd.DataFrame],
    models: pd.DataFrame,
    block_summary: pd.DataFrame,
) -> Path:
    outcomes = descriptives["outcome_descriptives"]
    lines = [
        "# ABAS-II/BRIEF-2 PRE-only analysis",
        "",
        "- Source participants: 101",
        "- Raven-screened exclusions: 2",
        "- Primary analytic sample: 99",
        "- PRE-only outcomes: 9",
        "- Model terms tested: 216",
        f"- FDR-significant model terms: {int(models['significant_fdr05'].sum())}",
        f"- Negative PLANNING_EFF observations retained: {int(outcomes.loc[outcomes['variable'].eq('PLANNING_EFF'), 'n_negative'].iloc[0])}",
        f"- Negative PWM_EFF observations retained: {int(outcomes.loc[outcomes['variable'].eq('PWM_EFF'), 'n_negative'].iloc[0])}",
        "- Model: standardized outcome ~ standardized questionnaire predictor + standardized age",
        "- Standard errors: HC3",
        "- Multiple testing: Benjamini–Hochberg within five prespecified blocks",
        "- No POST, training group, Raven covariate, or sensitivity model was used.",
        "",
        "## Block counts",
        "",
        block_summary.to_markdown(index=False),
    ]
    path = cfg.OUTPUT_ROOT / "analysis_summary.md"
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def run() -> dict:
    np.random.seed(cfg.RANDOM_SEED)
    clean_generated_outputs()
    full, analysis, screening_summary = load_and_prepare()
    descriptives = build_descriptives(full, analysis, screening_summary)
    models = run_primary_models(analysis)
    block_summary = build_block_summary(models)
    significant = models.loc[models["significant_fdr05"].fillna(False)].copy()
    diagnostics = significant[
        [
            "block",
            "predictor",
            "outcome",
            "n",
            "residual_skewness",
            "residual_excess_kurtosis",
            "max_abs_studentized_residual",
            "max_leverage",
            "max_cooks_distance",
            "cooks_threshold_4_over_n",
            "n_cooks_above_threshold",
            "breusch_pagan_p",
            "vif",
        ]
    ].copy()

    manifest = pd.DataFrame(
        [
            {"item": "input_database", "value": str(cfg.INPUT_DATABASE.relative_to(cfg.PROJECT_ROOT))},
            {"item": "input_sha256", "value": file_sha256(cfg.INPUT_DATABASE)},
            {"item": "generated_utc", "value": datetime.now(timezone.utc).isoformat()},
            {"item": "source_n", "value": len(full)},
            {"item": "raven_excluded_n", "value": len(full) - len(analysis)},
            {"item": "primary_n", "value": len(analysis)},
            {"item": "outcomes", "value": len(cfg.OUTCOMES)},
            {"item": "model_terms", "value": int(models["status"].eq("ok").sum())},
            {"item": "fdr_significant_terms", "value": int(models["significant_fdr05"].sum())},
            {"item": "scope", "value": "PRE-only; Raven screening only; no sensitivity analyses"},
            {"item": "python", "value": sys.version},
            {"item": "platform", "value": platform.platform()},
            {"item": "pandas", "value": pd.__version__},
            {"item": "numpy", "value": np.__version__},
            {"item": "scipy", "value": scipy.__version__},
            {"item": "statsmodels", "value": statsmodels.__version__},
            {"item": "matplotlib", "value": matplotlib.__version__},
            {"item": "seaborn", "value": sns.__version__},
        ]
    )

    workbook = cfg.OUTPUT_ROOT / "analysis_results.xlsx"
    write_workbook(
        workbook,
        {
            "manifest": manifest,
            "participants": descriptives["participants"],
            "screening_by_age": descriptives["screening_by_age"],
            "variable_dictionary": descriptives["variable_dictionary"],
            "outcome_descriptives": descriptives["outcome_descriptives"],
            "predictor_descriptives": descriptives["predictor_descriptives"],
            "correlations": descriptives["correlations"],
            "model_results": models,
            "block_summary": block_summary,
            "significant_results": significant,
            "significant_diagnostics": diagnostics,
        },
    )

    models.to_csv(cfg.TABLE_ROOT / "all_primary_models.csv", index=False, encoding="utf-8-sig")
    significant.to_csv(
        cfg.TABLE_ROOT / "fdr_significant_models.csv", index=False, encoding="utf-8-sig"
    )
    descriptives["outcome_descriptives"].to_csv(
        cfg.TABLE_ROOT / "outcome_descriptives.csv", index=False, encoding="utf-8-sig"
    )
    block_summary.to_csv(
        cfg.TABLE_ROOT / "block_summary.csv", index=False, encoding="utf-8-sig"
    )
    diagnostics.to_csv(
        cfg.TABLE_ROOT / "significant_model_diagnostics.csv",
        index=False,
        encoding="utf-8-sig",
    )
    paper_files = write_paper_tables(descriptives, models, block_summary)
    figure_files = create_figures(models)
    summary = write_summary(descriptives, models, block_summary)
    return {
        "full": full,
        "analysis": analysis,
        "descriptives": descriptives,
        "models": models,
        "block_summary": block_summary,
        "files": [workbook, summary, *paper_files, *figure_files],
    }


def main() -> None:
    result = run()
    print("ABAS-II/BRIEF-2 PRE-only analysis complete")
    print(f"Primary sample: {len(result['analysis'])}")
    print(f"Model terms: {int(result['models']['status'].eq('ok').sum())}")
    print(f"FDR-significant terms: {int(result['models']['significant_fdr05'].sum())}")
    print(f"Workbook: {cfg.OUTPUT_ROOT / 'analysis_results.xlsx'}")


if __name__ == "__main__":
    main()

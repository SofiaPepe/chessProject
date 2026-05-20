from pathlib import Path
import warnings

import openpyxl
import pandas as pd
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from scipy.stats import f_oneway, ttest_ind

from add_composites import build_composite_dataframe


ROOT = Path(__file__).resolve().parents[2]
OUTPUT_DIR = ROOT / "output" / "descriptives"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
OUT_PATH = OUTPUT_DIR / "descriptive_statistics.xlsx"


def main() -> None:
    warnings.filterwarnings("ignore", category=RuntimeWarning)
    df = build_composite_dataframe()

    if "class" in df.columns:
        df["class"] = df["class"].astype(str).str.strip().str.upper()
        df["class"] = df["class"].replace(
            {
                "1A": "1",
                "1B": "1",
                "2A": "2",
                "2B": "2",
                "INFANZIA": "INF",
            }
        )

    df_num = df.select_dtypes(include="number")
    statistics = df_num.describe().T
    statistics["skewness"] = df_num.skew()
    statistics["kurtosis"] = df_num.kurtosis()

    statistics_g1 = df[df["group"] == 1][df_num.columns].describe().T
    statistics_g1["skewness"] = df[df["group"] == 1][df_num.columns].skew()
    statistics_g1["kurtosis"] = df[df["group"] == 1][df_num.columns].kurtosis()

    statistics_g2 = df[df["group"] == 2][df_num.columns].describe().T
    statistics_g2["skewness"] = df[df["group"] == 2][df_num.columns].skew()
    statistics_g2["kurtosis"] = df[df["group"] == 2][df_num.columns].kurtosis()

    crosstab = pd.crosstab(df["group"], df["age"])
    class_age_crosstab = pd.crosstab(df["class"], df["age"]) if "class" in df.columns else None
    group_age_sex_crosstab = (
        pd.crosstab([df["group"], df["age"]], df["sex"]) if "sex" in df.columns else None
    )

    sex_results = []
    if "sex" in df.columns:
        for col in df_num.columns:
            groups = df[["sex", col]].dropna()
            if groups["sex"].nunique() == 2:
                g1, g2 = groups["sex"].unique()
                vals1 = groups[groups["sex"] == g1][col]
                vals2 = groups[groups["sex"] == g2][col]
                _, p_value = ttest_ind(vals1, vals2, equal_var=False)
                sex_results.append(
                    {
                        "variable": col,
                        f"mean_value_{g1}": vals1.mean(),
                        f"mean_value_{g2}": vals2.mean(),
                        "p_value": p_value,
                    }
                )
    sex_results = pd.DataFrame(sex_results)

    age_results = []
    for col in df_num.columns:
        groups = [group[col].dropna() for _, group in df.groupby("age") if len(group[col].dropna()) > 0]
        if len(groups) > 1:
            _, p_value = f_oneway(*groups)
            row = {"variable": col, "p_value": p_value}
            for age, mean_value in df.groupby("age")[col].mean().to_dict().items():
                row[f"mean_value_age_{age}"] = mean_value
            age_results.append(row)
    age_results = pd.DataFrame(age_results)

    class_results = []
    if "class" in df.columns:
        n_classes = df["class"].nunique()
        for col in df_num.columns:
            groups = [group[col].dropna() for _, group in df.groupby("class") if len(group[col].dropna()) > 0]
            if n_classes == 2 and len(groups) == 2:
                _, p_value = ttest_ind(groups[0], groups[1], equal_var=False)
            elif len(groups) > 1:
                _, p_value = f_oneway(*groups)
            else:
                continue

            row = {"variable": col, "p_value": p_value}
            for class_value, mean_value in df.groupby("class")[col].mean().to_dict().items():
                row[f"mean_value_class_{class_value}"] = mean_value
            class_results.append(row)
    class_results = pd.DataFrame(class_results)

    with pd.ExcelWriter(OUT_PATH) as writer:
        statistics.to_excel(writer, sheet_name="All")
        statistics_g1.to_excel(writer, sheet_name="Group 1")
        statistics_g2.to_excel(writer, sheet_name="Group 2")
        crosstab.to_excel(writer, sheet_name="Group age crosstab")
        if group_age_sex_crosstab is not None:
            group_age_sex_crosstab.to_excel(writer, sheet_name="Group age sex")
        if class_age_crosstab is not None:
            class_age_crosstab.to_excel(writer, sheet_name="Class age")
        if not sex_results.empty:
            sex_results.to_excel(writer, sheet_name="Sex differences", index=False)
        age_results.to_excel(writer, sheet_name="Age differences", index=False)
        if not class_results.empty:
            class_results.to_excel(writer, sheet_name="Class differences", index=False)

    wb = openpyxl.load_workbook(OUT_PATH)
    fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")
    for sheet_name in ["Sex differences", "Age differences", "Class differences"]:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        for col in ws.iter_cols(1, ws.max_column):
            if col[0].value == "p_value":
                col_letter = col[0].column_letter
                ws.conditional_formatting.add(
                    f"{col_letter}2:{col_letter}{ws.max_row}",
                    CellIsRule(operator="lessThan", formula=["0.05"], fill=fill),
                )
                break
    wb.save(OUT_PATH)

    print(f"Descriptive statistics saved to: {OUT_PATH}")
    print("Input database: data/FINAL_DATABASE.xlsx")


if __name__ == "__main__":
    main()

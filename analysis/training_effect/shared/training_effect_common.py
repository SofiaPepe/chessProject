from pathlib import Path
import math
import sys
import warnings

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill


ROOT = Path(__file__).resolve().parents[3]
OFFICIAL_DB = ROOT / "data" / "FINAL_DATABASE.xlsx"
TRAINING_OUTPUT = ROOT / "output" / "training_effect"
PLOTS_DIR = TRAINING_OUTPUT / "plots"

DESCRIPTIVES_DIR = ROOT / "analysis" / "descriptives"
if str(DESCRIPTIVES_DIR) not in sys.path:
    sys.path.append(str(DESCRIPTIVES_DIR))

from add_composites import build_composite_dataframe  # noqa: E402

ID_COL = "ID"
GROUP_COL = "group"
COVARIATES = ["age", "sex", "class"]
ALPHA = 0.05

GROUP_LABELS = {
    1: "experimental",
    2: "control",
}

GROUP_COLORS = {
    "experimental": "#3566A5",
    "control": "#C45A11",
}


def normalize_name(value) -> str:
    return str(value).strip().upper().replace(" ", "")


def safe_filename(value, max_len: int = 120) -> str:
    safe = "".join(ch if ch.isalnum() else "_" for ch in str(value)).strip("_")
    while "__" in safe:
        safe = safe.replace("__", "_")
    return (safe or "variable")[:max_len]


def to_numeric(series: pd.Series) -> pd.Series:
    if pd.api.types.is_numeric_dtype(series):
        return pd.to_numeric(series, errors="coerce")
    cleaned = series.astype(str).str.replace(",", ".", regex=False)
    cleaned = cleaned.replace({"nan": np.nan, "None": np.nan, "": np.nan})
    return pd.to_numeric(cleaned, errors="coerce")


def load_official_database() -> pd.DataFrame:
    if not OFFICIAL_DB.exists():
        raise FileNotFoundError(f"Official database not found: {OFFICIAL_DB}")
    df = pd.read_excel(OFFICIAL_DB)
    df.columns = df.columns.str.strip()
    return standardize_metadata(df)


def load_analysis_dataset() -> pd.DataFrame:
    df = build_composite_dataframe()
    return standardize_metadata(df)


def standardize_metadata(df: pd.DataFrame) -> pd.DataFrame:
    missing = [col for col in [ID_COL, GROUP_COL] if col not in df.columns]
    if missing:
        raise KeyError(f"Missing required official column(s): {', '.join(missing)}")

    df = df.copy()
    df[GROUP_COL] = to_numeric(df[GROUP_COL]).astype("Int64")
    df["group_label"] = df[GROUP_COL].map(GROUP_LABELS)
    df["group_experimental"] = (df[GROUP_COL] == 1).astype("Int64")

    if "age" in df.columns:
        df["age"] = to_numeric(df["age"])
    if "sex" in df.columns:
        df["sex"] = df["sex"].astype(str).str.strip()
    if "class" in df.columns:
        df["class"] = df["class"].astype(str).str.strip()
        df["class_cov"] = df["class"].replace({"nan": np.nan, "None": np.nan, "": np.nan})

    return df


def variable_name_from_pre(pre_col: str, is_log_pair: bool = False) -> str:
    if is_log_pair:
        return str(pre_col)[:-7] + "_ln"
    return str(pre_col)[:-4]


def find_prepost_pairs(df: pd.DataFrame) -> list[dict[str, str]]:
    lookup = {normalize_name(col): col for col in df.columns}
    pairs = []
    seen = set()

    for col in df.columns:
        col_norm = normalize_name(col)
        is_log_pair = False

        if col_norm.endswith("_PRE_LN"):
            base_norm = col_norm[:-7]
            post_norm = f"{base_norm}_POST_LN"
            is_log_pair = True
        elif col_norm.endswith("_PRE"):
            base_norm = col_norm[:-4]
            post_norm = f"{base_norm}_POST"
        else:
            continue

        post_col = lookup.get(post_norm)
        variable_norm = f"{base_norm}_LN" if is_log_pair else base_norm
        if not post_col or variable_norm in seen:
            continue

        seen.add(variable_norm)
        pairs.append(
            {
                "variable": variable_name_from_pre(col, is_log_pair=is_log_pair),
                "pre_col": col,
                "post_col": post_col,
            }
        )

    return sorted(pairs, key=lambda item: normalize_name(item["variable"]))


def add_model_codes(long_df: pd.DataFrame) -> pd.DataFrame:
    long_df = long_df.copy()
    long_df["time"] = pd.Categorical(long_df["time"], categories=["PRE", "POST"], ordered=True)
    long_df["time_post"] = (long_df["time"].astype(str) == "POST").astype(int)
    long_df[GROUP_COL] = to_numeric(long_df[GROUP_COL]).astype("Int64")
    long_df["group_experimental"] = (long_df[GROUP_COL] == 1).astype(int)
    long_df["group_label"] = long_df[GROUP_COL].map(GROUP_LABELS)
    long_df["group_x_time"] = long_df["group_experimental"] * long_df["time_post"]
    return long_df


def create_long_format(
    df: pd.DataFrame,
    pair: dict[str, str],
    extra_cols: list[str] | tuple[str, ...] = (),
    paired_only: bool = True,
) -> pd.DataFrame:
    id_vars = [ID_COL, GROUP_COL]
    for col in extra_cols:
        if col in df.columns and col not in id_vars:
            id_vars.append(col)

    needed = id_vars + [pair["pre_col"], pair["post_col"]]
    missing = [col for col in needed if col not in df.columns]
    if missing:
        raise KeyError(f"Missing columns for {pair['variable']}: {', '.join(missing)}")

    wide = df[needed].copy()
    wide[pair["pre_col"]] = to_numeric(wide[pair["pre_col"]])
    wide[pair["post_col"]] = to_numeric(wide[pair["post_col"]])
    if paired_only:
        wide = wide.dropna(subset=[pair["pre_col"], pair["post_col"]])

    long_df = pd.concat(
        [
            wide[id_vars + [pair["pre_col"]]]
            .rename(columns={pair["pre_col"]: "value"})
            .assign(time="PRE"),
            wide[id_vars + [pair["post_col"]]]
            .rename(columns={pair["post_col"]: "value"})
            .assign(time="POST"),
        ],
        ignore_index=True,
    )
    long_df = long_df.dropna(subset=[ID_COL, GROUP_COL, "value"])
    return add_model_codes(long_df)


def add_fdr(results: pd.DataFrame, p_col: str, q_col: str = "q_fdr_bh") -> pd.DataFrame:
    results = results.copy()
    results[q_col] = np.nan
    ok = results[p_col].notna()
    if not ok.any():
        return results

    try:
        from statsmodels.stats.multitest import multipletests

        results.loc[ok, q_col] = multipletests(results.loc[ok, p_col], method="fdr_bh")[1]
    except Exception:
        warnings.warn("Could not calculate FDR correction", RuntimeWarning)
    return results


def p_to_stars(p_value) -> str:
    if pd.isna(p_value):
        return ""
    if p_value < 0.001:
        return "***"
    if p_value < 0.01:
        return "**"
    if p_value < 0.05:
        return "*"
    return "ns"


def format_p(p_value) -> str:
    if pd.isna(p_value):
        return "p = NA"
    if p_value < 0.001:
        return "p < .001"
    return f"p = {p_value:.3f}"


def summary_mean_ci(long_df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for (group_label, time), subset in long_df.groupby(["group_label", "time"], observed=True):
        values = to_numeric(subset["value"]).dropna()
        n = int(values.count())
        mean = values.mean() if n else np.nan
        sd = values.std(ddof=1) if n > 1 else np.nan
        se = sd / math.sqrt(n) if n > 1 else np.nan
        ci = 1.96 * se if pd.notna(se) else np.nan
        rows.append(
            {
                "group_label": group_label,
                "time": str(time),
                "n": n,
                "mean": mean,
                "sd": sd,
                "se": se,
                "ci95": ci,
            }
        )
    return pd.DataFrame(rows)


def plot_prepost_by_group(
    long_df: pd.DataFrame,
    variable: str,
    path: Path,
    p_value=None,
    p_label: str = "group x time",
    title: str | None = None,
    ylabel: str = "Value",
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    plot_df = long_df.dropna(subset=["value", "group_label", "time"]).copy()
    if plot_df.empty:
        return

    summary = summary_mean_ci(plot_df)
    fig, ax = plt.subplots(figsize=(7.2, 4.8))
    x_lookup = {"PRE": 0, "POST": 1}

    for group_label in ["experimental", "control"]:
        group_summary = summary[summary["group_label"] == group_label].copy()
        if group_summary.empty:
            continue
        group_summary["x"] = group_summary["time"].map(x_lookup)
        group_summary = group_summary.sort_values("x")
        color = GROUP_COLORS.get(group_label, "#333333")
        ax.errorbar(
            group_summary["x"],
            group_summary["mean"],
            yerr=group_summary["ci95"],
            marker="o",
            linewidth=2,
            capsize=4,
            color=color,
            label=group_label,
        )

    ax.set_xticks([0, 1])
    ax.set_xticklabels(["PRE", "POST"])
    ax.set_xlabel("")
    ax.set_ylabel(ylabel)
    ax.set_title(title or str(variable))
    ax.grid(axis="y", alpha=0.25)
    ax.legend(title="group", frameon=False)

    if p_value is not None and pd.notna(p_value):
        ax.text(
            0.99,
            0.98,
            f"{p_label} {format_p(p_value)}",
            transform=ax.transAxes,
            ha="right",
            va="top",
            fontsize=9,
            color="#222222",
        )

    fig.tight_layout()
    fig.savefig(path, dpi=220)
    plt.close(fig)


def write_excel_with_highlights(path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        wrote = False
        for sheet_name, frame in sheets.items():
            if frame is None:
                continue
            frame = pd.DataFrame(frame)
            safe_sheet = sheet_name[:31]
            frame.to_excel(writer, sheet_name=safe_sheet, index=False)
            wrote = True
        if not wrote:
            pd.DataFrame({"status": ["no results"]}).to_excel(writer, sheet_name="results", index=False)
    style_workbook(path)


def style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    green_fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        headers = [cell.value for cell in ws[1]]

        for col_idx, header in enumerate(headers, start=1):
            header_text = str(header or "")
            max_len = len(header_text)
            for cell in ws.iter_cols(min_col=col_idx, max_col=col_idx, min_row=2, max_row=ws.max_row):
                for item in cell:
                    if item.value is not None:
                        max_len = max(max_len, min(len(str(item.value)), 45))
            col_letter = ws.cell(row=1, column=col_idx).column_letter
            ws.column_dimensions[col_letter].width = min(max_len + 2, 48)

            header_lower = header_text.lower()
            is_p_col = (
                header_lower in {"p", "pval", "p_value", "p-unc", "p-corr", "p-adj", "q_fdr_bh"}
                or header_lower.startswith("p-")
                or header_lower.endswith("_p")
                or header_lower.endswith("_p_value")
            )
            if is_p_col:
                ws.conditional_formatting.add(
                    f"{col_letter}2:{col_letter}{ws.max_row}",
                    CellIsRule(operator="lessThan", formula=["0.05"], fill=green_fill),
                )
                for cell in ws[col_letter][1:]:
                    cell.number_format = "0.00000"

    wb.save(path)

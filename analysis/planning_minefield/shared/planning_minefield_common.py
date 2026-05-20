from pathlib import Path
from datetime import datetime
import warnings

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from scipy import stats
from statsmodels.formula.api import mixedlm
import statsmodels.api as sm
from statsmodels.stats.multitest import multipletests
from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill

warnings.filterwarnings("ignore", category=pd.errors.PerformanceWarning)


ROOT = Path(__file__).resolve().parents[3]
INPUT_XLSX = ROOT / "data" / "FINAL_DATABASE.xlsx"
OUT_DIR = ROOT / "output" / "planning_minefield"
RESULTS_XLSX = OUT_DIR / "planning_minefield_results.xlsx"

GROUP_LABELS = {
    1: "experimental",
    2: "control",
}

ALPHA = 0.05

OUTCOME_DEFS = {
    "planning_diff_pre": {
        "task": "Planning",
        "time": "PRE",
        "matt": "PLANNING_MATT_PRE",
        "min": "PLANNING_MATTI_MIN_PRE",
    },
    "planning_diff_post": {
        "task": "Planning",
        "time": "POST",
        "matt": "PLANNING_MATT_POST",
        "min": "PLANNING_MATTI_MIN_POST",
    },
    "pwm_diff_pre": {
        "task": "PWM",
        "time": "PRE",
        "matt": "PWM_MATT_PRE",
        "min": "PWM_MATT_MIN_PRE",
    },
    "pwm_diff_post": {
        "task": "PWM",
        "time": "POST",
        "matt": "PWM_MATT_POST",
        "min": "PWM_MATT_MIN_POST",
    },
}

SUM_OUTCOMES = {
    "planning_diff_sum": ("planning_diff_pre", "planning_diff_post"),
    "pwm_diff_sum": ("pwm_diff_pre", "pwm_diff_post"),
}


def norm_name(value):
    return str(value).upper().replace(" ", "").strip()


def find_col(df, target):
    target_norm = norm_name(target)
    matches = [col for col in df.columns if norm_name(col) == target_norm]
    if not matches:
        raise KeyError(f"Column not found: {target}")
    return matches[0]


def to_numeric(series):
    if pd.api.types.is_numeric_dtype(series):
        return pd.to_numeric(series, errors="coerce")
    return pd.to_numeric(
        series.astype(str).str.replace(",", ".", regex=False),
        errors="coerce",
    )


def add_outcomes(df):
    invalid_rows = []
    checks = []

    for outcome, spec in OUTCOME_DEFS.items():
        matt_col = find_col(df, spec["matt"])
        min_col = find_col(df, spec["min"])

        matt = to_numeric(df[matt_col])
        min_tiles = to_numeric(df[min_col])
        raw = matt - min_tiles

        invalid_negative = raw < 0
        invalid_high = raw > 50
        invalid = invalid_negative | invalid_high

        df[f"{outcome}_raw"] = raw
        df[f"{outcome}_excluded"] = invalid.fillna(False)
        df[f"{outcome}_excluded_reason"] = np.select(
            [invalid_negative.fillna(False), invalid_high.fillna(False)],
            ["negative", "above_50"],
            default="",
        )
        df[outcome] = raw.where(~invalid)

        for idx in df.index[invalid.fillna(False)]:
            invalid_rows.append(
                {
                    "ID": df.loc[idx, "ID"],
                    "group": df.loc[idx, "group"],
                    "group_label": df.loc[idx, "group_label"],
                    "sex": df.loc[idx, "sex"],
                    "age": df.loc[idx, "age"],
                    "task": spec["task"],
                    "time": spec["time"],
                    "outcome": outcome,
                    "raw_diff": raw.loc[idx],
                    "reason": df.loc[idx, f"{outcome}_excluded_reason"],
                    "matt_col": matt_col,
                    "matt_value": matt.loc[idx],
                    "min_col": min_col,
                    "min_value": min_tiles.loc[idx],
                }
            )

        checks.append(
            {
                "outcome": outcome,
                "task": spec["task"],
                "time": spec["time"],
                "raw_n": int(raw.notna().sum()),
                "valid_n": int(df[outcome].notna().sum()),
                "negative_n": int(invalid_negative.sum()),
                "above_50_n": int(invalid_high.sum()),
                "missing_after_cleaning": int(df[outcome].isna().sum()),
                "raw_min": raw.min(),
                "raw_max": raw.max(),
            }
        )

    return pd.DataFrame(invalid_rows), pd.DataFrame(checks)


def add_sum_outcomes(df):
    for out_col, source_cols in SUM_OUTCOMES.items():
        df[out_col] = df[list(source_cols)].sum(axis=1, min_count=2)


def summarize_numeric(values):
    values = pd.Series(values).dropna()
    if values.empty:
        return {
            "n": 0,
            "mean": np.nan,
            "sd": np.nan,
            "median": np.nan,
            "q1": np.nan,
            "q3": np.nan,
            "min": np.nan,
            "max": np.nan,
        }
    return {
        "n": int(values.count()),
        "mean": values.mean(),
        "sd": values.std(ddof=1),
        "median": values.median(),
        "q1": values.quantile(0.25),
        "q3": values.quantile(0.75),
        "min": values.min(),
        "max": values.max(),
    }


def build_descriptives(df, checks):
    rows = []

    rows.append(
        {
            "section": "sample",
            "variable": "N_total",
            "level": "all",
            "n": len(df),
        }
    )

    for variable in ["group_label", "sex", "age", "class", "section"]:
        if variable not in df.columns:
            continue
        counts = df[variable].value_counts(dropna=False).sort_index()
        for level, n in counts.items():
            rows.append(
                {
                    "section": "sample_counts",
                    "variable": variable,
                    "level": level,
                    "n": int(n),
                }
            )

    for group_label in ["all"] + list(GROUP_LABELS.values()):
        subset = df if group_label == "all" else df[df["group_label"] == group_label]
        for outcome, spec in OUTCOME_DEFS.items():
            stats_row = summarize_numeric(subset[outcome])
            check = checks.loc[checks["outcome"] == outcome].iloc[0]
            rows.append(
                {
                    "section": "outcome_descriptives",
                    "task": spec["task"],
                    "time": spec["time"],
                    "group": group_label,
                    "variable": outcome,
                    "missing": int(subset[outcome].isna().sum()),
                    "excluded_invalid": int(subset[f"{outcome}_excluded"].sum()),
                    "raw_n_total": int(check["raw_n"]) if group_label == "all" else np.nan,
                    **stats_row,
                }
            )
        for outcome, source_cols in SUM_OUTCOMES.items():
            stats_row = summarize_numeric(subset[outcome])
            rows.append(
                {
                    "section": "outcome_sum_descriptives",
                    "task": "Planning" if outcome.startswith("planning") else "PWM",
                    "time": "PRE_PLUS_POST",
                    "group": group_label,
                    "variable": outcome,
                    "missing": int(subset[outcome].isna().sum()),
                    "source": " + ".join(source_cols),
                    **stats_row,
                }
            )

    for sex, group in pd.crosstab(df["group_label"], df["sex"]).stack().items():
        rows.append(
            {
                "section": "group_x_sex",
                "variable": "group_x_sex",
                "level": f"{sex[0]} | {sex[1]}",
                "n": int(group),
            }
        )

    for age, group in pd.crosstab(df["group_label"], df["age"]).stack().items():
        rows.append(
            {
                "section": "group_x_age",
                "variable": "group_x_age",
                "level": f"{age[0]} | {age[1]}",
                "n": int(group),
            }
        )

    return pd.DataFrame(rows)


def paired_test(df, pre_col, post_col, task, group_label="all"):
    subset = pd.DataFrame(
        {
            pre_col: to_numeric(df[pre_col]),
            post_col: to_numeric(df[post_col]),
        }
    ).dropna()
    if len(subset) < 3:
        return {
            "task": task,
            "group": group_label,
            "n_pairs": len(subset),
            "status": "skipped_insufficient_pairs",
        }

    pre = subset[pre_col]
    post = subset[post_col]
    diff = post - pre

    t_stat, t_p = stats.ttest_rel(post, pre, nan_policy="omit")
    try:
        w_stat, w_p = stats.wilcoxon(post, pre, zero_method="wilcox")
    except ValueError:
        w_stat, w_p = np.nan, np.nan

    return {
        "task": task,
        "group": group_label,
        "status": "ok",
        "n_pairs": int(len(subset)),
        "pre_mean": pre.mean(),
        "pre_sd": pre.std(ddof=1),
        "post_mean": post.mean(),
        "post_sd": post.std(ddof=1),
        "change_mean_post_minus_pre": diff.mean(),
        "change_sd": diff.std(ddof=1),
        "cohen_dz": diff.mean() / diff.std(ddof=1) if diff.std(ddof=1) else np.nan,
        "diff_skewness": diff.skew(),
        "paired_t": t_stat,
        "paired_t_df": int(len(subset) - 1),
        "paired_t_p": t_p,
        "paired_t_significant_p05": bool(t_p < ALPHA) if pd.notna(t_p) else np.nan,
        "wilcoxon_w": w_stat,
        "wilcoxon_p": w_p,
        "wilcoxon_significant_p05": bool(w_p < ALPHA) if pd.notna(w_p) else np.nan,
    }


def run_prepost_ttests(df):
    rows = []
    for group_label, subset in [
        ("all", df),
        ("experimental", df[df["group_label"] == "experimental"]),
        ("control", df[df["group_label"] == "control"]),
    ]:
        rows.append(
            paired_test(
                subset,
                "planning_diff_pre",
                "planning_diff_post",
                "Planning",
                group_label,
            )
        )
        rows.append(
            paired_test(
                subset,
                "pwm_diff_pre",
                "pwm_diff_post",
                "PWM",
                group_label,
            )
        )

    results = pd.DataFrame(rows)
    ok_mask = results["status"].eq("ok") & results["paired_t_p"].notna()
    results["paired_t_q_fdr_bh"] = np.nan
    if ok_mask.any():
        results.loc[ok_mask, "paired_t_q_fdr_bh"] = multipletests(
            results.loc[ok_mask, "paired_t_p"],
            method="fdr_bh",
        )[1]
    results["paired_t_significant_fdr05"] = results["paired_t_q_fdr_bh"] < ALPHA
    return results


def variable_name_from_pre(pre_col):
    text = str(pre_col).strip()
    return text[:-4] if text.upper().endswith("_PRE") else norm_name(text)[:-4]


def find_prepost_pairs(df):
    lookup = {norm_name(col): col for col in df.columns}
    pairs = []
    seen = set()

    for col in df.columns:
        col_norm = norm_name(col)
        if not col_norm.endswith("_PRE"):
            continue
        post_norm = f"{col_norm[:-4]}_POST"
        post_col = lookup.get(post_norm)
        if not post_col or col_norm in seen:
            continue
        seen.add(col_norm)
        pairs.append(
            {
                "variable": variable_name_from_pre(col),
                "pre_col": col,
                "post_col": post_col,
            }
        )

    return pairs


def run_all_prepost_ttests(df, prepost_pairs):
    rows = []
    for pair in prepost_pairs:
        for group_label, subset in [
            ("all", df),
            ("experimental", df[df["group_label"] == "experimental"]),
            ("control", df[df["group_label"] == "control"]),
        ]:
            result = paired_test(
                subset,
                pair["pre_col"],
                pair["post_col"],
                pair["variable"],
                group_label,
            )
            result["variable"] = pair["variable"]
            result["pre_col"] = pair["pre_col"]
            result["post_col"] = pair["post_col"]
            rows.append(result)

    results = pd.DataFrame(rows)
    first_cols = ["variable", "pre_col", "post_col", "group"]
    other_cols = [col for col in results.columns if col not in first_cols + ["task"]]
    results = results[first_cols + other_cols]

    ok_mask = results["status"].eq("ok") & results["paired_t_p"].notna()
    results["paired_t_q_fdr_bh"] = np.nan
    if ok_mask.any():
        results.loc[ok_mask, "paired_t_q_fdr_bh"] = multipletests(
            results.loc[ok_mask, "paired_t_p"],
            method="fdr_bh",
        )[1]
    results["paired_t_significant_fdr05"] = results["paired_t_q_fdr_bh"] < ALPHA
    return results


def run_prepost_without_group(df):
    return pd.DataFrame(
        [
            paired_test(
                df,
                "planning_diff_pre",
                "planning_diff_post",
                "Planning",
            ),
            paired_test(
                df,
                "pwm_diff_pre",
                "pwm_diff_post",
                "PWM",
            ),
        ]
    )


def long_task_df(df, task):
    if task == "Planning":
        pre_col = "planning_diff_pre"
        post_col = "planning_diff_post"
    else:
        pre_col = "pwm_diff_pre"
        post_col = "pwm_diff_post"

    long_df = pd.concat(
        [
            df[["ID", "group", "group_label", "sex", "age", pre_col]].rename(
                columns={pre_col: "value"}
            ).assign(time="PRE", time_post=0),
            df[["ID", "group", "group_label", "sex", "age", post_col]].rename(
                columns={post_col: "value"}
            ).assign(time="POST", time_post=1),
        ],
        ignore_index=True,
    )
    long_df["group_control"] = (long_df["group"] == 2).astype(int)
    long_df["interaction"] = long_df["time_post"] * long_df["group_control"]
    return long_df.dropna(subset=["value", "group_control", "time_post"])


def fit_repeated_model(long_df, task):
    terms = ["Intercept", "time_post", "group_control", "interaction"]

    try:
        last_error = None
        for method in ["lbfgs", "powell", "cg", "bfgs", "nm"]:
            try:
                with warnings.catch_warnings():
                    warnings.simplefilter("ignore")
                    model = mixedlm(
                        "value ~ time_post + group_control + interaction",
                        long_df,
                        groups=long_df["ID"],
                    )
                    fit = model.fit(reml=False, method=method, maxiter=2000, disp=False)
                break
            except Exception as exc:
                fit = None
                last_error = exc
        if fit is None:
            raise last_error
        model_type = "mixedlm_random_intercept"
        params = fit.params
        bse = fit.bse
        pvalues = fit.pvalues
        conf = fit.conf_int()
        statistic_label = "z_or_t"
        statistic = params / bse
    except Exception as exc:
        x = sm.add_constant(long_df[["time_post", "group_control", "interaction"]])
        fit = sm.OLS(long_df["value"], x).fit(
            cov_type="cluster",
            cov_kwds={"groups": long_df["ID"]},
        )
        model_type = f"ols_cluster_fallback: {type(exc).__name__}"
        params = fit.params
        bse = fit.bse
        pvalues = fit.pvalues
        conf = fit.conf_int()
        statistic_label = "z_or_t"
        statistic = params / bse
        terms = ["const", "time_post", "group_control", "interaction"]

    rows = []
    for term in terms:
        if term not in params.index:
            continue
        rows.append(
            {
                "analysis": "mixed_repeated",
                "task": task,
                "model_type": model_type,
                "term": "Intercept" if term == "const" else term,
                "estimate": params.loc[term],
                "se": bse.loc[term],
                statistic_label: statistic.loc[term],
                "p_value": pvalues.loc[term],
                "ci_low": conf.loc[term].iloc[0],
                "ci_high": conf.loc[term].iloc[1],
                "n_obs": int(long_df.shape[0]),
                "n_ids": int(long_df["ID"].nunique()),
            }
        )
    return rows


def run_prepost_with_group(df):
    rows = []
    for task in ["Planning", "PWM"]:
        long_df = long_task_df(df, task)
        rows.extend(fit_repeated_model(long_df, task))

        pre_col = "planning_diff_pre" if task == "Planning" else "pwm_diff_pre"
        post_col = "planning_diff_post" if task == "Planning" else "pwm_diff_post"
        for group_label in GROUP_LABELS.values():
            group_df = df[df["group_label"] == group_label]
            result = paired_test(group_df, pre_col, post_col, task, group_label)
            result.update({"analysis": "paired_within_group"})
            rows.append(result)

    return pd.DataFrame(rows)


def predictor_candidates(df, time):
    time = time.upper()
    suffix = f"_{time}"
    timed_prefixes = ("CBT_", "TOL_", "RAVEN_", "WM_")
    stable_prefixes = ("ABAS_", "BRIEF_")
    excluded_fragments = ("PLANNING", "PWM")

    timed = []
    stable = []
    for col in df.columns:
        col_norm = norm_name(col)
        if any(fragment in col_norm for fragment in excluded_fragments):
            continue
        if col_norm.endswith(suffix) and col_norm.startswith(timed_prefixes):
            timed.append(col)
        elif col_norm.startswith(stable_prefixes):
            stable.append(col)

    return sorted(set(timed + stable), key=lambda x: norm_name(x))


def add_mixed_predictors(df):
    timed_prefixes = ("CBT_", "TOL_", "RAVEN_", "WM_")
    stable_prefixes = ("ABAS_", "BRIEF_")
    excluded_fragments = ("PLANNING", "PWM")
    metadata = []

    columns_by_norm = {norm_name(col): col for col in df.columns}
    for pre_col in df.columns:
        pre_norm = norm_name(pre_col)
        if not pre_norm.endswith("_PRE"):
            continue
        if not pre_norm.startswith(timed_prefixes):
            continue
        if any(fragment in pre_norm for fragment in excluded_fragments):
            continue

        post_norm = f"{pre_norm[:-4]}_POST"
        post_col = columns_by_norm.get(post_norm)
        if not post_col:
            continue

        base_name = str(pre_col).strip()[:-4]
        safe_base = "".join(ch if ch.isalnum() else "_" for ch in base_name).strip("_")
        model_col = f"mixsum_{safe_base}"
        df[model_col] = pd.concat(
            [to_numeric(df[pre_col]), to_numeric(df[post_col])],
            axis=1,
        ).sum(axis=1, min_count=2)
        metadata.append(
            {
                "model_col": model_col,
                "predictor": f"{base_name}_PRE_PLUS_POST",
                "predictor_type": "time_varying_sum",
                "predictor_pre": pre_col,
                "predictor_post": post_col,
            }
        )

    for col in df.columns:
        col_norm = norm_name(col)
        if col_norm.startswith(stable_prefixes):
            metadata.append(
                {
                    "model_col": col,
                    "predictor": col,
                    "predictor_type": "stable",
                    "predictor_pre": "",
                    "predictor_post": "",
                }
            )

    return sorted(metadata, key=lambda item: norm_name(item["predictor"]))


def fit_predictor_model(df, outcome, predictor):
    needed = ["group", "sex", "age", outcome, predictor]
    model_df = df[needed].copy()
    model_df[outcome] = to_numeric(model_df[outcome])
    model_df[predictor] = to_numeric(model_df[predictor])
    model_df["age"] = to_numeric(model_df["age"])
    model_df = model_df.dropna()

    if len(model_df) < 20 or model_df[predictor].nunique() < 2:
        return {
            "outcome": outcome,
            "predictor": predictor,
            "n": int(len(model_df)),
            "status": "skipped_insufficient_data_or_constant",
        }

    x = pd.DataFrame(index=model_df.index)
    x["const"] = 1.0
    x["age"] = model_df["age"].astype(float)
    x[predictor] = model_df[predictor].astype(float)

    group_dummies = pd.get_dummies(model_df["group"].map(GROUP_LABELS), prefix="group", drop_first=True)
    sex_dummies = pd.get_dummies(model_df["sex"].astype(str), prefix="sex", drop_first=True)
    x = pd.concat([x, group_dummies, sex_dummies], axis=1)
    x = x.astype(float)
    y = model_df[outcome].astype(float)

    try:
        fit = sm.OLS(y, x).fit()
        ci_low, ci_high = fit.conf_int().loc[predictor].tolist()
        return {
            "outcome": outcome,
            "predictor": predictor,
            "n": int(len(model_df)),
            "status": "ok",
            "beta": fit.params[predictor],
            "se": fit.bse[predictor],
            "t": fit.tvalues[predictor],
            "p_value": fit.pvalues[predictor],
            "ci_low": ci_low,
            "ci_high": ci_high,
            "r_squared": fit.rsquared,
            "adj_r_squared": fit.rsquared_adj,
            "covariates": "group + sex + age",
        }
    except Exception as exc:
        return {
            "outcome": outcome,
            "predictor": predictor,
            "n": int(len(model_df)),
            "status": f"error: {type(exc).__name__}",
        }


def run_predictor_models(df, time):
    time = time.upper()
    outcomes = (
        ["planning_diff_pre", "pwm_diff_pre"]
        if time == "PRE"
        else ["planning_diff_post", "pwm_diff_post"]
    )
    rows = []
    for outcome in outcomes:
        for predictor in predictor_candidates(df, time):
            rows.append(fit_predictor_model(df, outcome, predictor))

    results = pd.DataFrame(rows)
    ok_mask = results["status"].eq("ok") & results["p_value"].notna()
    results["q_fdr_bh"] = np.nan
    if ok_mask.any():
        results.loc[ok_mask, "q_fdr_bh"] = multipletests(
            results.loc[ok_mask, "p_value"],
            method="fdr_bh",
        )[1]
    results.insert(0, "time", time)
    return results


def run_predictor_models_mix(df, mixed_predictors):
    rows = []
    for outcome in SUM_OUTCOMES:
        for predictor_info in mixed_predictors:
            result = fit_predictor_model(df, outcome, predictor_info["model_col"])
            result["predictor"] = predictor_info["predictor"]
            result["predictor_model_col"] = predictor_info["model_col"]
            result["predictor_type"] = predictor_info["predictor_type"]
            result["predictor_pre"] = predictor_info["predictor_pre"]
            result["predictor_post"] = predictor_info["predictor_post"]
            rows.append(result)

    results = pd.DataFrame(rows)
    ok_mask = results["status"].eq("ok") & results["p_value"].notna()
    results["q_fdr_bh"] = np.nan
    if ok_mask.any():
        results.loc[ok_mask, "q_fdr_bh"] = multipletests(
            results.loc[ok_mask, "p_value"],
            method="fdr_bh",
        )[1]
    results.insert(0, "time", "PRE_PLUS_POST")
    return results


def p_to_stars(p_value):
    if pd.isna(p_value):
        return ""
    if p_value < 0.001:
        return "***"
    if p_value < 0.01:
        return "**"
    if p_value < 0.05:
        return "*"
    return "ns"


def safe_filename(value):
    safe = "".join(ch if ch.isalnum() else "_" for ch in str(value)).strip("_")
    while "__" in safe:
        safe = safe.replace("__", "_")
    return safe[:120] or "variable"


def add_sig_bracket(ax, x1, x2, y, height, label, text_color="#222222"):
    ax.plot([x1, x1, x2, x2], [y, y + height, y + height, y], color="#222222", linewidth=1)
    ax.text(
        (x1 + x2) / 2,
        y + height,
        label,
        ha="center",
        va="bottom",
        color=text_color,
        fontsize=11,
    )


def annotation_lookup(results, variable_col):
    lookup = {}
    for _, row in results.iterrows():
        key = (str(row[variable_col]), str(row["group"]))
        lookup[key] = {
            "p_value": row.get("paired_t_p", np.nan),
            "q_fdr_bh": row.get("paired_t_q_fdr_bh", np.nan),
        }
    return lookup


def annotation_color(p_value, q_fdr_bh):
    if pd.isna(p_value) or p_value >= ALPHA:
        return "#555555"
    if pd.notna(q_fdr_bh) and q_fdr_bh < ALPHA:
        return "#222222"
    return "#C00000"


def make_plot(
    df,
    task,
    pre_col,
    post_col,
    path,
    ylabel="Tile difference",
    title=None,
    annotations=None,
    annotation_key=None,
):
    groups = [("experimental", 1), ("control", 2)]
    plot_data = []
    labels = []
    colors = []
    bracket_info = []

    for label, group_code in groups:
        subset = df[df["group"] == group_code]
        pre_values = to_numeric(subset[pre_col])
        post_values = to_numeric(subset[post_col])
        plot_data.extend([pre_values.dropna(), post_values.dropna()])
        labels.extend([f"{label}\nPRE", f"{label}\nPOST"])
        colors.extend(["#4C78A8", "#72B7B2"] if group_code == 1 else ["#F58518", "#E45756"])

        paired = pd.DataFrame({pre_col: pre_values, post_col: post_values}).dropna()
        ann = annotations.get((annotation_key, label), {}) if annotations and annotation_key else {}
        if ann:
            p_value = ann.get("p_value", np.nan)
            q_fdr_bh = ann.get("q_fdr_bh", np.nan)
            label_text = p_to_stars(p_value)
        elif len(paired) >= 3:
            _, p_value = stats.ttest_rel(paired[post_col], paired[pre_col], nan_policy="omit")
            q_fdr_bh = np.nan
            label_text = p_to_stars(p_value)
        else:
            p_value = np.nan
            q_fdr_bh = np.nan
            label_text = ""
        bracket_info.append(
            {
                "group": label,
                "p_value": p_value,
                "q_fdr_bh": q_fdr_bh,
                "label": label_text,
                "text_color": annotation_color(p_value, q_fdr_bh),
                "local_max": paired[[pre_col, post_col]].max().max() if not paired.empty else np.nan,
            }
        )

    fig, ax = plt.subplots(figsize=(9, 5.5))
    positions = np.arange(1, len(plot_data) + 1)
    violins = ax.violinplot(plot_data, positions=positions, showmeans=False, showmedians=True)
    for body, color in zip(violins["bodies"], colors):
        body.set_facecolor(color)
        body.set_edgecolor("#333333")
        body.set_alpha(0.35)

    box = ax.boxplot(
        plot_data,
        positions=positions,
        widths=0.22,
        patch_artist=True,
        showfliers=True,
        medianprops={"color": "#111111", "linewidth": 1.5},
    )
    for patch, color in zip(box["boxes"], colors):
        patch.set_facecolor(color)
        patch.set_alpha(0.65)
        patch.set_edgecolor("#333333")

    ax.set_title(title or f"{task}: tiles traveled - minimum tiles")
    ax.set_ylabel(ylabel)
    ax.set_xticks(positions)
    ax.set_xticklabels(labels)
    ax.axhline(0, color="#666666", linewidth=0.8)
    ax.grid(axis="y", alpha=0.25)

    all_values = pd.concat([pd.Series(values) for values in plot_data]).dropna()
    y_min = all_values.min() if not all_values.empty else 0
    y_max = all_values.max() if not all_values.empty else 1
    y_range = max(y_max - y_min, 1)
    bracket_height = y_range * 0.035
    bracket_offset = y_range * 0.08

    for idx, info in enumerate(bracket_info):
        if not info["label"]:
            continue
        x1 = 1 + idx * 2
        x2 = x1 + 1
        local_max = info["local_max"] if pd.notna(info["local_max"]) else y_max
        y = local_max + bracket_offset
        add_sig_bracket(ax, x1, x2, y, bracket_height, info["label"], info["text_color"])
        y_max = max(y_max, y + bracket_height + bracket_offset)

    ax.set_ylim(y_min - y_range * 0.08, y_max + y_range * 0.08)
    fig.text(
        0.5,
        0.01,
        "Black asterisks: q FDR < .05; red asterisks: uncorrected p < .05 only",
        ha="center",
        va="bottom",
        fontsize=8,
        color="#333333",
    )
    fig.tight_layout(rect=(0, 0.04, 1, 1))
    fig.savefig(path, dpi=200)
    plt.close(fig)


def make_all_prepost_plots(df, prepost_pairs, all_prepost_ttests):
    plot_dir = OUT_DIR / "prepost_plots_all_variables"
    plot_dir.mkdir(exist_ok=True)
    manifest = []
    annotations = annotation_lookup(all_prepost_ttests, "variable")

    for pair in prepost_pairs:
        filename = f"{safe_filename(pair['variable'])}_prepost_by_group.png"
        path = plot_dir / filename
        make_plot(
            df,
            pair["variable"],
            pair["pre_col"],
            pair["post_col"],
            path,
            ylabel="Value",
            title=f"{pair['variable']}: PRE vs POST by group",
            annotations=annotations,
            annotation_key=pair["variable"],
        )
        manifest.append(
            {
                "variable": pair["variable"],
                "pre_col": pair["pre_col"],
                "post_col": pair["post_col"],
                "plot_path": str(path),
            }
        )

    return pd.DataFrame(manifest)


def write_outputs(
    df,
    descriptives,
    invalid_values,
    ttests_pre_post,
    all_prepost_ttests,
    all_prepost_plots,
    prepost_without_group,
    prepost_with_group,
    predictors_pre,
    predictors_post,
    predictors_mix,
):
    results_path = RESULTS_XLSX
    try:
        write_results_workbook(
            results_path,
            descriptives,
            invalid_values,
            ttests_pre_post,
            all_prepost_ttests,
            all_prepost_plots,
            prepost_without_group,
            prepost_with_group,
            predictors_pre,
            predictors_post,
            predictors_mix,
        )
    except PermissionError:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        results_path = OUT_DIR / f"{RESULTS_XLSX.stem}_updated_{timestamp}.xlsx"
        write_results_workbook(
            results_path,
            descriptives,
            invalid_values,
            ttests_pre_post,
            all_prepost_ttests,
            all_prepost_plots,
            prepost_without_group,
            prepost_with_group,
            predictors_pre,
            predictors_post,
            predictors_mix,
        )

    style_workbook(results_path)
    return results_path


def write_results_workbook(
    path,
    descriptives,
    invalid_values,
    ttests_pre_post,
    all_prepost_ttests,
    all_prepost_plots,
    prepost_without_group,
    prepost_with_group,
    predictors_pre,
    predictors_post,
    predictors_mix,
):
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        descriptives.to_excel(writer, sheet_name="descriptives", index=False)
        invalid_values.to_excel(writer, sheet_name="excluded_values", index=False)
        ttests_pre_post.to_excel(writer, sheet_name="t_test_pre_post", index=False)
        all_prepost_ttests.to_excel(writer, sheet_name="t_test_pre_post_all", index=False)
        all_prepost_plots.to_excel(writer, sheet_name="all_variable_plots", index=False)
        prepost_without_group.to_excel(writer, sheet_name="prepost_without_group", index=False)
        prepost_with_group.to_excel(writer, sheet_name="prepost_with_group", index=False)
        predictors_pre.to_excel(writer, sheet_name="predictors_pre", index=False)
        predictors_post.to_excel(writer, sheet_name="predictors_post", index=False)
        predictors_mix.to_excel(writer, sheet_name="predictors_mixed_sum", index=False)


def style_workbook(path):
    wb = load_workbook(path)
    green_fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")
    p_headers = {"p_value", "paired_t_p", "paired_t_q_fdr_bh", "wilcoxon_p", "q_fdr_bh"}

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        headers = [cell.value for cell in ws[1]]
        for col_idx, header in enumerate(headers, start=1):
            header_text = str(header)
            max_len = len(header_text)
            for cell in ws.iter_cols(min_col=col_idx, max_col=col_idx, min_row=2, max_row=ws.max_row):
                for item in cell:
                    if item.value is not None:
                        max_len = max(max_len, min(len(str(item.value)), 40))
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(max_len + 2, 42)

            if header_text in p_headers or header_text.endswith("_p"):
                col_letter = ws.cell(row=1, column=col_idx).column_letter
                ws.conditional_formatting.add(
                    f"{col_letter}2:{col_letter}{ws.max_row}",
                    CellIsRule(operator="lessThan", formula=["0.05"], fill=green_fill),
                )
                for cell in ws[col_letter][1:]:
                    cell.number_format = "0.00000"

    wb.save(path)


def load_planning_dataset():
    df = pd.read_excel(INPUT_XLSX)
    df.columns = df.columns.str.strip()
    df["group"] = to_numeric(df["group"]).astype("Int64")
    df["group_label"] = df["group"].map(GROUP_LABELS)
    df["age"] = to_numeric(df["age"])
    df["sex"] = df["sex"].astype(str).str.strip()
    return df


def main():
    OUT_DIR.mkdir(exist_ok=True)

    df = load_planning_dataset()

    prepost_pairs = find_prepost_pairs(df)
    invalid_values, checks = add_outcomes(df)
    add_sum_outcomes(df)
    mixed_predictors = add_mixed_predictors(df)
    descriptives = build_descriptives(df, checks)
    ttests_pre_post = run_prepost_ttests(df)
    all_prepost_ttests = run_all_prepost_ttests(df, prepost_pairs)
    all_prepost_plots = make_all_prepost_plots(df, prepost_pairs, all_prepost_ttests)
    prepost_without_group = run_prepost_without_group(df)
    prepost_with_group = run_prepost_with_group(df)
    predictors_pre = run_predictor_models(df, "PRE")
    predictors_post = run_predictor_models(df, "POST")
    predictors_mix = run_predictor_models_mix(df, mixed_predictors)
    planning_pwm_annotations = annotation_lookup(ttests_pre_post, "task")

    results_path = write_outputs(
        df,
        descriptives,
        invalid_values,
        ttests_pre_post,
        all_prepost_ttests,
        all_prepost_plots,
        prepost_without_group,
        prepost_with_group,
        predictors_pre,
        predictors_post,
        predictors_mix,
    )

    make_plot(
        df,
        "Planning",
        "planning_diff_pre",
        "planning_diff_post",
        OUT_DIR / "planning_diff_prepost_by_group.png",
        annotations=planning_pwm_annotations,
        annotation_key="Planning",
    )
    make_plot(
        df,
        "PWM",
        "pwm_diff_pre",
        "pwm_diff_post",
        OUT_DIR / "pwm_diff_prepost_by_group.png",
        annotations=planning_pwm_annotations,
        annotation_key="PWM",
    )

    complete_planning = df[["planning_diff_pre", "planning_diff_post"]].dropna().shape[0]
    complete_pwm = df[["pwm_diff_pre", "pwm_diff_post"]].dropna().shape[0]

    print("Planning Minefield analysis complete")
    print(f"Input N: {len(df)}")
    print(checks.to_string(index=False))
    print(f"Complete pairs Planning: {complete_planning}")
    print(f"Complete pairs PWM: {complete_pwm}")
    print(f"Results workbook: {results_path}")


if __name__ == "__main__":
    main()

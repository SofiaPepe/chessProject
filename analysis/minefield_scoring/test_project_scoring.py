from __future__ import annotations

import unittest
import csv
from pathlib import Path

from openpyxl import load_workbook

from scoring import (
    COMPLETE_DATABASE,
    OFFICIAL_DATABASE,
    canonical_project_id,
    is_legacy_minefield_column,
)
from shortest_path import parse_coordinate, shortest_path


class ProjectScoringUnitTests(unittest.TestCase):
    def test_project_id_normalization(self) -> None:
        self.assertEqual(canonical_project_id("ppr7"), "PPR07")
        self.assertEqual(canonical_project_id("SPR01"), "SPR1")
        self.assertEqual(canonical_project_id("AC 09"), "AC09")
        self.assertEqual(canonical_project_id("PROVA"), "")

    def test_shortest_path_uses_orthogonal_grid(self) -> None:
        start = parse_coordinate("A1")
        stop = parse_coordinate("C1")
        path = shortest_path(start, stop, {parse_coordinate("B1")})
        self.assertIsNotNone(path)
        self.assertEqual(len(path) - 1, 4)


class CompleteDatabaseIntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        if not Path(COMPLETE_DATABASE).exists():
            raise unittest.SkipTest("Run scoring.py before integration tests")

        source_book = load_workbook(
            OFFICIAL_DATABASE,
            read_only=True,
            data_only=True,
        )
        source_sheet = source_book[source_book.sheetnames[0]]
        cls.source_headers = [
            str(value) if value is not None else ""
            for value in next(source_sheet.iter_rows(values_only=True))
        ]
        cls.source_rows = [
            dict(zip(cls.source_headers, row))
            for row in source_sheet.iter_rows(min_row=2, values_only=True)
            if row[0]
        ]
        source_book.close()

        output_book = load_workbook(
            COMPLETE_DATABASE,
            read_only=True,
            data_only=True,
        )
        cls.output_sheet_names = output_book.sheetnames
        output_sheet = output_book[output_book.sheetnames[0]]
        cls.output_headers = list(
            next(output_sheet.iter_rows(values_only=True))
        )
        cls.output_rows = [
            dict(zip(cls.output_headers, row))
            for row in output_sheet.iter_rows(min_row=2, values_only=True)
            if row[0]
        ]
        output_book.close()
        update_path = Path(COMPLETE_DATABASE).parents[1] / "output" / "minefield_scoring" / "database_route_efficiency_updates.csv"
        with update_path.open("r", encoding="utf-8-sig", newline="") as handle:
            cls.update_rows = list(csv.DictReader(handle))

    def test_all_official_ids_are_retained_once(self) -> None:
        source_ids = [row["ID"] for row in self.source_rows]
        output_ids = [row["ID"] for row in self.output_rows]
        self.assertEqual(output_ids, source_ids)
        self.assertEqual(len(output_ids), 101)
        self.assertEqual(len(set(output_ids)), 101)

    def test_database_schema_matches_original(self) -> None:
        self.assertEqual(self.output_sheet_names, ["Foglio1"])
        self.assertEqual(self.output_headers, self.source_headers)
        self.assertEqual(len(self.output_headers), 104)
        self.assertFalse(any(str(header).startswith("MF_") for header in self.output_headers))

    def test_only_route_efficiency_columns_can_change(self) -> None:
        target_columns = {"PLANNING_PERC_PRE", "PLANNING_PERC_POST", "PWM_PERC_PRE", "PWM_PERC_POST"}
        output_by_id = {row["ID"]: row for row in self.output_rows}
        for source in self.source_rows:
            output = output_by_id[source["ID"]]
            for header in self.source_headers:
                if header in target_columns:
                    continue
                self.assertEqual(output[header], source[header])

    def test_update_audit_matches_database_and_preserves_missing_raw(self) -> None:
        output_by_id = {row["ID"]: row for row in self.output_rows}
        source_by_id = {row["ID"]: row for row in self.source_rows}
        self.assertEqual(len(self.update_rows), 404)
        for update in self.update_rows:
            participant_id = update["participant_id"]
            column = update["database_column"]
            actual = output_by_id[participant_id][column]
            expected = float(update["output_value"]) if update["output_value"] else None
            self.assertEqual(actual, expected)
            if update["source"] == "historical_original":
                self.assertEqual(actual, source_by_id[participant_id][column])


if __name__ == "__main__":
    unittest.main()
